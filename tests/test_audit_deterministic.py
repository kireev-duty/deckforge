"""audit/deterministic: каждая проверка ловит свою фикстуру-нарушение и молчит на чистой колоде."""

from __future__ import annotations

from pathlib import Path

import pytest

from deckforge.audit import audit_deck, report_markdown, summary
from deckforge.audit.deterministic import CHECKS, check_by_prefix
from tests.fixtures.bad_slides import FIXTURES, make_clean, make_L02_title_under_card

PREFIXES = [c[:3] for c in CHECKS]


def test_registry_matches_audit_md() -> None:
    assert len(CHECKS) == 26 and set(PREFIXES) == set(FIXTURES)
    assert PREFIXES == sorted(PREFIXES, key=lambda p: ("LTDIN".index(p[0]), p))


def test_clean_deck_has_no_findings(tmp_path: Path) -> None:
    case = make_clean(tmp_path)
    report = audit_deck(case.pptx, case.dna)
    assert report.checks_run == list(CHECKS)
    assert report.findings == [], report_markdown(report)
    assert summary(report)["errors"] == 0


@pytest.mark.parametrize("prefix", PREFIXES)
def test_check_catches_its_fixture(prefix: str, tmp_path: Path) -> None:
    case = FIXTURES[prefix](tmp_path)
    check_id, _ = check_by_prefix(prefix)
    report = audit_deck(case.pptx, case.dna, case.ir, checks=[prefix])
    hits = [f for f in report.findings if f.check_id == check_id]
    assert hits, f"{check_id}: нарушение не найдено\n{report_markdown(report)}"
    assert any(f.slide_idx == case.slide_idx for f in hits), [f.slide_idx for f in hits]
    for f in hits:
        assert f.kind == "deterministic" and f.message
        if f.slide_idx >= 0 and check_id not in ("D05_fill", "I06_duplicate_slides", "I03_empty_slide", "T04_layout",
                                                 "N01_notes_missing"):
            assert f.element_id and f.box is not None


def test_placeholder_dictionary_catches_template_notes() -> None:
    """Обращения шаблона к автору («не забудьте удалить этот слайд») — такие же заглушки, как lorem ipsum.

    Встретилось в `data/wild/шаблон-макет МТУСИ.pptx` в режиме «на входе только шаблон»:
    слайды-инструкции годятся как образцы, но их служебные надписи в колоду попадать не должны.
    """
    from deckforge.core.placeholders import is_placeholder_text

    assert is_placeholder_text("P.S. Не забудьте удалить этот слайд из финальной версии вашей презентации")
    assert is_placeholder_text("Удалите этот слайд перед показом")
    assert is_placeholder_text("Этот слайд нужно удалить")
    assert is_placeholder_text("Delete this slide before presenting")
    # не заглушки: обычный текст колоды про удаление данных и про слайды
    assert not is_placeholder_text("Удалите дубликаты записей перед загрузкой в витрину")
    assert not is_placeholder_text("Слайд с архитектурой решения")


def test_broken_file_reports_only_i01(tmp_path: Path) -> None:
    case = FIXTURES["I01"](tmp_path)
    report = audit_deck(case.pptx, case.dna)
    assert report.checks_run == ["I01_file"] and [f.check_id for f in report.findings] == ["I01_file"]
    assert report.findings[0].slide_idx == -1 and report.errors == 1


def test_l03_severity_depends_on_autofit(tmp_path: Path) -> None:
    from pptx import Presentation
    from pptx.enum.text import MSO_AUTO_SIZE

    case = FIXTURES["L03"](tmp_path)
    assert audit_deck(case.pptx, case.dna, checks=["L03"]).findings[0].severity == "error"
    prs = Presentation(str(case.pptx))
    for sh in prs.slides[0].shapes:
        if sh.has_text_frame and "согласование" in sh.text_frame.text:
            sh.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    prs.save(str(case.pptx))
    f = audit_deck(case.pptx, case.dna, checks=["L03"]).findings[0]
    assert f.severity == "warning" and f.evidence["autofit"] == "normAutofit"


def test_i03_empty_slide_is_error_and_title_only_is_warning_without_ir(tmp_path: Path) -> None:
    case = FIXTURES["I03"](tmp_path)
    by_slide = {f.slide_idx: f for f in audit_deck(case.pptx, case.dna, checks=["I03"]).findings}
    assert by_slide[0].severity == "warning" and by_slide[0].evidence["blocks"] == 1
    assert by_slide[1].severity == "error" and by_slide[1].evidence["blocks"] == 0


def test_summary_and_markdown(tmp_path: Path) -> None:
    case = FIXTURES["D01"](tmp_path)
    report = audit_deck(case.pptx, case.dna)
    s = summary(report)
    assert s["checks_run"] == 26 and s["by_check"].get("D01_bullets") == 1 and 0 in s["slides_with_issues"]
    md = report_markdown(report)
    assert "D01_bullets" in md and md.startswith("**D01.pptx**")


def test_l02_our_text_on_exemplar_block_is_error(tmp_path: Path) -> None:
    """Боксы заголовка и карточки пересекались ещё в образце — warning с in_exemplar; когда наш длинный
    заголовок лёг строкой на карточку — error: это уже не дизайн шаблона (VK WorkSpace slide5)."""
    case = make_L02_title_under_card(tmp_path)
    [f] = audit_deck(case.pptx, case.dna, case.ir, checks=["L02"]).findings
    assert f.severity == "error" and f.evidence["text_reached"] == 1 and f.evidence["in_exemplar"] == 0
    short = tmp_path / "L02_card_template.pptx"  # тот же образец с коротким заголовком
    [f] = audit_deck(short, case.dna, case.ir, checks=["L02"]).findings
    assert f.severity == "warning" and f.evidence["in_exemplar"] == 1 and f.evidence["text_reached"] == 0
