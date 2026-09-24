"""layout: planner, picker, fitting, builder."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from pptx import Presentation

from deckforge.core.ir import (
    Archetype,
    Box,
    ChartSpec,
    DeckOutline,
    Exemplar,
    ImageSpec,
    OutlineSlide,
    Slot,
    SlotKind,
    TableSpec,
)
from deckforge.core.strategy import Strategy, load_strategy
from deckforge.layout import build_deck_ir, plan
from deckforge.layout.exemplar_picker import candidate_archetypes, pick_exemplar
from deckforge.layout.fitting import fit_number, shorten, shorten_words, split_label_body
from deckforge.layout.planner import parse_num, table_to_chart
from deckforge.parsing.exemplars import load_exemplars
from deckforge.parsing.extract_tokens import extract_tokens
from deckforge.render import render_pptx

REPO = Path(__file__).resolve().parents[1]
OUTLINE = REPO / "examples" / "content_pack" / "outline.json"
SLIDE_W, SLIDE_H = 9144000, 5143500  # 10 × 5.625"


def outline() -> DeckOutline:
    return DeckOutline.model_validate_json(OUTLINE.read_text("utf-8"))


def strategy(**over) -> Strategy:
    base = dict(name="t", target_slides={"min": 8, "max": 12}, density={"max_bullets": 4, "max_words_per_bullet": 10, "max_fill_ratio": 0.7},
                archetype_priority=[Archetype.BULLETS, Archetype.CARDS])
    base.update(over)
    return Strategy.model_validate(base)


# ──────────────────────────── planner ────────────────────────────


def test_executive_turns_chart_into_table_and_drops_sections() -> None:
    res = plan(outline(), load_strategy("executive"), set(Archetype))
    archs = res.archetypes
    assert "section" not in archs and "chart" not in archs
    assert archs.count("table") == 2  # исходная таблица + таблица из графика
    assert load_strategy("executive").target_slides.min <= len(res.slides) <= load_strategy("executive").target_slides.max
    assert not res.warnings


def test_narrative_inserts_sections_and_splits_dense_slides() -> None:
    res = plan(outline(), load_strategy("narrative"), set(Archetype))
    archs = res.archetypes
    assert archs.count("section") >= 2 and archs[1] == "section"
    titles = [s.title for s in res.slides]
    assert any(t.endswith("(1/2)") for t in titles) and any(t.endswith("(2/2)") for t in titles)
    assert all(len(s.bullets) <= 4 for s in res.slides)
    assert len(res.slides) <= 15


def test_narrative_without_section_exemplars_has_no_sections() -> None:
    res = plan(outline(), load_strategy("narrative"), set(Archetype) - {Archetype.SECTION})
    assert "section" not in res.archetypes


def test_visual_turns_numeric_table_into_chart_but_keeps_text_table() -> None:
    res = plan(outline(), load_strategy("visual"), set(Archetype))
    assert res.archetypes.count("chart") == 2 and "table" not in res.archetypes
    text_table = TableSpec(header=["Роль", "Задача"], rows=[["Лид", "планирует"], ["Инженер", "делает"]])
    assert table_to_chart(text_table, "t") is None


def test_kpi_split_by_capacity_only_when_room() -> None:
    o = outline()
    roomy = plan(o, strategy(target_slides={"min": 8, "max": 15}, data_visualization={"key_metrics": "kpi"}), set(Archetype), kpi_capacity=2)
    assert roomy.archetypes.count("kpi") == 2
    tight = plan(o, strategy(target_slides={"min": 8, "max": 10}, data_visualization={"key_metrics": "kpi"}), set(Archetype), kpi_capacity=2)
    assert tight.archetypes.count("kpi") == 1


def test_kpi_split_parts_get_distinct_titles() -> None:
    res = plan(outline(), strategy(target_slides={"min": 8, "max": 15}, data_visualization={"key_metrics": "kpi"}), set(Archetype), kpi_capacity=2)
    kpi_titles = [s.title for s in res.slides if s.archetype == Archetype.KPI]
    assert len(kpi_titles) == len(set(kpi_titles)) and any(t.endswith("(1/2)") for t in kpi_titles)


def test_executive_compacts_real_outline_without_losing_words() -> None:
    """Outline из кассеты ужимается в 10–11 executive: шаги, KPI и цитата становятся карточками без потерь."""
    from deckforge.content.outline_writer import repair_outline
    from tests.conftest import cassette

    o, _ = repair_outline(cassette("outline_writer_pulse"), set(Archetype), set())
    assert len(o.slides) >= 12
    res = plan(o, load_strategy("executive"), set(Archetype), kpi_capacity=3)
    assert 10 <= len(res.slides) <= 11 and not res.warnings
    planned = " ".join(" ".join(s.bullets + s.steps + [k.value for k in s.kpis] + [k.label for k in s.kpis] + [s.quote or ""]) for s in res.slides)
    for s in o.slides:
        for word in " ".join(s.steps + [k.value for k in s.kpis] + [s.quote or ""]).split():
            assert word in planned, word
    # текстовые пары сливаются раньше KPI и цитаты
    vis = plan(o, load_strategy("visual"), set(Archetype), kpi_capacity=3)
    assert vis.archetypes.count("kpi") == sum(1 for s in o.slides if s.archetype == Archetype.KPI)


def test_parse_num() -> None:
    assert parse_num("1 250") == 1250 and parse_num("−89 %") == -89 and parse_num("4,7") == 4.7
    assert parse_num("6 ч") is None and parse_num("×3") == 3


# ──────────────────────────── picker ────────────────────────────


def _slot(id_: str, kind: SlotKind, x: float, y: float, w: float, h: float, max_chars: int | None = 100, size: float = 14) -> Slot:
    box = Box(x=int(x * SLIDE_W), y=int(y * SLIDE_H), w=int(w * SLIDE_W), h=int(h * SLIDE_H))
    return Slot(id=id_, kind=kind, box=box, max_chars=max_chars, max_lines=1 if kind != SlotKind.BODY else 8,
                max_items=8 if kind == SlotKind.BODY else None, size_pt=size)


def _exemplars() -> list[Exemplar]:
    title = _slot("t", SlotKind.TITLE, 0.05, 0.05, 0.9, 0.12, max_chars=60, size=24)
    bullets = Exemplar(id="slide1", source_index=0, layout_name="a", archetype=Archetype.BULLETS,
                       slots=[title, _slot("b", SlotKind.BODY, 0.05, 0.25, 0.9, 0.6, max_chars=600)])
    cards = Exemplar(id="slide2", source_index=1, layout_name="a", archetype=Archetype.CARDS, slots=[title] + [
        _slot(f"l{i}", SlotKind.LABEL, 0.05 + 0.3 * i, 0.3, 0.25, 0.05, max_chars=20) for i in range(3)
    ] + [_slot(f"b{i}", SlotKind.BODY, 0.05 + 0.3 * i, 0.36, 0.25, 0.4, max_chars=200) for i in range(3)])
    pic = Exemplar(id="slide3", source_index=2, layout_name="a", archetype=Archetype.IMAGE_TEXT,
                   slots=[title, _slot("b", SlotKind.BODY, 0.5, 0.25, 0.45, 0.6, max_chars=400),
                          _slot("p", SlotKind.PICTURE, 0.05, 0.25, 0.4, 0.6, max_chars=None)])
    sec = Exemplar(id="slide4", source_index=3, layout_name="a", archetype=Archetype.TITLE,
                   slots=[_slot("t", SlotKind.TITLE, 0.1, 0.4, 0.8, 0.2, max_chars=60, size=40)])
    return [bullets, cards, pic, sec]


def test_priority_switches_between_bullets_and_cards() -> None:
    slide = OutlineSlide(idx=1, archetype=Archetype.BULLETS, title="Три тезиса", bullets=["А — а", "Б — б", "В — в"])
    ex = _exemplars()
    e_b, _ = pick_exemplar(slide, ex, strategy(archetype_priority=[Archetype.BULLETS, Archetype.CARDS]), SLIDE_W * SLIDE_H)
    e_c, _ = pick_exemplar(slide, ex, strategy(archetype_priority=[Archetype.CARDS, Archetype.BULLETS]), SLIDE_W * SLIDE_H)
    assert e_b.archetype == Archetype.BULLETS and e_c.archetype == Archetype.CARDS


def test_section_falls_back_to_title_exemplar() -> None:
    slide = OutlineSlide(idx=1, archetype=Archetype.SECTION, title="Контекст")
    e, _ = pick_exemplar(slide, _exemplars(), strategy(), SLIDE_W * SLIDE_H)
    assert e is not None and e.archetype == Archetype.TITLE
    assert candidate_archetypes(Archetype.SECTION, strategy())[0] == Archetype.SECTION


def test_minimal_images_avoid_picture_exemplar_without_image(tmp_path) -> None:
    slide = OutlineSlide(idx=1, archetype=Archetype.IMAGE_TEXT, title="Т", bullets=["один", "два"])
    e, _ = pick_exemplar(slide, _exemplars(), strategy(images="minimal", archetype_priority=[Archetype.IMAGE_TEXT, Archetype.BULLETS]),
                         SLIDE_W * SLIDE_H)
    assert e.archetype != Archetype.IMAGE_TEXT
    slide.image = ImageSpec(path=str(tmp_path / "missing.png"))  # файла нет — как без картинки
    e_missing, _ = pick_exemplar(slide, _exemplars(), strategy(images="minimal", archetype_priority=[Archetype.IMAGE_TEXT, Archetype.BULLETS]),
                                 SLIDE_W * SLIDE_H)
    assert e_missing.archetype != Archetype.IMAGE_TEXT
    (tmp_path / "x.png").write_bytes(b"png")
    slide.image = ImageSpec(path=str(tmp_path / "x.png"))
    e2, _ = pick_exemplar(slide, _exemplars(), strategy(images="minimal", archetype_priority=[Archetype.IMAGE_TEXT, Archetype.BULLETS]),
                          SLIDE_W * SLIDE_H)
    assert e2.archetype == Archetype.IMAGE_TEXT


def test_chart_without_data_slot_goes_to_body_or_nowhere() -> None:
    """Без data-слота диаграмма встаёт на место самого крупного текстового блока, но не на структурные образцы."""
    slide = OutlineSlide(idx=1, archetype=Archetype.CHART, title="Т",
                         chart=ChartSpec(kind="bar", title="c", categories=["a"], series={"s": [1]}))
    e, score = pick_exemplar(slide, _exemplars(), strategy(), SLIDE_W * SLIDE_H)
    assert e is not None and e.archetype == Archetype.BULLETS and score < 0
    title_only = [x for x in _exemplars() if x.archetype == Archetype.TITLE]
    assert pick_exemplar(slide, title_only, strategy(), SLIDE_W * SLIDE_H)[0] is None
    from deckforge.layout.builder import build_slide

    ir, leftover = build_slide(0, slide, e, SLIDE_H, {})
    assert leftover is None and [el.kind for el in ir.elements if el.chart] == [SlotKind.CHART]
    assert ir.elements[-1].slot_id == "b"


def test_quote_needs_body_or_big_title() -> None:
    """Цитата на образец без body — отказ; section/title годятся (цитата уходит в крупный заголовок)."""
    quote = OutlineSlide(idx=1, archetype=Archetype.QUOTE, title="Т", quote="За месяц мы впервые увидели, куда уходит время",
                         quote_author="Лид")
    title = _slot("t", SlotKind.TITLE, 0.05, 0.05, 0.9, 0.12, max_chars=60, size=24)
    ladder = Exemplar(id="s27", source_index=26, layout_name="a", archetype=Archetype.BULLETS,
                      slots=[title] + [_slot(f"l{i}", SlotKind.LABEL, 0.05 + 0.1 * i, 0.3 + 0.15 * i, 0.5, 0.1, max_chars=40) for i in range(4)])
    e, score = pick_exemplar(quote, [ladder], strategy(), SLIDE_W * SLIDE_H)
    assert e is None
    e, _ = pick_exemplar(quote, [ladder] + _exemplars(), strategy(), SLIDE_W * SLIDE_H)
    assert e is not None and e.id != "s27"
    sec = [x for x in _exemplars() if x.archetype == Archetype.TITLE]
    e, _ = pick_exemplar(quote, sec, strategy(), SLIDE_W * SLIDE_H)
    assert e is not None and e.archetype == Archetype.TITLE


def test_quote_prefers_section_over_tiny_image_text() -> None:
    """Для цитаты section не «последний резерв», а первый после quote-образца."""
    quote = OutlineSlide(idx=1, archetype=Archetype.QUOTE, title="Т", quote_author="Лид",
                         quote="За месяц мы впервые увидели, где реально теряется время, — и перестали спорить об этом на ретро")
    title = _slot("t", SlotKind.TITLE, 0.05, 0.05, 0.9, 0.12, max_chars=60, size=24)
    mockup = Exemplar(id="mock", source_index=38, layout_name="a", archetype=Archetype.IMAGE_TEXT,
                      slots=[title, _slot("lab", SlotKind.LABEL, 0.05, 0.25, 0.3, 0.03, max_chars=16)]
                      + [_slot(f"b{i}", SlotKind.BODY, 0.05, 0.3 + 0.05 * i, 0.3, 0.04, max_chars=21) for i in range(10)])
    section = Exemplar(id="sec", source_index=2, layout_name="a", archetype=Archetype.SECTION,
                       slots=[_slot("t", SlotKind.TITLE, 0.1, 0.35, 0.8, 0.2, max_chars=80, size=36),
                              _slot("s", SlotKind.SUBTITLE, 0.1, 0.6, 0.8, 0.08, max_chars=60)])
    used = {"sec": 2}
    e, _ = pick_exemplar(quote, [mockup, section], strategy(), SLIDE_W * SLIDE_H, used=used)
    assert e is not None and e.id == "sec"
    from deckforge.layout.builder import build_slide

    ir, _ = build_slide(0, quote, section, SLIDE_H, {})  # цитата в заголовок, автор в подзаголовок
    texts = {el.slot_id: "".join(r.text for p in el.paragraphs for r in p.runs) for el in ir.elements}
    assert texts["t"].startswith("«За месяц") and texts["s"] == "Лид", texts
    q = Exemplar(id="q", source_index=5, layout_name="a", archetype=Archetype.QUOTE,
                 slots=[_slot("b", SlotKind.BODY, 0.1, 0.3, 0.8, 0.3, max_chars=300), _slot("c", SlotKind.CAPTION, 0.1, 0.65, 0.8, 0.05)])
    e, _ = pick_exemplar(quote, [mockup, section, q], strategy(), SLIDE_W * SLIDE_H, used=used)
    assert e is not None and e.id == "q"


def test_kpis_fall_back_to_bullets_when_no_numbers_or_cards() -> None:
    kpi = OutlineSlide(idx=1, archetype=Archetype.KPI, title="Т",
                       kpis=[{"value": "42%", "label": "перегружены"}, {"value": "12", "label": "команд"}])
    bullets = [x for x in _exemplars() if x.archetype == Archetype.BULLETS]
    e, score = pick_exemplar(kpi, bullets, strategy(), SLIDE_W * SLIDE_H)
    assert e is not None and score > -10
    from deckforge.layout.builder import build_slide

    ir, left = build_slide(0, kpi, e, SLIDE_H, {})
    assert left is None
    body = next(el for el in ir.elements if el.kind == SlotKind.BODY)
    assert [p.runs[0].text for p in body.paragraphs] == ["42% — перегружены", "12 — команд"]


def test_paragraphs_are_laid_out_like_bullets() -> None:
    """Абзац — такой же контент: без этого слайд из режима «только шаблон» уезжал с одним заголовком."""
    from deckforge.layout.builder import build_slide

    slide = OutlineSlide(idx=1, archetype=Archetype.CARDS, title="Т",
                         paragraphs=["Первый абзац про продукт", "Второй абзац про внедрение"])
    cards = [x for x in _exemplars() if x.archetype == Archetype.CARDS]
    e, _ = pick_exemplar(slide, cards, strategy(), SLIDE_W * SLIDE_H)
    assert e is not None
    ir, left = build_slide(0, slide, e, SLIDE_H, {})

    placed = [" ".join(r.text for p in el.paragraphs for r in p.runs) for el in ir.elements
              if el.kind in (SlotKind.BODY, SlotKind.LABEL)]
    assert any("Первый абзац" in t for t in placed), placed
    # маркеров у абзацев нет
    assert all(not p.bullet for el in ir.elements for p in el.paragraphs)
    if left is not None:  # остаток возвращается абзацами, а не буллетами
        assert left.paragraphs and not left.bullets


# ──────────────────────────── fitting ────────────────────────────


def test_shorten_prefers_separators_then_words() -> None:
    assert shorten("Статусы вместо работы — до 9,5 часов в неделю", 30) == "Статусы вместо работы"
    assert shorten("Очень длинное предложение без разделителей внутри", 25).endswith("…")
    assert shorten("короткий", 100) == "короткий"
    # одно-два слова в тесном слоте не режем
    assert shorten("Октябрь", 5) == "Октябрь"
    assert shorten("0,6 дня", 5) == "0,6 дня"
    assert shorten("Подключение трекера и календаря", 8) == "Подключение…"
    # висящий предлог перед многоточием не оставляем
    assert shorten("Бесшовная интеграция с текущими системами", 23) == "Бесшовная интеграция…"
    assert shorten_words("Риски заранее — дайджест приходит в понедельник, до планирования недели", 5) == "Риски заранее"


def test_list_element_shrinks_font_before_cutting() -> None:
    from deckforge.layout.builder import list_element

    slot = _slot("b", SlotKind.BODY, 0.1, 0.2, 0.8, 0.6, max_chars=120, size=18)
    slot.max_items = 4
    items = ["Статусы съедают до 9,5 часов в неделю у каждого разработчика", "Трекер показывает задачи, а не людей",
             "О выгорании узнают слишком поздно"]
    el = list_element(slot, items, {}, bullet=True)
    texts = [p.runs[0].text for p in el.paragraphs]
    assert texts == items, "три пункта по 40 знаков при 120 знаках слота влезают после уменьшения кегля"
    assert 18 * 0.7 <= el.style_overrides["size_pt"] < 18
    # минимальный кегль не спасает — режем по разделителю
    el = list_element(slot, items + ["Ручные отчёты устаревают за день — к утру данные уже неверны"] * 3, {}, bullet=True)
    assert el.style_overrides["size_pt"] == round(18 * 0.7, 1)
    assert all(not p.runs[0].text.endswith("…") or " " in p.runs[0].text for p in el.paragraphs)


def test_number_element_label_inside_shape() -> None:
    """Цифра без label-слота: подпись вторым абзацем, цифра ужимается; в однострочный бокс подпись не лезет."""
    from deckforge.layout.builder import LABEL_MAX_PT, number_element

    tall = _slot("n", SlotKind.NUMBER, 0.1, 0.2, 0.3, 0.15, max_chars=6, size=60)  # ≈ 61 pt высоты
    el = number_element(tall, "118", {}, label="участников пилота")
    assert [p.runs[0].text for p in el.paragraphs] == ["118", "участников пилота"]
    assert el.paragraphs[1].runs[0].size_pt <= LABEL_MAX_PT
    assert (el.paragraphs[0].runs[0].size_pt or 60) < 60
    flat = _slot("n", SlotKind.NUMBER, 0.1, 0.2, 0.3, 0.05, max_chars=6, size=28)  # ≈ 20 pt высоты
    el = number_element(flat, "118", {}, label="участников пилота")
    assert len(el.paragraphs) == 1


def test_kpi_labels_pair_with_numbers_by_geometry() -> None:
    """Подпись KPI — label под своей цифрой, а не по порядку чтения."""
    from deckforge.layout.builder import build_slide

    title = _slot("t", SlotKind.TITLE, 0.05, 0.05, 0.9, 0.12, max_chars=60, size=24)
    nums = [_slot(f"n{i}", SlotKind.NUMBER, 0.05 + 0.32 * i, 0.5, 0.28, 0.15, max_chars=8, size=48) for i in range(3)]
    labels = [_slot("l0a", SlotKind.LABEL, 0.05, 0.66, 0.28, 0.06, max_chars=40),
              _slot("l0b", SlotKind.LABEL, 0.05, 0.74, 0.28, 0.06, max_chars=40),
              _slot("l1", SlotKind.LABEL, 0.37, 0.66, 0.28, 0.06, max_chars=40)]
    kpi_ex = Exemplar(id="k", source_index=0, layout_name="a", archetype=Archetype.KPI, slots=[title] + nums + labels)
    kpis = [{"value": "42%", "label": "команд перегружены"}, {"value": "9,5 ч/нед", "label": "на статусы и встречи"},
            {"value": "1,8 дня", "label": "средний срок ответа"}]
    ir, left = build_slide(0, OutlineSlide(idx=1, archetype=Archetype.KPI, title="Т", kpis=kpis), kpi_ex, SLIDE_H, {})
    assert left is None
    texts = {el.slot_id: ["".join(r.text for r in p.runs) for p in el.paragraphs] for el in ir.elements if el.paragraphs}
    assert texts["l0a"] == ["команд перегружены"] and texts["l1"] == ["на статусы и встречи"], texts
    assert "l0b" not in texts  # вторая строка под первой цифрой пустая
    assert texts["n2"] == ["1,8 дня", "средний срок ответа"]  # подпись внутри фигуры


def test_kpis_in_cards_stay_in_their_card() -> None:
    """Карточки под KPI: значение → label, подпись → body той же карточки по геометрии, а не по порядку чтения."""
    from deckforge.layout.builder import build_slide

    title = _slot("t", SlotKind.TITLE, 0.05, 0.05, 0.9, 0.12, max_chars=60, size=24)
    slots = [title]
    for i in range(4):  # label чуть ниже у чётных карточек — порядок чтения перемешан
        x = 0.05 + 0.23 * i
        slots.append(_slot(f"l{i}", SlotKind.LABEL, x, 0.30 + (0.03 if i % 2 else 0), 0.2, 0.06, max_chars=12))
        slots.append(_slot(f"b{i}", SlotKind.BODY, x, 0.40, 0.2, 0.3, max_chars=120))
    cards = Exemplar(id="c", source_index=0, layout_name="a", archetype=Archetype.CARDS, slots=slots)
    kpis = [{"value": "42%", "label": "перегружены"}, {"value": "12", "label": "команд"},
            {"value": "1,8 дня", "label": "срок ответа"}, {"value": "4,7", "label": "оценка пилота"}]
    ir, left = build_slide(0, OutlineSlide(idx=1, archetype=Archetype.KPI, title="Т", kpis=kpis), cards, SLIDE_H, {})
    assert left is None
    by_slot = {el.slot_id: el.paragraphs[0].runs[0].text for el in ir.elements}
    for i, k in enumerate(kpis):
        assert by_slot[f"l{i}"] == k["value"] and by_slot[f"b{i}"] == k["label"], (i, by_slot)


def _mixed_cards_exemplar() -> Exemplar:
    """Два буллета сверху без label + «лестница» из 4 ступенек, у первой — label без body."""
    title = _slot("t", SlotKind.TITLE, 0.05, 0.05, 0.9, 0.12, max_chars=60, size=24)
    slots = [title, _slot("top0", SlotKind.BODY, 0.1, 0.25, 0.8, 0.08), _slot("top1", SlotKind.BODY, 0.1, 0.35, 0.8, 0.08),
             _slot("l0", SlotKind.LABEL, 0.05, 0.90, 0.2, 0.06, max_chars=12)]
    for i in range(1, 4):
        x, y = 0.05 + 0.23 * i, 0.75 - 0.08 * i
        slots.append(_slot(f"l{i}", SlotKind.LABEL, x, y, 0.2, 0.06, max_chars=12))
        slots.append(_slot(f"b{i}", SlotKind.BODY, x, y + 0.06, 0.2, 0.15, max_chars=80))
    return Exemplar(id="mixed", source_index=0, layout_name="a", archetype=Archetype.CARDS, slots=slots)


def test_pair_labels_keeps_foreign_labels_for_unpaired_bodies() -> None:
    """Часть карточек спарилась по геометрии — одинокий label не уезжает к чужому body."""
    from deckforge.layout.builder import build_slide, pair_labels

    ex = _mixed_cards_exemplar()
    bodies = [s for s in ex.slots if s.kind == SlotKind.BODY]
    labels = [s for s in ex.slots if s.kind == SlotKind.LABEL]
    paired = pair_labels(bodies, labels)
    assert {b: l.id for b, l in paired.items()} == {"b1": "l1", "b2": "l2", "b3": "l3"}

    items = ["1.8 — дня (до пилота)", "0.6 — дня (после пилота)", "3.4 — оценка до", "4.7 — оценка после", "«ц» — Автор"]
    ir, left = build_slide(0, OutlineSlide(idx=1, archetype=Archetype.CARDS, title="Т", bullets=items), ex, SLIDE_H, {})
    assert left is None
    by_slot = {el.slot_id: el.paragraphs[0].runs[0].text for el in ir.elements if el.paragraphs}
    assert by_slot["top0"] == "1.8 — дня (до пилота)" and by_slot["top1"] == "0.6 — дня (после пилота)"
    assert "l0" not in by_slot
    assert by_slot["l3"] == "3.4" and by_slot["b3"] == "оценка до"


def test_pair_labels_order_fallback_without_geometry() -> None:
    """Без единой геометрической пары label раздаются по порядку, как раньше."""
    from deckforge.layout.builder import pair_labels

    bodies = [_slot(f"b{i}", SlotKind.BODY, 0.05 + 0.3 * i, 0.3, 0.25, 0.3) for i in range(2)]
    labels = [_slot(f"l{i}", SlotKind.LABEL, 0.05 + 0.3 * i, 0.8, 0.25, 0.06) for i in range(2)]
    assert {b: l.id for b, l in pair_labels(bodies, labels).items()} == {"b0": "l0", "b1": "l1"}


def test_long_lead_without_rest_goes_to_body_not_label() -> None:
    """Цитата длиннее плашки карточки — пункт целиком в тело, плашка пустая."""
    from deckforge.layout.builder import build_slide

    ex = _mixed_cards_exemplar()
    quote = "«За месяц мы впервые увидели, где реально теряется время, — и перестали спорить об этом на ретро»"
    for item in (quote, f"{quote} — Руководитель разработки"):
        ir, _ = build_slide(0, OutlineSlide(idx=1, archetype=Archetype.CARDS, title="Т", bullets=["1 — а", "2 — б", item]), ex, SLIDE_H, {})
        texts = {el.slot_id: "".join(r.text for p in el.paragraphs for r in p.runs) for el in ir.elements if el.paragraphs}
        assert texts["b3"] == item and "l3" not in texts, texts


def test_lone_leftover_goes_to_continuation_not_dropped() -> None:
    """Пятый KPI в образец на четыре цифры уходит на слайд-продолжение."""
    title = _slot("t", SlotKind.TITLE, 0.05, 0.05, 0.9, 0.12, max_chars=60, size=24)
    nums = [_slot(f"n{i}", SlotKind.NUMBER, 0.05 + 0.23 * i, 0.3, 0.2, 0.2, max_chars=6, size=48) for i in range(4)]
    kpi_ex = Exemplar(id="k", source_index=0, layout_name="a", archetype=Archetype.KPI, slots=[title] + nums)
    kpis = [{"value": str(i), "label": f"п{i}"} for i in range(5)]
    o = DeckOutline(title="t", purpose="other", slides=[
        OutlineSlide(idx=0, archetype=Archetype.TITLE, title="Т"),
        OutlineSlide(idx=1, archetype=Archetype.KPI, title="Пять цифр", kpis=kpis),
    ])
    res = build_deck_ir(o, strategy(target_slides={"min": 3, "max": 20}), [kpi_ex] + _exemplars(), "t", SLIDE_W, SLIDE_H)
    assert not [w for w in res.warnings if "отброшен" in w]
    values = [r.text for s in res.ir.slides for el in s.elements if el.kind == SlotKind.NUMBER for p in el.paragraphs for r in p.runs]
    assert "4" in values, values


def test_over_budget_kpi_compacts_into_list_instead_of_cascade() -> None:
    """Колода сверх объёма: 8 KPI на образце с четырьмя цифрами сворачиваются в список на одном слайде."""
    title = _slot("t", SlotKind.TITLE, 0.05, 0.05, 0.9, 0.12, max_chars=60, size=24)
    nums = [_slot(f"n{i}", SlotKind.NUMBER, 0.05 + 0.23 * i, 0.3, 0.2, 0.2, max_chars=6, size=48) for i in range(4)]
    kpi_ex = Exemplar(id="k", source_index=0, layout_name="a", archetype=Archetype.KPI, slots=[title] + nums)
    kpis = [{"value": str(i), "label": f"п{i}"} for i in range(8)]
    slides = [OutlineSlide(idx=0, archetype=Archetype.TITLE, title="Т")]
    slides += [OutlineSlide(idx=i, archetype=Archetype.KPI, title=f"Метрики {i}", kpis=kpis) for i in range(1, 7)]
    o = DeckOutline(title="t", purpose="other", slides=slides)
    res = build_deck_ir(o, strategy(target_slides={"min": 3, "max": 5}), [kpi_ex] + _exemplars(), "t", SLIDE_W, SLIDE_H)
    assert len(res.ir.slides) == 7, [s.exemplar_id for s in res.ir.slides]
    assert all("свёрнут" in w for w in res.warnings if "Метрики" in w), res.warnings
    texts = " ".join(r.text for s in res.ir.slides for el in s.elements for p in el.paragraphs for r in p.runs)
    assert all(f"п{i}" in texts for i in range(8))
    assert res.plan.slides[1].archetype == Archetype.KPI  # план стратегии не переписан
    # в пределах объёма — крупные цифры и продолжение
    res2 = build_deck_ir(o, strategy(target_slides={"min": 3, "max": 25}), [kpi_ex] + _exemplars(), "t", SLIDE_W, SLIDE_H)
    assert len(res2.ir.slides) >= 13 and not [w for w in res2.warnings if "свёрнут" in w]


def test_over_budget_process_compacts_into_numbered_list() -> None:
    title = _slot("t", SlotKind.TITLE, 0.05, 0.05, 0.9, 0.12, max_chars=60, size=24)
    proc = Exemplar(id="p", source_index=0, layout_name="a", archetype=Archetype.PROCESS, slots=[title] + [
        _slot(f"s{i}", SlotKind.BODY, 0.05 + 0.23 * i, 0.4, 0.2, 0.3, max_chars=80) for i in range(4)])
    steps = [f"Шаг {i}: сделать" for i in range(12)]
    slides = [OutlineSlide(idx=0, archetype=Archetype.TITLE, title="Т")]
    slides += [OutlineSlide(idx=i, archetype=Archetype.PROCESS, title=f"Процесс {i}", steps=steps) for i in range(1, 7)]
    o = DeckOutline(title="t", purpose="other", slides=slides)
    res = build_deck_ir(o, strategy(target_slides={"min": 3, "max": 5}), [proc] + _exemplars(), "t", SLIDE_W, SLIDE_H)
    assert len(res.ir.slides) == 7, [s.exemplar_id for s in res.ir.slides]
    assert not [w for w in res.warnings if "продолжение" in w]
    texts = " ".join(r.text for s in res.ir.slides for el in s.elements for p in el.paragraphs for r in p.runs)
    assert "Шаг 11" in texts and "12." in texts


def test_continuation_title_suffix_not_stacked() -> None:
    title = _slot("t", SlotKind.TITLE, 0.05, 0.05, 0.9, 0.12, max_chars=60, size=24)
    nums = [_slot(f"n{i}", SlotKind.NUMBER, 0.05 + 0.3 * i, 0.3, 0.2, 0.2, max_chars=6, size=48) for i in range(2)]
    kpi_ex = Exemplar(id="k", source_index=0, layout_name="a", archetype=Archetype.KPI, slots=[title] + nums)
    kpis = [{"value": str(i), "label": f"п{i}"} for i in range(8)]
    o = DeckOutline(title="t", purpose="other", slides=[OutlineSlide(idx=1, archetype=Archetype.KPI, title="Восемь", kpis=kpis)])
    res = build_deck_ir(o, strategy(target_slides={"min": 3, "max": 3}), [kpi_ex], "t", SLIDE_W, SLIDE_H)
    titles = [r.text for s in res.ir.slides for el in s.elements if el.kind == SlotKind.TITLE for p in el.paragraphs for r in p.runs]
    assert titles == ["Восемь"] + ["Восемь (продолжение)"] * 3, titles


def test_split_label_body_and_numbers() -> None:
    assert split_label_body("Лид — пояснение текста") == ("Лид", "пояснение текста")
    assert split_label_body("Просто пункт") == ("Просто пункт", "")
    assert split_label_body("«А — б» — Автор") == ("«А — б»", "Автор")
    assert split_label_body("«А — б»") == ("«А — б»", "")
    slot = Slot(id="n", kind=SlotKind.NUMBER, box=Box(x=0, y=0, w=1, h=1), max_chars=2, max_lines=1, size_pt=160)
    num, unit, size = fit_number("1,8 дня", slot)
    assert (num, unit) == ("1,8", "дня") and size is not None and size < 160
    assert fit_number("42 %", slot)[:2] == ("42%", "")
    # пустое / нечисловое значение KPI
    assert fit_number("", slot) == ("", "", None)
    assert fit_number("   ", slot) == ("", "", None)
    assert fit_number("—", slot)[0] == "—" and fit_number("N/A", slot)[0] == "N/A"


# ──────────────────────────── builder e2e ────────────────────────────


@pytest.mark.parametrize("name", ["executive", "narrative", "visual"])
def test_build_and_render_on_vk_tech(template_path, tmp_path, name: str) -> None:
    pptx = template_path("VK Tech")
    exemplars = load_exemplars(pptx)
    tokens = extract_tokens(pptx)
    st = load_strategy(name)
    res = build_deck_ir(outline(), st, exemplars, tokens.template_id, tokens.slide_w, tokens.slide_h)
    ir = res.ir
    assert ir.strategy == st.id
    assert st.target_slides.min <= len(ir.slides) <= st.target_slides.max, res.warnings
    ids = {e.id for e in exemplars}
    assert all(s.exemplar_id in ids for s in ir.slides)
    assert not any(c.exemplar_id is None for c in res.choices)
    # весь контент разложен
    dropped = [w for w in res.warnings if "отброшен" in w]
    assert not dropped, dropped
    out = render_pptx(ir, pptx, exemplars, tmp_path / f"{name}.pptx")
    prs = Presentation(str(out))
    assert len(prs.slides) == len(ir.slides)
    kinds = json.loads(ir.model_dump_json())["slides"]
    assert kinds


def test_strategies_differ_on_same_content(template_path) -> None:
    pptx = template_path("VK Tech")
    exemplars = load_exemplars(pptx)
    tokens = extract_tokens(pptx)
    seqs = {}
    for name in ("executive", "narrative", "visual"):
        res = build_deck_ir(outline(), load_strategy(name), exemplars, tokens.template_id, tokens.slide_w, tokens.slide_h)
        seqs[name] = res.plan.archetypes
    assert len(seqs["executive"]) < len(seqs["narrative"])
    assert "chart" not in seqs["executive"] and "chart" in seqs["visual"]
    assert "section" in seqs["narrative"] and "section" not in seqs["visual"] and "section" not in seqs["executive"]


# ──────────────────── режим «только шаблон»: текст без данных на VK Tech ────────────────────


def test_structural_title_shrinks_before_cut() -> None:
    """Заголовок финала 51 знак в слоте на 17: ужимается до STRUCTURAL_TITLE_MIN_SCALE, а не «…»."""
    from deckforge.layout.builder import text_element
    from deckforge.layout.fitting import STRUCTURAL_TITLE_MIN_SCALE

    slot = Slot(id="1", kind=SlotKind.TITLE, box=Box(x=0, y=0, w=4_373_113, h=1_085_547), max_chars=17, max_lines=1,
                size_pt=36.0)
    text = "Следующий шаг: консультация для оценки применимости"
    el = text_element(slot, text, {}, min_scale=STRUCTURAL_TITLE_MIN_SCALE)
    assert el.paragraphs[0].runs[0].text == text
    assert 36.0 * STRUCTURAL_TITLE_MIN_SCALE <= el.style_overrides["size_pt"] < 36.0


def test_shorten_title_keeps_lead_before_colon() -> None:
    from deckforge.layout.fitting import shorten_title

    assert shorten_title("VK Tech: корпоративные ИТ-решения для цифровой трансформации", 20) == "VK Tech"
    assert shorten_title("Короткий", 20) == "Короткий"
    assert shorten_title("Заголовок без разделителей и очень длинный хвост", 20).endswith("…")


def _template_only_outline() -> DeckOutline:
    """Как outline из template_brief: длинные заголовки, по 3 пункта, абзацы без картинок, данных нет."""
    bullets = ["Рост сложности ИТ-ландшафта создаёт риски для процессов", "Разрозненность систем снижает эффективность команд",
               "Увеличение операционных расходов из-за дублирования функций"]
    para = ("Сохранение накопленных данных и отлаженных процессов. API-шлюзы и коннекторы для лёгкого подключения "
            "к текущему стеку.")
    slides = [OutlineSlide(idx=0, archetype=Archetype.TITLE, title="VK Tech: корпоративные ИТ-решения для цифровой трансформации",
                           subtitle="Единая экосистема для автоматизации и масштабирования бизнеса")]
    for i in range(1, 13):  # карточек больше, чем повторов одного образца до потолка штрафа
        if i % 3:
            slides.append(OutlineSlide(idx=i, archetype=Archetype.CARDS, title=f"Тезис {i}: почему это важно", bullets=bullets))
        else:
            slides.append(OutlineSlide(idx=i, archetype=Archetype.IMAGE_TEXT, title="Бесшовная интеграция с существующими системами",
                                       paragraphs=[para]))
    slides.append(OutlineSlide(idx=13, archetype=Archetype.CLOSING, title="Следующий шаг: консультация для оценки применимости"))
    return DeckOutline(title="VK Tech", purpose="other", slides=slides)


@pytest.mark.parametrize("name", ["executive", "narrative", "visual"])
def test_template_only_text_fits_without_empty_cards(template_path, name: str) -> None:
    """D1–D3 на VK Tech: нет «…», сетка не заполнена меньше чем на половину и без пустых карточек,
    заголовок не уходит под контентный блок (slide31), подписи не пишутся в картинки."""
    from deckforge.layout.exemplar_picker import _overlaps_content, _visible_title_chars
    from deckforge.layout.fitting import MIN_SIZE_SCALE

    pptx = template_path("VK Tech")
    exemplars = load_exemplars(pptx)
    tokens = extract_tokens(pptx)
    res = build_deck_ir(_template_only_outline(), load_strategy(name), exemplars, tokens.template_id,
                        tokens.slide_w, tokens.slide_h)
    by_id = {e.id: e for e in exemplars}
    for s in res.ir.slides:
        texts = ["".join(r.text for r in p.runs) for e in s.elements for p in e.paragraphs]
        assert not any(t.endswith("…") for t in texts), (s.idx, s.exemplar_id, texts)
        ex = by_id[s.exemplar_id]
        bodies = [x for x in ex.slots if x.kind == SlotKind.BODY]
        filled = {e.slot_id for e in s.elements}
        if len(bodies) > 1 and not ex.card_frames:
            # сетка нарисована общим фоном (slide19 — картинкой лейаута): пустая карточка осталась бы видна;
            # карточка заполнена, если занят её body или подпись
            texts_filled = sum(1 for x in ex.slots if x.kind in (SlotKind.BODY, SlotKind.LABEL) and x.id in filled)
            assert texts_filled >= len(bodies), (s.idx, s.exemplar_id, "пустая карточка")
        title = next((x for x in ex.slots if x.kind == SlotKind.TITLE), None)
        if title is not None and s.archetype not in (Archetype.TITLE, Archetype.CLOSING, Archetype.SECTION) \
                and _overlaps_content(title, ex.slots):
            placed = next(t for e, t in zip(s.elements, texts) if e.kind == SlotKind.TITLE)
            assert len(placed) <= _visible_title_chars(title, ex.slots, MIN_SIZE_SCALE), (s.idx, s.exemplar_id)


def test_title_drops_empty_speaker_photo_frame(template_path, tmp_path) -> None:
    """Титул VK Tech: рамка под фото докладчика с подсказкой «Вставить фото» без картинки убирается,
    и T05 не считает это удалением фиксированного элемента."""
    from deckforge.audit import audit_deck
    from deckforge.parsing.dna import build_dna

    pptx = template_path("VK Tech")
    dna = build_dna(pptx)
    outline = DeckOutline(title="VK Tech", purpose="other", slides=[_template_only_outline().slides[0]])
    res = build_deck_ir(outline, load_strategy("executive"), dna.exemplars, dna.template_id, dna.slide_w, dna.slide_h)
    assert res.ir.slides[0].exemplar_id == "slide2"
    out = render_pptx(res.ir, pptx, dna.exemplars, tmp_path / "title.pptx")
    ids = {str(sh.shape_id) for sh in Presentation(str(out)).slides[0].shapes}
    assert not ids & {"411", "412"}, ids
    report = audit_deck(out, dna, res.ir, checks=["T05"])
    assert not report.findings, [f.message for f in report.findings]


def test_image_text_without_picture_merges_to_fit_volume(tmp_path) -> None:
    """Outline из шаблона чередует cards и image_text по 3 пункта: без готовой картинки image_text — такой же
    текст и сливается, иначе executive не укладывается в свой диапазон; с картинкой разворот не трогаем."""
    bullets = ["Пункт первый про платформу", "Пункт второй про внедрение", "Пункт третий про поддержку"]
    body = [OutlineSlide(idx=i, archetype=Archetype.CARDS if i % 2 else Archetype.IMAGE_TEXT, title=f"Тезис {i}",
                         bullets=list(bullets)) for i in range(1, 11)]
    slides = [OutlineSlide(idx=0, archetype=Archetype.TITLE, title="Т"), *body,
              OutlineSlide(idx=11, archetype=Archetype.CLOSING, title="Ф")]
    st = load_strategy("executive")
    res = plan(DeckOutline(title="Т", purpose="other", slides=slides), st, set(Archetype))
    assert st.target_slides.min <= len(res.slides) <= st.target_slides.max, res.warnings

    pic = tmp_path / "p.png"
    pic.write_bytes(b"x")
    for s in slides:
        if s.archetype == Archetype.IMAGE_TEXT:
            s.image = ImageSpec(prompt="p", path=str(pic))
    res = plan(DeckOutline(title="Т", purpose="other", slides=slides), st, set(Archetype))
    assert sum(1 for s in res.slides if s.archetype == Archetype.IMAGE_TEXT) == 5


def test_closing_text_never_goes_after_closing(template_path) -> None:
    """Финал с телом («Обсудим подбор решения…») на шаблоне, где у closing-образца нет тела (VK Education):
    текст — в свободный подзаголовок или опускается, но колода кончается финалом, а не продолжением."""
    pptx = template_path("VK Education")
    exemplars = load_exemplars(pptx)
    tokens = extract_tokens(pptx)
    outline = _template_only_outline()
    outline.slides[-1].bullets = ["Обсудим подбор решения под ваши задачи", "Свяжитесь с нами для начала сотрудничества"]
    for name in ("executive", "narrative", "visual"):
        res = build_deck_ir(outline, load_strategy(name), exemplars, tokens.template_id, tokens.slide_w, tokens.slide_h)
        texts = ["".join(r.text for p in e.paragraphs for r in p.runs) for e in res.ir.slides[-1].elements]
        assert texts and texts[0].startswith("Следующий шаг"), (name, texts)
        assert not any("продолжение" in w and "Следующий шаг" in w for w in res.warnings), res.warnings


def test_title_min_scale_floor() -> None:
    from deckforge.layout.fitting import MIN_SIZE_SCALE, STRUCTURAL_TITLE_MIN_SCALE, title_min_scale

    slot = lambda pt: Slot(id="t", kind=SlotKind.TITLE, box=Box(x=0, y=0, w=1, h=1), max_chars=20, max_lines=1, size_pt=pt)
    assert title_min_scale(slot(24.0), False) == MIN_SIZE_SCALE  # 24 pt не ужимаем сильнее обычного
    assert title_min_scale(slot(36.0), False) == STRUCTURAL_TITLE_MIN_SCALE  # 36 pt — до ~20 pt
    assert abs(title_min_scale(slot(30.0), False) - 0.6) < 1e-9  # 18 pt
    assert title_min_scale(slot(24.0), True) == STRUCTURAL_TITLE_MIN_SCALE
    # текст: шаг процесса по 20 pt (VK WorkSpace) — до 12 pt, мелкий текст — прежние 70 %
    from deckforge.layout.fitting import text_min_scale

    body = lambda pt: Slot(id="b", kind=SlotKind.BODY, box=Box(x=0, y=0, w=1, h=1), max_chars=18, max_lines=1, size_pt=pt)
    assert abs(text_min_scale(body(20.0)) - 0.6) < 1e-9 and text_min_scale(body(14.0)) == MIN_SIZE_SCALE


def test_single_line_body_is_one_item() -> None:
    """Однострочное тело-примечание без max_items (VK WorkSpace slide15) не вмещает три пункта."""
    from deckforge.layout.exemplar_picker import Needs, score_exemplar

    def ex(lines: int) -> Exemplar:
        return Exemplar(id=f"e{lines}", source_index=0, layout_name="", archetype=Archetype.CARDS, slots=[
            Slot(id="t", kind=SlotKind.TITLE, box=Box(x=0, y=0, w=8_000_000, h=500_000), max_chars=60, max_lines=1, size_pt=24),
            Slot(id="b", kind=SlotKind.BODY, box=Box(x=0, y=900_000, w=8_000_000, h=300_000 * lines), max_chars=66 * lines,
                 max_lines=lines, size_pt=14)])
    slide = OutlineSlide(idx=1, archetype=Archetype.CARDS, title="Т", bullets=["Комплексность", "Готовность", "Этапы"])
    n = Needs.of(slide)
    assert score_exemplar(ex(1), n, strategy(), SLIDE_W * SLIDE_H) < score_exemplar(ex(6), n, strategy(), SLIDE_W * SLIDE_H)


def test_paragraph_slides_merge_into_cards() -> None:
    """image_text с одним абзацем (ответ модели «text» → paragraphs) сливается с соседними карточками,
    абзац становится пунктом, а в слитом слайде не остаётся абзацев рядом с пунктами."""
    from deckforge.layout.planner import merge_pair

    a = OutlineSlide(idx=1, archetype=Archetype.CARDS, title="А", bullets=["один", "два", "три"])
    b = OutlineSlide(idx=2, archetype=Archetype.IMAGE_TEXT, title="Б", paragraphs=["Абзац про интеграцию."])
    merged = merge_pair([a, b], max_bullets=6)
    assert merged is not None and len(merged) == 1
    assert merged[0].bullets == ["один", "два", "три", "Абзац про интеграцию."] and not merged[0].paragraphs


def test_paragraph_goes_to_roomiest_text_slot() -> None:
    """VK WorkSpace slide15: крупные рамки карточек размечены подписями, единственное тело — однострочное
    примечание внизу. Абзац идёт в самую вместительную рамку, а не режется «…» в примечании."""
    from deckforge.layout.builder import build_slide

    slots = [_slot("t", SlotKind.TITLE, 0.05, 0.05, 0.6, 0.1, max_chars=60, size=36),
             *[_slot(f"card{i}", SlotKind.LABEL, 0.04 + 0.31 * i, 0.27, 0.3, 0.29, max_chars=200) for i in range(3)],
             _slot("note", SlotKind.BODY, 0.03, 0.81, 0.54, 0.06, max_chars=66)]
    slots[-1].max_lines = slots[-1].max_items = 1
    ex = Exemplar(id="slide15", source_index=0, layout_name="a", archetype=Archetype.CARDS, slots=slots)
    para = ("Система легко подключается к текущим корпоративным сервисам. Это позволяет избежать разрыва "
            "бизнес-процессов и сохранить привычные инструменты команд.")
    ir, left = build_slide(0, OutlineSlide(idx=1, archetype=Archetype.IMAGE_TEXT, title="Т", paragraphs=[para]),
                           ex, SLIDE_H, {})
    texts = {el.slot_id: "".join(r.text for p in el.paragraphs for r in p.runs) for el in ir.elements if el.paragraphs}
    assert left is None and texts.get("card0") == para and "note" not in texts, texts
    # короткий абзац влезает в тело — порядок чтения не меняется
    ir, _ = build_slide(0, OutlineSlide(idx=1, archetype=Archetype.IMAGE_TEXT, title="Т", paragraphs=["Коротко."]),
                        ex, SLIDE_H, {})
    assert {el.slot_id for el in ir.elements if el.paragraphs} == {"t", "note"}
