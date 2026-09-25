"""pipeline.run на фейковом LLM: outline → колоды + аудит/автофикс + manifest.json."""

import json
import time
from pathlib import Path

import pytest

from deckforge.pipeline import RunConfig, load_config, run
from tests.conftest import FakeClient, cassette

REPO = Path(__file__).resolve().parents[1]
NO_JUDGE = {"deterministic": True, "contextual": False, "autofix": True}
# ответ template_brief в режиме «на входе только шаблон»
TEMPLATE_BRIEF = {
    "topic": "VK Tech: корпоративные сервисы для команд",
    "brand": "VK Tech",
    "purpose": "product",
    "audience": "ИТ-руководители крупных компаний",
    "summary": "Показать, из чего состоит платформа и зачем она корпоративным командам.",
    "key_points": ["Команды теряют время на переключение между сервисами",
                   "Единая платформа закрывает коммуникации и документы",
                   "Разворачивается в контуре заказчика"],
    "tone": "светлый корпоративный, синий акцент",
}


def _soffice() -> bool:
    from deckforge.export.render import find_soffice

    try:
        find_soffice()
        return True
    except RuntimeError:
        return False


def test_example_config_loads() -> None:
    """В датасете только шаблоны: конфиги-примеры работают без контент-пакета."""
    cfg = load_config(REPO / "configs" / "run.example.yaml")
    assert cfg.template.is_absolute() and cfg.output_dir.is_absolute() and cfg.content_pack is None
    assert cfg.strategies == ["executive", "narrative", "visual"] and cfg.purpose == "other"
    assert cfg.audit.autofix is True and cfg.audit.contextual is True and "pdf" in cfg.export


def test_final_configs_run_on_templates_alone() -> None:
    """9 витринных колод + holdout собираются из шаблонов; контент один на все — через --outline.
    Сценарий жюри — контекст репозитория и задача на шаблоне VK Tech."""
    finals = sorted((REPO / "configs" / "final").glob("*.yaml"))
    assert [p.stem for p in finals] == ["jury_scenario", "lct2026_holdout", "vk_education", "vk_tech", "vk_workspace"]
    for path in finals:
        cfg = load_config(path)
        assert cfg.template.is_absolute() and cfg.strategies == ["executive", "narrative", "visual"]
        assert cfg.output_dir.name == path.stem and cfg.output_dir.parent.name == "output"
        if path.stem == "jury_scenario":
            assert cfg.context is not None and cfg.context.parent == cfg.output_dir and cfg.topic
            assert cfg.template.name.startswith("VK Tech") and cfg.talk_minutes == 7
        else:
            assert cfg.content_pack is None and cfg.context is None and not cfg.topic, \
                f"{path.name}: контент-пакета в датасете нет"


def test_empty_pack_falls_back_to_template_brief(template_path, tmp_path: Path) -> None:
    """Пустой пакет — не ошибка: бриф выводится из шаблона (ТЗ: на входе может быть только шаблон)."""
    (tmp_path / "pack").mkdir()
    cfg = RunConfig(template=template_path("VK Tech"), content_pack=tmp_path / "pack", strategies=["executive"],
                    output_dir=tmp_path / "run", audit=NO_JUDGE)
    client = FakeClient(by_skill={"template_brief": [TEMPLATE_BRIEF], "outline_writer": [cassette("outline_writer_pulse")]})
    res = run(cfg, client=client)

    assert [c.skill for c in client.calls] == ["template_brief@v2", "outline_writer@v3"]
    assert res.content_source == "template"
    brief = (tmp_path / "run" / "brief.md").read_text("utf-8")
    assert TEMPLATE_BRIEF["topic"] in brief and TEMPLATE_BRIEF["key_points"][0] in brief
    m = json.loads(res.decks[0].manifest.read_text("utf-8"))
    assert m["content_source"] == "template" and m["brief"] == "brief.md"
    assert m["skills"]["template_brief"] == "v2" and m["timings_s"]["brief"] >= 0
    assert any("только шаблон" in w for w in res.warnings)
    # слепок шаблона дошёл до модели цельным, а бриф — в outline_writer
    digest = client.inputs[0]["template_digest"]
    assert "Файл шаблона:" in digest and "Слайды-образцы" in digest
    assert TEMPLATE_BRIEF["topic"] in client.inputs[1]["brief"]


def test_topic_only_run(template_path, tmp_path: Path) -> None:
    """Тема одной строкой: до вёрстки LLM зовётся только за outline, бриф — сама тема."""
    cfg = RunConfig(template=template_path("VK Tech"), topic="Пульс команды: как мерить вовлечённость",
                    strategies=["executive"], output_dir=tmp_path / "run", audit=NO_JUDGE, speaker_notes=False)
    client = FakeClient(cassette("outline_writer_pulse"))
    res = run(cfg, client=client)

    assert [c.skill for c in client.calls] == ["outline_writer@v3"]
    assert res.content_source == "topic"
    assert (tmp_path / "run" / "brief.md").read_text("utf-8").startswith("Пульс команды")
    assert client.inputs[0]["brief"].startswith("Пульс команды") and client.inputs[0]["content_pack"] == ""
    m = json.loads(res.decks[0].manifest.read_text("utf-8"))
    assert m["content_source"] == "topic" and m["content_pack"] == ""


def test_ready_outline_carries_brief_to_next_template(template_path, tmp_path: Path) -> None:
    """Один контент на несколько шаблонов: с готовым outline рядом лежащий brief.md едет в прогон."""
    from deckforge.core.ir import DeckOutline

    src = tmp_path / "first"
    src.mkdir()
    outline_path = src / "outline.json"
    outline = DeckOutline.model_validate_json(
        (REPO / "examples" / "content_pack" / "outline.json").read_text("utf-8"))
    outline_path.write_text(outline.model_dump_json(indent=1), "utf-8")
    (src / "brief.md").write_text("# Тема из шаблона\n\nТезис один.\n", "utf-8")

    cfg = RunConfig(template=template_path("VK Tech"), strategies=["executive"], output_dir=tmp_path / "run",
                    audit=NO_JUDGE)
    res = run(cfg, outline=outline, outline_path=outline_path)

    assert res.content_source == "outline"
    assert (tmp_path / "run" / "brief.md").read_text("utf-8").startswith("# Тема из шаблона")
    m = json.loads(res.decks[0].manifest.read_text("utf-8"))
    assert m["content_source"] == "outline" and m["brief"] == "brief.md" and m["content_pack"] == ""


def test_run_with_fake_llm(template_path, tmp_path: Path) -> None:
    cfg = RunConfig(
        template=template_path("VK Tech"), content_pack=REPO / "examples" / "content_pack", purpose="product",
        audience="руководители", target_slides=12, strategies=["executive", "visual"], output_dir=tmp_path / "run",
        audit=NO_JUDGE,
    )
    from tests.conftest import fake_notes

    client = FakeClient(cassette("outline_writer_pulse"), by_skill={"speaker_notes": [fake_notes]})
    messages: list[str] = []
    res = run(cfg, client=client, progress=messages.append)

    # outline — один на прогон, текст выступления — по вызову на колоду
    assert [c.skill for c in client.calls].count("outline_writer@v3") == 1 and len(client.calls) == 3
    assert (tmp_path / "run" / "outline.json").exists() and (tmp_path / "run" / "outline.raw.json").exists()
    assert [d.strategy for d in res.decks] == ["executive", "visual"]
    for d in res.decks:
        assert d.pptx.exists() and d.pptx.stat().st_size > 100_000 and d.ir_json.exists()
        m = json.loads(d.manifest.read_text("utf-8"))
        assert m["skills"] == {"outline_writer": "v3", "speaker_notes": "v3"}
        assert m["models"]["text"] == "fake-text"
        assert [c["skill"] for c in m["llm_calls"]] == ["outline_writer@v3", "speaker_notes@v3"]
        assert {"parse", "outline", "layout", "render", "audit", "autofix", "notes"} <= set(m["timings_s"])
        assert d.audit is not None and d.audit.exists() and m["audit"]["checks_run"] == 26
        assert m["audit"]["errors"] == d.audit_summary["errors"] and "by_check" in m["audit"]
        assert m["audit"]["kind"] == "deterministic" and m["audit"]["contextual"] == 0
        fix = m["audit"]["autofix"]
        assert {"applied", "skipped", "before", "after", "items"} <= set(fix)
        assert fix["after"]["errors"] <= fix["before"]["errors"]
        if fix["applied"]:
            ir = json.loads(d.ir_json.read_text("utf-8"))
            assert any("size_pt" in el["style_overrides"] or any(r["size_pt"] for p in el["paragraphs"] for r in p["runs"])
                       for s in ir["slides"] for el in s["elements"])
        assert m["strategy"]["name"] == d.strategy and m["template"]["sha1"]
        assert len(m["plan"]) == len(m["choices"]) >= 8 and d.stats["skipped"] == 0
    run_json = json.loads(res.run_json.read_text("utf-8"))
    assert run_json["timings_s"]["total"] > 0
    assert run_json["not_implemented"] == []
    assert all(json.loads(d.manifest.read_text("utf-8"))["images"].get("generated", 0) == 0 for d in res.decks)
    assert (tmp_path / "run" / "dna.json").exists() and run_json["decks"][0]["audit"]["checks_run"] == 26
    assert (tmp_path / "run" / "compare.md").read_text("utf-8").count("\n") >= 3
    assert any(m.startswith("outline:") for m in messages)


def test_autofix_removes_our_overflow_errors(template_path, tmp_path: Path, no_fitting) -> None:
    """Без подгонки текста есть L03; с фиксами их нет."""
    from deckforge.core.ir import DeckOutline

    outline = DeckOutline.model_validate_json((REPO / "examples" / "content_pack" / "outline.json").read_text("utf-8"))
    base = dict(template=template_path("VK Tech"), content_pack=REPO / "examples" / "content_pack",
                strategies=["narrative"], images="off")
    raw = run(RunConfig(**base, output_dir=tmp_path / "raw", audit={**NO_JUDGE, "autofix": False}), outline=outline)
    fixed = run(RunConfig(**base, output_dir=tmp_path / "fixed", audit=NO_JUDGE), outline=outline)
    raw_report = json.loads(raw.decks[0].audit.read_text("utf-8"))
    fixed_report = json.loads(fixed.decks[0].audit.read_text("utf-8"))
    raw_l03 = [f for f in raw_report["findings"] if f["check_id"] == "L03_text_overflow" and f["severity"] == "error"]
    fixed_l03 = [f for f in fixed_report["findings"] if f["check_id"] == "L03_text_overflow" and f["severity"] == "error"]
    assert raw_l03 and not fixed_l03
    m = json.loads(fixed.decks[0].manifest.read_text("utf-8"))
    assert m["audit"]["autofix"]["applied"] >= len(raw_l03)
    assert m["audit"]["autofix"]["after"]["errors"] < m["audit"]["autofix"]["before"]["errors"]
    assert all({"slide_idx", "check_id", "fix", "before", "after"} <= set(it) for it in m["audit"]["autofix"]["items"]
               if "reason" not in it)
    assert "autofix" not in json.loads(raw.decks[0].manifest.read_text("utf-8"))["timings_s"]


def test_run_with_ready_outline_skips_llm(template_path, tmp_path: Path) -> None:
    from deckforge.core.ir import DeckOutline

    outline = DeckOutline.model_validate_json((REPO / "examples" / "content_pack" / "outline.json").read_text("utf-8"))
    cfg = RunConfig(template=template_path("VK Tech"), content_pack=REPO / "examples" / "content_pack",
                    strategies=["narrative"], output_dir=tmp_path, images="off",
                    audit={"deterministic": False, "contextual": False, "autofix": False})
    res = run(cfg, outline=outline)
    m = json.loads(res.decks[0].manifest.read_text("utf-8"))
    assert m["skills"] == {} and m["llm_calls"] == [] and m["timings_s"]["outline"] == 0
    assert m["audit"] == {} and res.decks[0].audit is None and "audit" not in m["timings_s"]
    assert json.loads(res.run_json.read_text("utf-8"))["not_implemented"] == []


class ImageFakeClient(FakeClient):
    """FakeClient с text-to-image: пишет 1×1 PNG и считает вызовы."""

    images_enabled = True
    image_model = "fake-t2i"

    def generate_image(self, prompt: str, out_path: Path, size: str = "1024x576", deadline: float | None = None) -> Path:
        from PIL import Image

        from deckforge.llm.client import LLMCall

        Image.new("RGB", (64, 36), (0, 119, 255)).save(out_path)
        self._log(LLMCall("image_gen", self.image_model, 0.01))
        return out_path


def test_images_generated_cached_and_rendered(template_path, tmp_path: Path) -> None:
    """Иллюстрации генерируются до вёрстки и попадают в .pptx; повтор — из кэша; images: off — ни одного вызова."""
    from deckforge.content.images import MAX_IMAGES

    cfg = RunConfig(template=template_path("VK Tech"), content_pack=REPO / "examples" / "content_pack",
                    strategies=["visual"], output_dir=tmp_path / "img", audit=NO_JUDGE)
    client = ImageFakeClient(by_skill={"image_prompter": [{"prompt": "abstract blue gradient, no text"}]})
    deck = run(cfg, client=client, outline=_outline()).decks[0]
    m = json.loads(deck.manifest.read_text("utf-8"))
    gen = [c for c in client.calls if c.skill == "image_gen"]
    assert 1 <= len(gen) <= MAX_IMAGES and m["images"]["generated"] == len(gen) and m["images"]["mode"] == "always"
    assert m["skills"]["image_prompter"] == "v1" and m["timings_s"]["images"] >= 0
    ir = json.loads(deck.ir_json.read_text("utf-8"))
    with_pic = [el for s in ir["slides"] for el in s["elements"] if el["image_path"]]
    assert with_pic, "сгенерированная картинка должна встать в picture-слот"
    from pptx import Presentation

    prs = Presentation(str(deck.pptx))
    assert sum(1 for sl in prs.slides for sh in sl.shapes if sh.shape_type is not None and "PICTURE" in str(sh.shape_type)
               or sh._element.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip") is not None) >= 1
    # повтор — из кэша
    again = run(cfg, client=client, outline=_outline()).decks[0]
    m2 = json.loads(again.manifest.read_text("utf-8"))
    assert m2["images"]["cache"] == len(gen) and m2["images"]["generated"] == 0
    assert len([c for c in client.calls if c.skill == "image_gen"]) == len(gen)
    off = run(cfg.model_copy(update={"images": "off", "output_dir": tmp_path / "off"}), client=client, outline=_outline()).decks[0]
    assert json.loads(off.manifest.read_text("utf-8"))["images"] == {}


def test_image_generated_once_for_parallel_decks(tmp_path: Path) -> None:
    """narrative и visual иллюстрируют одни слайды одновременно: картинку генерирует одна колода, вторая ждёт
    её по ключу и берёт из кэша — без второй оплаты и без записи в тот же файл."""
    import threading

    from deckforge.content.images import _illustrate_one
    from deckforge.core.ir import Archetype, OutlineSlide
    from deckforge.llm.skills import load_skill

    class SlowImages(ImageFakeClient):
        def generate_image(self, prompt: str, out_path: Path, size: str = "1024x576", deadline: float | None = None):
            time.sleep(0.2)
            return super().generate_image(prompt, out_path, size, deadline)

    client = SlowImages(by_skill={"image_prompter": [{"prompt": "abstract blue gradient"}]})
    slide = OutlineSlide(idx=3, archetype=Archetype.IMAGE_TEXT, title="Платформа", paragraphs=["Один абзац текста"])
    skill = load_skill("image_prompter")
    results = []

    def work(c) -> None:
        results.append(_illustrate_one(slide, c, skill, {"palette": "0077FF"}, "light background", tmp_path))

    threads = [threading.Thread(target=work, args=(client.fork(),)) for _ in range(2)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert sorted(r.source for r in results) == ["cache", "generated"]
    assert results[0].path == results[1].path and Path(results[0].path).exists()
    assert len([c for c in client.calls if c.skill == "image_gen"]) == 1
    assert not list(tmp_path.glob("*.tmp"))


def test_illustrate_does_not_wait_past_deadline(tmp_path: Path) -> None:
    """Начатую генерацию, не ответившую к дедлайну, колода не ждёт — как и судью: бюджет прогона важнее картинки."""
    import threading

    from deckforge.content.images import illustrate
    from deckforge.core.strategy import load_strategy

    release = threading.Event()

    class Hanging(ImageFakeClient):
        def generate_image(self, prompt: str, out_path: Path, size: str = "1024x576", deadline: float | None = None):
            release.wait(10)  # ответа к дедлайну не будет
            return super().generate_image(prompt, out_path, size, deadline)

    client = Hanging(by_skill={"image_prompter": [{"prompt": "abstract blue gradient"}]})
    try:
        t0 = time.monotonic()
        res = illustrate(_outline(), load_strategy("visual"), {"palette": "0077FF"}, client, tmp_path,
                         cfg_mode="always", deadline=t0 + 0.5)
        assert time.monotonic() - t0 < 5  # вернулись сразу после дедлайна, а не через 10 с
    finally:
        release.set()  # брошенные потоки держат локи ключей кэша — отпустить до следующих тестов
    assert res.items and all(i.source == "failed" and "бюджет" in (i.error or "") for i in res.items)
    assert len(res.warnings) == len(res.items)
    assert not any(s.image and s.image.path for s in res.outline.slides)


def test_speaker_notes_for_every_slide_of_every_deck(template_path, tmp_path: Path) -> None:
    """Вводная жюри: у каждого слайда каждого варианта есть текст выступления под заданную длительность —
    в заметках .pptx, в speech.md, в manifest; N01/N02 чистые. Разделители narrative тоже с текстом."""
    from pptx import Presentation

    from tests.conftest import fake_notes

    cfg = RunConfig(template=template_path("VK Tech"), content_pack=REPO / "examples" / "content_pack",
                    strategies=["executive", "narrative"], output_dir=tmp_path, images="off", talk_minutes=5,
                    audit=NO_JUDGE)
    client = FakeClient(by_skill={"speaker_notes": [fake_notes]})
    res = run(cfg, client=client, outline=_outline())
    assert len(client.calls) == 2 and all(c.skill == "speaker_notes@v3" for c in client.calls)
    for d in res.decks:
        m = d.load_manifest()
        prs = Presentation(str(d.pptx))
        assert all(s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip() for s in prs.slides)
        assert m["skills"]["speaker_notes"] == "v3" and "notes" in m["timings_s"]
        assert m["speech"]["talk_minutes"] == 5 and 3.75 <= m["speech"]["minutes"] <= 6.25
        assert m["speech"]["slides_without_notes"] == [] and m["exports"]["speech"] == f"{d.strategy}.speech.md"
        assert d.speech is not None and d.speech.read_text("utf-8").count("\n## ") == m["stats"]["slides"]
        report = d.load_report()
        assert report is not None and {"N01_notes_missing", "N02_talk_duration"} <= set(report.checks_run)
        assert not [f for f in report.findings if f.check_id.startswith("N0")]
        ir = d.load_ir()
        assert ir.talk_minutes == 5 and all(s.notes for s in ir.slides)
    assert any(s.archetype.value == "section" for s in res.decks[1].load_ir().slides)
    assert "| речь, мин |" in (tmp_path / "compare.md").read_text("utf-8")


def test_speaker_notes_do_not_outlive_deadline(template_path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Модель не ответила к дедлайну прогона — колода готова без текста, с предупреждением и N01, а не ждёт ответа."""
    import importlib
    import threading

    run_mod = importlib.import_module("deckforge.pipeline.run")
    monkeypatch.setattr(run_mod, "EXPORT_RESERVE_S", 0.0)
    monkeypatch.setattr(run_mod, "NOTES_MIN_LEFT_S", 0.0)
    release = threading.Event()

    def hang(inputs: dict) -> dict:
        release.wait(30)  # ответа к дедлайну не будет
        return {"notes": []}

    cfg = RunConfig(template=template_path("VK Tech"), content_pack=REPO / "examples" / "content_pack",
                    strategies=["executive"], output_dir=tmp_path, images="off", time_budget_s=10, audit=NO_JUDGE)
    try:
        res = run(cfg, client=FakeClient(by_skill={"speaker_notes": [hang]}), outline=_outline())
    finally:
        release.set()
    d = res.decks[0]
    m = d.load_manifest()
    assert m["timings_s"]["deck_total"] < 15  # дедлайн 10 с + экспорт, а не 30 с ожидания
    assert any("не готов к дедлайну" in w for w in m["warnings"])
    # остались только черновики заметок из outline «Пульса» — у остальных слайдов N01
    missing = m["speech"]["slides_without_notes"]
    assert 0 < len(missing) < m["stats"]["slides"] and "notes" not in m["timings_s"]
    report = d.load_report()
    assert report is not None
    assert sorted(f.slide_idx + 1 for f in report.findings if f.check_id == "N01_notes_missing") == missing


def test_ready_outline_without_judge_and_images_still_gets_notes(template_path, tmp_path: Path,
                                                                 monkeypatch: pytest.MonkeyPatch) -> None:
    """`run --outline … --no-judge --no-images`: судье и картинкам клиент не нужен, но текст выступления пишет
    модель — клиент создаётся по ключу. Без ключа колоды собираются без текста, с предупреждением."""
    import importlib

    from tests.conftest import fake_notes

    run_mod = importlib.import_module("deckforge.pipeline.run")
    made: list[FakeClient] = []
    monkeypatch.setattr(run_mod, "LLMClient",
                        lambda: made.append(FakeClient(by_skill={"speaker_notes": [fake_notes]})) or made[-1])
    cfg = RunConfig(template=template_path("VK Tech"), strategies=["executive"], output_dir=tmp_path / "key",
                    images="off", audit=NO_JUDGE)

    monkeypatch.setenv("LLM_API_KEY", "sk-test")
    m = run(cfg, outline=_outline()).decks[0].load_manifest()
    assert len(made) == 1 and m["speech"]["slides_without_notes"] == [] and "notes" in m["timings_s"]

    monkeypatch.setenv("LLM_API_KEY", "")
    res = run(cfg.model_copy(update={"output_dir": tmp_path / "nokey"}), outline=_outline())
    assert len(made) == 1  # без ключа клиента нет — и запросов в API тоже
    assert any("LLM_API_KEY не задан" in w for w in res.warnings)
    report = res.decks[0].load_report()
    assert report is not None and not [f for f in report.findings if f.check_id.startswith("N0")]


def test_keep_awake_sets_and_clears_execution_state(monkeypatch: pytest.MonkeyPatch) -> None:
    """Прогон держит систему бодрствующей (Modern Standby посреди прогона съел бюджет) и снимает флаг даже при ошибке."""
    import importlib

    run_mod = importlib.import_module("deckforge.pipeline.run")  # в пакете имя run — функция
    calls: list[int] = []
    monkeypatch.setattr(run_mod, "_set_execution_state", calls.append)
    with pytest.raises(RuntimeError), run_mod._keep_awake():
        assert calls[-1] & run_mod.ES_DISPLAY_REQUIRED and calls[-1] & run_mod.ES_SYSTEM_REQUIRED
        raise RuntimeError("колода упала")
    assert calls == [run_mod.ES_CONTINUOUS | run_mod.ES_SYSTEM_REQUIRED | run_mod.ES_DISPLAY_REQUIRED,
                     run_mod.ES_CONTINUOUS]


def test_set_execution_state_is_noop_off_windows(monkeypatch: pytest.MonkeyPatch) -> None:
    import ctypes
    import importlib

    run_mod = importlib.import_module("deckforge.pipeline.run")
    called: list[int] = []

    class Kernel32:
        def SetThreadExecutionState(self, flags: int) -> int:  # noqa: N802 — имя WinAPI
            called.append(flags)
            return 1

    monkeypatch.setattr(ctypes, "windll", type("WinDLL", (), {"kernel32": Kernel32()})(), raising=False)
    monkeypatch.setattr(run_mod.sys, "platform", "linux")
    run_mod._set_execution_state(run_mod.ES_CONTINUOUS)
    assert called == []
    monkeypatch.setattr(run_mod.sys, "platform", "win32")
    run_mod._set_execution_state(run_mod.ES_CONTINUOUS)
    assert called == [run_mod.ES_CONTINUOUS]


def _outline():
    from deckforge.core.ir import DeckOutline

    return DeckOutline.model_validate_json((REPO / "examples" / "content_pack" / "outline.json").read_text("utf-8"))


def test_build_deck_matches_run(template_path, tmp_path: Path) -> None:
    """Этапы по отдельности (parse_template → make_outline → build_deck) дают то же, что run()."""
    from deckforge.pipeline import RunContext, build_deck, make_outline, parse_template

    cfg = RunConfig(template=template_path("VK Tech"), content_pack=REPO / "examples" / "content_pack",
                    strategies=["executive"], output_dir=tmp_path / "steps", images="off", audit=NO_JUDGE)
    parsed = parse_template(cfg.template, cfg.output_dir)
    assert parsed.meta["sha1"] and parsed.dna.exemplars and (cfg.output_dir / "dna.json").exists()
    step = make_outline(cfg, parsed, cfg.output_dir, outline=_outline())
    assert step.path.exists() and step.seconds == 0 and step.skills_used == {}
    ctx = RunContext.prepare(cfg, parsed, step)
    assert ctx.contextual_on is False and ctx.client is None
    deck = build_deck(ctx, "executive")
    whole = run(cfg.model_copy(update={"output_dir": tmp_path / "whole"}), outline=_outline()).decks[0]
    assert deck.pptx.exists() and deck.audit_summary["errors"] == whole.audit_summary["errors"]
    assert set(deck.load_manifest()) == set(whole.load_manifest())
    m = deck.load_manifest()
    # без LLM текст выступления — черновики заметок outline «Пульса»
    assert m["exports"] == {"pptx": "executive.pptx", "speech": "executive.speech.md"} and deck.pdf is None
    # пути в manifest — относительные
    assert m["outline"] == "outline.json" and m["audit"]["path"] == "executive.audit.json"
    assert json.loads(deck.audit.read_text("utf-8"))["deck_path"] == "executive.pptx"
    assert ":" not in json.loads((cfg.output_dir / "dna.json").read_text("utf-8"))["source_path"]
    assert deck.load_report().deck_path == "executive.pptx" and parsed.dna.source_path == str(parsed.template)
    if Path.cwd().resolve() == REPO.resolve():
        assert m["template"]["path"] == "data/templates/VK Tech шаблон.pptx" and m["content_pack"] == "examples/content_pack"
    assert deck.load_report() is not None and len(deck.load_ir().slides) == deck.stats["slides"]


def test_refine_deck_applies_user_fixes(template_path, tmp_path: Path, no_fitting) -> None:
    """Фиксы по выбору пользователя: выбранные L03 после refine исчезают."""
    from deckforge.pipeline import refine_deck

    cfg = RunConfig(template=template_path("VK Tech"), content_pack=REPO / "examples" / "content_pack",
                    strategies=["narrative"], output_dir=tmp_path, images="off",
                    audit={**NO_JUDGE, "autofix": False})
    res = run(cfg, outline=_outline())
    deck = res.decks[0]
    report = deck.load_report()
    selected = [i for i, f in enumerate(report.findings) if f.check_id == "L03_text_overflow" and f.severity == "error"]
    assert selected
    before_errors = deck.audit_summary["errors"]
    fixed = refine_deck(deck, res.parsed, selected)
    new_report = fixed.load_report()
    assert not [f for f in new_report.findings if f.check_id == "L03_text_overflow" and f.severity == "error"]
    assert fixed.audit_summary["errors"] < before_errors
    m = fixed.load_manifest()
    fix = m["audit"]["autofix"]
    assert fix["user_applied"] >= 1 and fix["applied"] == fix["user_applied"] and fix["after"]["errors"] < before_errors
    assert "refine" in m["timings_s"] and "contextual_stale" not in m["audit"] and m["audit"]["contextual"] == 0
    # выбор без фиксируемых находок ничего не меняет
    same = refine_deck(fixed, res.parsed, [])
    assert same.audit_summary["errors"] == fixed.audit_summary["errors"]
    assert same.load_manifest()["audit"]["autofix"]["applied"] == fix["applied"]


@pytest.mark.skipif(not _soffice(), reason="нужен LibreOffice для PDF")
def test_pdf_export(template_path, tmp_path: Path) -> None:
    """PDF без PNG конвертируется во временную папку колоды; колоды экспортируются параллельно и не мешают друг другу."""
    cfg = RunConfig(template=template_path("VK Tech"), content_pack=REPO / "examples" / "content_pack",
                    strategies=["executive", "narrative"], output_dir=tmp_path, images="off",
                    export=["pptx", "pdf", "html"], max_parallel_decks=2,
                    audit={"deterministic": False, "contextual": False, "autofix": False})
    res = run(cfg, outline=_outline())
    d = res.decks[0]
    assert d.pdf is not None, d.warnings
    assert d.pdf.exists() and d.pdf.name == "executive.pdf" and d.pdf.stat().st_size > 10_000
    assert res.decks[1].pdf is not None, res.decks[1].warnings
    assert res.decks[1].pdf.name == "narrative.pdf" and res.decks[1].pdf.exists()
    assert not list(tmp_path.glob("_pdf*"))
    assert d.html is not None and d.html.exists() and d.html.name == "executive.html"
    m = d.load_manifest()
    assert m["exports"] == {"pptx": "executive.pptx", "pdf": "executive.pdf", "html": "executive.html",
                            "speech": "executive.speech.md"}
    assert "export_pdf" in m["timings_s"] and "export_html" in m["timings_s"]
    run_json = json.loads(res.run_json.read_text("utf-8"))
    assert run_json["not_implemented"] == [] and run_json["decks"][0]["exports"]["html"] == "executive.html"
    assert run_json["decks"][0]["manifest"] == "executive.manifest.json" and run_json["outline"] == "outline.json"


@pytest.mark.skipif(not _soffice(), reason="нужен LibreOffice для PNG")
def test_run_with_contextual_judge(template_path, tmp_path: Path) -> None:
    """Судья вызывается по слайду, находки C.. попадают в audit.json, версия скилла — в manifest."""
    from deckforge.core.ir import DeckOutline

    outline = DeckOutline.model_validate_json((REPO / "examples" / "content_pack" / "outline.json").read_text("utf-8"))
    cfg = RunConfig(template=template_path("VK Tech"), content_pack=REPO / "examples" / "content_pack",
                    strategies=["executive"], output_dir=tmp_path, images="off", render_dpi=40, speaker_notes=False,
                    audit={"deterministic": True, "contextual": True, "autofix": False})
    client = FakeClient(by_skill={"audit_judge": cassette("audit_judge_pulse")})
    res = run(cfg, client=client, outline=outline)
    d = res.decks[0]
    m = json.loads(d.manifest.read_text("utf-8"))
    n = m["stats"]["slides"]
    assert len(client.calls) == n and all(c.skill == "audit_judge@v2" for c in client.calls)
    assert all(len(imgs) == 1 and imgs[0].suffix == ".png" for imgs in client.images)
    assert client.inputs[0]["prev_slide_title"].startswith("(нет") and client.inputs[0]["language"] == "ru"
    assert any("[brief" in i["source_facts"] or "metrics:" in i["source_facts"] for i in client.inputs)
    report = json.loads(d.audit.read_text("utf-8"))
    assert len(report["checks_run"]) == 24 + 11
    ctx = [f for f in report["findings"] if f["kind"] == "contextual"]
    assert ctx and m["audit"]["contextual"] == len(ctx) and m["audit"]["kind"] == "deterministic+contextual"
    assert m["skills"]["audit_judge"] == "v2" and "audit_contextual" in m["timings_s"] and "png" in m["timings_s"]
    assert len(d.pngs) == n and not (tmp_path / "executive" / "contact.png").exists()
    # бюджет времени: судья получил дедлайн, время колоды и бюджет — в manifest и compare.md
    assert all(dl is not None for dl in client.deadlines) and m["time_budget_s"] == 300
    assert 0 < m["timings_s"]["deck_total"] < 300 and not any("бюджет" in w for w in m["warnings"])
    compare = (tmp_path / "compare.md").read_text("utf-8")
    assert "| время, с |" in compare and "**executive** — руководителю" in compare
    assert m["strategy"]["audience_hint"].startswith("руководителю") and d.audience_hint == m["strategy"]["audience_hint"]


def test_time_budget_skips_judge_and_images(template_path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """ТЗ: три варианта ≤ 5 мин независимо от инференса — при исчерпанном бюджете прогона судья и картинки
    пропускаются, вёрстка, аудит и экспорт делаются всегда; бюджет читается из RUN_TIME_BUDGET_S."""
    from deckforge.core.ir import DeckOutline

    monkeypatch.setenv("RUN_TIME_BUDGET_S", "10")
    monkeypatch.setenv("DECK_MAX_PARALLEL_LLM", "2")
    outline = DeckOutline.model_validate_json((REPO / "examples" / "content_pack" / "outline.json").read_text("utf-8"))
    cfg = RunConfig(template=template_path("VK Tech"), content_pack=REPO / "examples" / "content_pack",
                    strategies=["visual"], output_dir=tmp_path, images="auto", render_dpi=40,
                    audit={"deterministic": True, "contextual": True, "autofix": True})
    assert cfg.time_budget_s == 10 and cfg.max_parallel_llm == 2
    client = FakeClient(by_skill={"audit_judge": cassette("audit_judge_pulse")})
    client.images_enabled = True
    res = run(cfg, client=client, outline=outline)
    d = res.decks[0]
    m = json.loads(d.manifest.read_text("utf-8"))
    assert client.calls == []  # ни судьи, ни промптов картинок: до дедлайна меньше резерва
    assert d.pptx.exists() and d.audit is not None and m["audit"]["kind"] == "deterministic"
    assert any("images: пропущено" in w for w in m["warnings"])
    assert any("VLM-судья пропущен" in w for w in m["warnings"])
    assert m["time_budget_s"] == 10 and 0 < m["timings_s"]["deck_build"] <= m["timings_s"]["deck_total"]
    # превышение бюджета не замалчивается: предупреждение есть ровно тогда, когда колода вышла за него
    assert (m["timings_s"]["deck_total"] > 10) == any("бюджет времени прогона превышен" in w for w in m["warnings"])
    assert m["strategy"]["target_slides"] == {"min": 10, "max": 12}
    run_json = json.loads(res.run_json.read_text("utf-8"))
    assert (run_json["timings_s"]["total"] > 10) == any("бюджет времени прогона превышен: " in w
                                                        for w in run_json["warnings"] if not w.startswith("visual"))


def test_run_config_budget_and_strategies(monkeypatch: pytest.MonkeyPatch) -> None:
    """Бюджет — на прогон (прежнее имя DECK_TIME_BUDGET_S читается как запасное); повторы стратегий убираются."""
    monkeypatch.delenv("RUN_TIME_BUDGET_S", raising=False)
    monkeypatch.delenv("RUN_MAX_PARALLEL_DECKS", raising=False)
    monkeypatch.setenv("DECK_TIME_BUDGET_S", "42")
    cfg = RunConfig(template=Path("t.pptx"), strategies=["visual", " executive", "visual"])
    assert cfg.time_budget_s == 42 and cfg.max_parallel_decks == 3
    assert cfg.strategies == ["visual", "executive"]
    monkeypatch.setenv("RUN_TIME_BUDGET_S", "240")
    monkeypatch.setenv("RUN_MAX_PARALLEL_DECKS", "1")
    cfg = RunConfig(template=Path("t.pptx"))
    assert cfg.time_budget_s == 240 and cfg.max_parallel_decks == 1


def test_decks_build_in_parallel(monkeypatch: pytest.MonkeyPatch) -> None:
    """Колоды прогона — в потоках одновременно; результат в порядке конфига, прогресс — только из вызывающего
    потока (Streamlit пишет в статус только из потока скрипта), ошибка колоды — после того, как досчитались остальные."""
    import importlib
    import threading

    run_mod = importlib.import_module("deckforge.pipeline.run")  # в пакете имя run — функция
    names = ["executive", "narrative", "visual"]
    active, peak, finished, lock = [0], [0], [], threading.Lock()
    barrier = threading.Barrier(len(names), timeout=10)

    def fake_build(ctx, name, progress=None):
        with lock:
            active[0] += 1
            peak[0] = max(peak[0], active[0])
        if ctx != "seq":
            barrier.wait()  # не дойдут все три одновременно — BrokenBarrierError
        time.sleep({"executive": 0.3, "narrative": 0.05, "visual": 0.15}[name])
        progress(f"{name}: готово")
        with lock:
            active[0] -= 1
            finished.append(name)
        if ctx == "boom" and name == "narrative":
            raise RuntimeError("narrative упала")
        return name

    monkeypatch.setattr(run_mod, "build_deck", fake_build)
    main = threading.current_thread()
    said: list[tuple[str, bool]] = []

    def say(msg: str) -> None:
        said.append((msg, threading.current_thread() is main))

    assert run_mod._build_decks("ok", names, say, workers=3) == names
    assert peak[0] == 3 and all(on_main for _, on_main in said)
    assert sorted(m for m, _ in said) == sorted(f"{n}: готово" for n in names)

    finished.clear()
    barrier.reset()
    with pytest.raises(RuntimeError, match="narrative упала"):
        run_mod._build_decks("boom", names, say, workers=3)
    assert sorted(finished) == sorted(names)

    peak[0] = 0
    assert run_mod._build_decks("seq", names, say, workers=1) == names and peak[0] == 1


@pytest.mark.skipif(not _soffice(), reason="нужен LibreOffice для PNG")
def test_parallel_decks_share_deadline_and_keep_own_calls(template_path, tmp_path: Path,
                                                         monkeypatch: pytest.MonkeyPatch) -> None:
    """Три колоды с судьёй параллельно: дедлайн один на прогон, в manifest колоды — только её вызовы LLM."""
    from deckforge.audit.contextual import judge as judge_mod
    from deckforge.pipeline.run import EXPORT_RESERVE_S

    # окно слайда (у каждого своё, test_audit_contextual) шире бюджета — в запрос уходит дедлайн прогона
    monkeypatch.setattr(judge_mod, "SLIDE_TIMEOUT_S", 10_000.0)

    cfg = RunConfig(template=template_path("VK Tech"), content_pack=REPO / "examples" / "content_pack",
                    strategies=["executive", "narrative", "visual"], output_dir=tmp_path, images="off", render_dpi=40,
                    time_budget_s=300, max_parallel_decks=3,
                    audit={"deterministic": True, "contextual": True, "autofix": False})
    client = FakeClient(by_skill={"audit_judge": cassette("audit_judge_pulse")})
    messages: list[str] = []
    t0 = time.monotonic()
    res = run(cfg, client=client, outline=_outline(), progress=messages.append)
    assert [d.strategy for d in res.decks] == cfg.strategies
    # дедлайн судьи у всех колод один: старт прогона + бюджет − резерв на экспорт
    assert len(set(client.deadlines)) == 1
    assert abs(client.deadlines[0] - (t0 + 300 - EXPORT_RESERVE_S)) < 5
    n_judge = 0
    for d in res.decks:
        m = d.load_manifest()
        judge = [c for c in m["llm_calls"] if c["skill"].startswith("audit_judge")]
        assert len(judge) == m["stats"]["slides"] and m["audit"]["kind"] == "deterministic+contextual"
        assert m["time_budget_s"] == 300 and 0 < m["timings_s"]["deck_build"] <= m["timings_s"]["deck_total"] < 300
        n_judge += len(judge)
    assert len(client.calls) == n_judge
    run_json = json.loads(res.run_json.read_text("utf-8"))
    assert run_json["config"]["max_parallel_decks"] == 3 and "decks" in run_json["timings_s"]
    assert max(d["timings_s"]["deck_total"] for d in run_json["decks"]) <= run_json["timings_s"]["total"]
    assert any(m.startswith("сборка:") and "параллельно (3 потока)" in m for m in messages)
    assert "параллельно (3 потока)" in (tmp_path / "compare.md").read_text("utf-8")


def test_target_slides_shifts_strategy_ranges(template_path, tmp_path: Path) -> None:
    """Объём, заданный пользователем, доходит до стратегий: 15 → executive 13–14, narrative 15–18."""
    from deckforge.core.ir import DeckOutline

    outline = DeckOutline.model_validate_json((REPO / "examples" / "content_pack" / "outline.json").read_text("utf-8"))
    cfg = RunConfig(template=template_path("VK Tech"), content_pack=REPO / "examples" / "content_pack",
                    strategies=["executive", "narrative"], output_dir=tmp_path, images="off", target_slides=15,
                    audit=NO_JUDGE)
    res = run(cfg, outline=outline)
    by = {d.strategy: json.loads(d.manifest.read_text("utf-8")) for d in res.decks}
    assert by["executive"]["strategy"]["target_slides"] == {"min": 13, "max": 14}
    assert by["narrative"]["strategy"]["target_slides"] == {"min": 15, "max": 18}
    assert 13 <= by["executive"]["stats"]["slides"] <= 14
    assert 15 <= by["narrative"]["stats"]["slides"] <= 18
