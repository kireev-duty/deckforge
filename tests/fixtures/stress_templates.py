"""Синтетические шаблоны-крайности для стресс-теста: пустая презентация, нестандартные размеры, сотня фигур,
вложенные группы, нативные chart/table и т. п. Реестр `TEMPLATES` читает `tests/test_stress.py`.
"""

from __future__ import annotations

from collections.abc import Callable
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

SW, SH = Inches(13.333), Inches(7.5)


def _prs(w=SW, h=SH):
    prs = Presentation()
    prs.slide_width, prs.slide_height = w, h
    return prs


def _text(slide, x, y, w, h, text: str, size: int = 18) -> None:
    tb = slide.shapes.add_textbox(x, y, w, h)
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)


def _png(tmp: Path, w: int = 400, h: int = 300) -> Path:
    p = tmp / f"img_{w}x{h}.png"
    if not p.exists():
        Image.new("RGB", (w, h), "#0077FF").save(p)
    return p


def empty_presentation(tmp: Path) -> Path:
    prs = _prs()
    prs.save(p := tmp / "empty.pptx")
    return p


def one_blank_slide(tmp: Path) -> Path:
    prs = _prs()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(p := tmp / "one_blank.pptx")
    return p


def four_by_three(tmp: Path) -> Path:
    prs = _prs(Inches(10), Inches(7.5))
    for i in range(3):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        _text(s, Inches(0.5), Inches(0.4), Inches(9), Inches(1), f"Заголовок {i + 1}", 32)
        _text(s, Inches(0.5), Inches(1.8), Inches(9), Inches(4.5), "Первый тезис\nВторой тезис\nТретий тезис")
    prs.save(p := tmp / "four_by_three.pptx")
    return p


def a4_portrait(tmp: Path) -> Path:
    prs = _prs(Inches(8.27), Inches(11.69))
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _text(s, Inches(0.5), Inches(0.5), Inches(7), Inches(1.2), "Портрет A4", 36)
    _text(s, Inches(0.5), Inches(2), Inches(7), Inches(8), "Пункт один\nПункт два\nПункт три\nПункт четыре")
    prs.save(p := tmp / "a4_portrait.pptx")
    return p


def pictures_only(tmp: Path) -> Path:
    prs = _prs()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    for i in range(4):
        s.shapes.add_picture(str(_png(tmp)), Inches(0.5 + 3.2 * i), Inches(2), Inches(3), Inches(2.25))
    prs.save(p := tmp / "pictures_only.pptx")
    return p


def hundred_shapes(tmp: Path) -> Path:
    prs = _prs()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _text(s, Inches(0.5), Inches(0.3), Inches(12), Inches(1), "Сто фигур", 32)
    for i in range(100):
        col, row = i % 10, i // 10
        _text(s, Inches(0.5 + 1.25 * col), Inches(1.5 + 0.55 * row), Inches(1.1), Inches(0.45), f"ф{i}", 9)
    prs.save(p := tmp / "hundred_shapes.pptx")
    return p


def nested_groups(tmp: Path) -> Path:
    prs = _prs()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _text(s, Inches(0.5), Inches(0.3), Inches(12), Inches(1), "Группы в группах", 32)
    outer = s.shapes.add_group_shape()
    inner = outer.shapes.add_group_shape()
    deepest = inner.shapes.add_group_shape()
    for k, grp in enumerate((outer, inner, deepest)):
        tb = grp.shapes.add_textbox(Inches(1 + 3 * k), Inches(2), Inches(2.5), Inches(1))
        tb.text_frame.text = f"уровень {k + 1}"
        grp.shapes.add_picture(str(_png(tmp)), Inches(1 + 3 * k), Inches(3.5), Inches(2.5), Inches(1.5))
    prs.save(p := tmp / "nested_groups.pptx")
    return p


def native_chart_and_table(tmp: Path) -> Path:
    prs = _prs()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _text(s, Inches(0.5), Inches(0.3), Inches(12), Inches(1), "Диаграмма и таблица", 32)
    data = CategoryChartData()
    data.categories = ["а", "б", "в"]
    data.add_series("ряд", (1, 2, 3))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.5), Inches(6), Inches(5), data)
    tbl = s.shapes.add_table(3, 3, Inches(7), Inches(1.5), Inches(5.5), Inches(3)).table
    for r in range(3):
        for c in range(3):
            tbl.cell(r, c).text = f"{r},{c}"
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    _text(s2, Inches(0.5), Inches(0.3), Inches(12), Inches(1), "Спасибо", 40)
    prs.save(p := tmp / "native_chart_table.pptx")
    return p


def placeholders_only(tmp: Path) -> Path:
    """Слайды из плейсхолдеров стандартных лейаутов python-pptx (title, title+content, two content, section)."""
    prs = _prs()
    for layout_idx, texts in ((0, ["Титул", "подзаголовок"]), (1, ["Список", "раз\nдва\nтри"]),
                              (3, ["Две колонки", "лево", "право"]), (2, ["Раздел", "описание"])):
        s = prs.slides.add_slide(prs.slide_layouts[layout_idx])
        for ph, t in zip(s.placeholders, texts):
            ph.text = t
    prs.save(p := tmp / "placeholders_only.pptx")
    return p


TEMPLATES: dict[str, Callable[[Path], Path]] = {
    "empty_presentation": empty_presentation,
    "one_blank_slide": one_blank_slide,
    "four_by_three": four_by_three,
    "a4_portrait": a4_portrait,
    "pictures_only": pictures_only,
    "hundred_shapes": hundred_shapes,
    "nested_groups": nested_groups,
    "native_chart_and_table": native_chart_and_table,
    "placeholders_only": placeholders_only,
}

__all__ = ["TEMPLATES"]
