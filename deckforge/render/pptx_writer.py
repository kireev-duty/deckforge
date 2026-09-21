"""DeckIR → .pptx клонированием слайдов-образцов шаблона.

Выходной файл — сам шаблон: исходные слайды отвязываются, для каждого SlideIR копируется p:cSld образца,
слоты заполняются по slot_id, связи переносятся с переписыванием rId. Ничего не растеризуется.
"""

from __future__ import annotations

import copy
import logging
import re
from collections.abc import Iterable
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part, XmlPart
from pptx.opc.packuri import PackURI
from pptx.parts.slide import SlidePart
from pptx.slide import Slide

from deckforge.core.ir import Box, DeckIR, Element, Exemplar, Paragraph, SlideIR, SlotKind, TemplateDNA
from deckforge.core.ooxml import NS, A, P, R, absolute_bbox, iter_shapes, localname, shape_id, shape_text
from deckforge.core.placeholders import is_placeholder_text
from deckforge.render.charts import add_chart
from deckforge.render.tables import add_table

log = logging.getLogger(__name__)

TEXT_KINDS = {
    SlotKind.TITLE, SlotKind.SUBTITLE, SlotKind.BODY, SlotKind.CAPTION, SlotKind.NUMBER, SlotKind.LABEL,
    SlotKind.FOOTER, SlotKind.DATE, SlotKind.OTHER,
}
PICTURE_KINDS = {SlotKind.PICTURE, SlotKind.ICON}
# поля (номер, колонтитулы) при отсутствии элемента не очищаем
KEEP_IF_UNFILLED = {SlotKind.SLIDE_NUMBER, SlotKind.FOOTER, SlotKind.DATE}
LAYOUT_FIELD_PH = {"sldNum", "dt", "ftr"}
# связи, которые в копии не нужны
SKIP_RELTYPES = {RT.SLIDE_LAYOUT, RT.NOTES_SLIDE, RT.SLIDE}
# бинарные части можно разделять между слайдами, а не копировать
SHARED_RELTYPES = {RT.IMAGE, RT.MEDIA, RT.VIDEO, RT.AUDIO, RT.FONT}
R_ATTRS = (R + "embed", R + "id", R + "link", R + "pict")
FILL_TAGS = ("noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill")
BULLET_TAGS = ("buNone", "buChar", "buAutoNum", "buBlip")
DEFAULT_BULLET_MARL = 285750  # 0.3125", как в Office
DARK_BG_LUMINANCE = 0.45  # ниже — фон тёмный, нативным объектам светлый текст
LIGHT_TEXT = "FFFFFF"
# схемные цвета фона по роли, без резолва темы
SCHEME_LIGHT = {"bg1", "lt1", "bg2", "lt2"}
SCHEME_DARK = {"tx1", "dk1", "tx2", "dk2"}


# ──────────────────────────── публичный API ────────────────────────────


def render_pptx(
    ir: DeckIR, template_path: str | Path, exemplars: Iterable[Exemplar], out_path: str | Path
) -> Path:
    """Собрать колоду из DeckIR на базе шаблона и сохранить в out_path."""
    writer = DeckWriter(template_path, exemplars)
    return writer.write(ir, out_path)


def render_deck(ir: DeckIR, dna: TemplateDNA, out_path: str | Path) -> Path:
    return render_pptx(ir, dna.source_path, dna.exemplars, out_path)


# ──────────────────────────── писатель ────────────────────────────


class DeckWriter:
    def __init__(self, template_path: str | Path, exemplars: Iterable[Exemplar]) -> None:
        self.template_path = Path(template_path)
        self.exemplars = {e.id: e for e in exemplars}
        self.prs = Presentation(str(self.template_path))
        self.package = self.prs.part.package
        self.src_parts: list[SlidePart] = [s.part for s in self.prs.slides]
        self._cloned: dict[Part, Part] = {}  # исходная часть → копия в пределах слайда
        self._n_source = len(self.src_parts)

    def _detach_source_slides(self) -> None:
        """Убрать исходные слайды из показа — только перед сохранением.

        Пока они в p:sldIdLst, их части достижимы и next_partname не выдаст занятое имя."""
        sld_id_lst = self.prs.slides._sldIdLst
        for sld_id in list(sld_id_lst)[: self._n_source]:
            self.prs.part.drop_rel(sld_id.rId)
            sld_id_lst.remove(sld_id)
        self._n_source = 0

    def _strip_layout_prompts(self) -> None:
        """Сделать контентные плейсхолдеры лейаутов невидимыми: без подсказки, заливки и обводки.

        PowerPoint их на слайдах не показывает, а LibreOffice рисует в PDF/PNG, если у плейсхолдера
        нет пары на слайде. Геометрия остаётся, наследуемую заливку слайды уже получили явно."""
        for layout in self.prs.slide_layouts:
            for sp in layout.shapes._spTree.iter(P + "sp"):
                ph = sp.find("p:nvSpPr/p:nvPr/p:ph", NS)
                if ph is None or ph.get("type") in LAYOUT_FIELD_PH:
                    continue
                for para in sp.findall("p:txBody/a:p", NS):
                    for run in para.findall("a:r", NS) + para.findall("a:fld", NS) + para.findall("a:br", NS):
                        para.remove(run)
                sp_pr = sp.find("p:spPr", NS)
                if sp_pr is None:
                    continue
                for child in list(sp_pr):
                    if localname(child) in FILL_TAGS or localname(child) == "ln":
                        sp_pr.remove(child)
                if sp.find("p:style", NS) is not None:  # заливка по ссылке на тему — тоже гасим
                    no_fill = etree.Element(A + "noFill")
                    _insert_after_geom(sp_pr, no_fill)
                    ln = etree.Element(A + "ln")
                    etree.SubElement(ln, A + "noFill")
                    no_fill.addnext(ln)

    @staticmethod
    def _drop_empty_placeholders(slide: Slide) -> None:
        """Убрать плейсхолдеры слайда, в которых после заполнения ничего нет.

        В показе их не видно, но в редакторе они показывают подсказку лейаута. Поля не трогаем."""
        for sp in list(slide.shapes._spTree.iter(P + "sp")):
            ph = sp.find("p:nvSpPr/p:nvPr/p:ph", NS)
            if ph is None or ph.get("type") in LAYOUT_FIELD_PH:
                continue
            has_text = any((t.text or "").strip() for t in sp.iter(A + "t")) or sp.find(".//a:fld", NS) is not None
            has_image = sp.find(".//a:blip", NS) is not None
            # своя заливка/обводка — уже элемент дизайна, а не пустая подсказка
            if not has_text and not has_image and not _has_visible_frame(sp):
                _remove(sp)

    @staticmethod
    def _materialize_placeholders(slide: Slide) -> None:
        """Перенести в плейсхолдеры слайда наследуемые от лейаута xfrm, заливку и обводку."""
        layout_tree = slide.slide_layout.shapes._spTree
        by_idx: dict[str, etree._Element] = {}
        by_type: dict[str, etree._Element] = {}
        for lsp in layout_tree.iter(P + "sp"):
            ph = lsp.find("p:nvSpPr/p:nvPr/p:ph", NS)
            if ph is None:
                continue
            if ph.get("idx"):
                by_idx[ph.get("idx")] = lsp
            by_type.setdefault(ph.get("type") or "body", lsp)
        for sp in slide.shapes._spTree.iter(P + "sp"):
            ph = sp.find("p:nvSpPr/p:nvPr/p:ph", NS)
            if ph is None or ph.get("type") in LAYOUT_FIELD_PH:
                continue
            lsp = by_idx.get(ph.get("idx") or "")
            if lsp is None:
                lsp = by_type.get(ph.get("type") or "body")
            lay_pr = lsp.find("p:spPr", NS) if lsp is not None else None
            if lay_pr is None:
                continue
            sp_pr = sp.find("p:spPr", NS)
            if sp_pr is None:
                sp_pr = etree.SubElement(sp, P + "spPr")
            have = {localname(c) for c in sp_pr}
            if "xfrm" not in have and (xfrm := lay_pr.find("a:xfrm", NS)) is not None:
                sp_pr.insert(0, copy.deepcopy(xfrm))
            if not have & set(FILL_TAGS):
                fill = next((c for c in lay_pr if localname(c) in FILL_TAGS), None)
                if fill is not None:
                    _insert_after_geom(sp_pr, copy.deepcopy(fill))
            if "ln" not in have and (ln := lay_pr.find("a:ln", NS)) is not None:
                anchor = next((c for c in sp_pr if localname(c) in FILL_TAGS), None)
                if anchor is None:
                    _insert_after_geom(sp_pr, copy.deepcopy(ln))
                else:
                    anchor.addnext(copy.deepcopy(ln))

    # ── основной цикл ──

    def write(self, ir: DeckIR, out_path: str | Path) -> Path:
        sw, sh = self.prs.slide_width, self.prs.slide_height
        if (ir.slide_w, ir.slide_h) != (sw, sh):
            log.warning("размер слайда в DeckIR %sx%s не совпадает с шаблоном %sx%s", ir.slide_w, ir.slide_h, sw, sh)
        for slide_ir in ir.slides:
            self.add_slide(slide_ir)
        self._detach_source_slides()
        self._strip_layout_prompts()
        out = Path(out_path)
        out.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(out))
        return out

    def add_slide(self, slide_ir: SlideIR) -> Slide:
        exemplar = self.exemplars.get(slide_ir.exemplar_id)
        if exemplar is None:
            raise KeyError(f"образец {slide_ir.exemplar_id!r} не найден среди {sorted(self.exemplars)}")
        src = self.src_parts[exemplar.source_index]
        slide = self._clone_slide(src)
        self._cloned = {}
        # связи переносим до заполнения, иначе rId новых частей коллидируют с rId образца
        self._copy_rels(src, slide.part, slide.part._element)
        self._fill(slide, slide_ir, exemplar)
        self._drop_empty_placeholders(slide)
        self._materialize_placeholders(slide)
        self._prune_rels(slide.part)
        if slide_ir.notes:
            slide.notes_slide.notes_text_frame.text = slide_ir.notes
        return slide

    def _clone_slide(self, src: SlidePart) -> Slide:
        layout_part = src.part_related_by(RT.SLIDE_LAYOUT)
        slide = self.prs.slides.add_slide(layout_part.slide_layout)
        root = slide.part._element
        for tag in ("p:cSld", "p:clrMapOvr"):
            old = root.find(tag, NS)
            new = src._element.find(tag, NS)
            if new is None:
                continue
            new = copy.deepcopy(new)
            if old is not None:
                root.replace(old, new)
            else:
                root.append(new)
        # у Slide закэширован shapes над старым spTree
        return Slide(root, slide.part)

    # ── связи ──

    def _copy_rels(self, src: Part, dst: Part, xml: etree._Element, only_used: bool = True) -> None:
        """Перенести связи src на dst и переписать rId в xml.

        Для слайда — только те, на которые ссылается XML; для чарта — все (chartStyle без r:id)."""
        used = {el.get(attr) for el in xml.iter() for attr in R_ATTRS if el.get(attr)}
        mapping: dict[str, str | None] = {}
        for rId, rel in list(src.rels.items()):
            if only_used and rId not in used:
                continue
            if rel.is_external:
                mapping[rId] = dst.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
            elif rel.reltype in SKIP_RELTYPES:
                mapping[rId] = None
            elif rel.reltype in SHARED_RELTYPES:
                mapping[rId] = dst.relate_to(rel.target_part, rel.reltype)
            else:  # chart, diagram, oleObject, xlsx — у каждой копии свой экземпляр
                mapping[rId] = dst.relate_to(self._clone_part(rel.target_part), rel.reltype)
        for el in list(xml.iter()):
            for attr in R_ATTRS:
                old = el.get(attr)
                if old is None or old not in mapping:
                    continue
                new = mapping[old]
                if new is None:
                    # ссылка на другой слайд исходника: гиперссылку убираем
                    if localname(el) in ("hlinkClick", "hlinkHover", "hlinkMouseOver"):
                        el.getparent().remove(el)
                    else:
                        del el.attrib[attr]
                else:
                    el.set(attr, new)

    @staticmethod
    def _prune_rels(part: Part) -> None:
        """Убрать связи, на которые после заполнения не ссылается XML; неявные (лейаут, notes) не трогаем."""
        used = {el.get(attr) for el in part._element.iter() for attr in R_ATTRS if el.get(attr)}
        for rId, rel in list(part.rels.items()):
            if rId not in used and rel.reltype not in SKIP_RELTYPES:
                part.rels.pop(rId)

    def _clone_part(self, part: Part) -> Part:
        """Глубокая копия части с новым именем; связи копируются рекурсивно."""
        if part in self._cloned:
            return self._cloned[part]
        new = type(part).load(self.package.next_partname(_partname_template(part.partname)),
                              part.content_type, self.package, part.blob)
        self._cloned[part] = new
        if isinstance(new, XmlPart):
            self._copy_rels(part, new, new._element, only_used=False)
        else:
            for rel in part.rels.values():
                if rel.is_external:
                    new.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
                else:
                    new.relate_to(self._clone_part(rel.target_part), rel.reltype)
        return new

    # ── заполнение слотов ──

    def _fill(self, slide: Slide, slide_ir: SlideIR, exemplar: Exemplar) -> None:
        sp_tree = slide.part._element.find("p:cSld/p:spTree", NS)
        shapes = {shape_id(sp): sp for sp in iter_shapes(sp_tree)}
        slot_ids = {s.id for s in exemplar.slots}
        filled: set[str] = set()
        for el in slide_ir.elements:
            sp = shapes.get(el.slot_id)
            if sp is None:
                log.warning("слайд %s: слот %s не найден в образце %s", slide_ir.idx, el.slot_id, exemplar.id)
                continue
            if el.kind in (SlotKind.CHART, SlotKind.TABLE) and (el.chart or el.table):
                self._replace_with_native(slide, sp, el, shapes, slot_ids, set(exemplar.fixed))
                filled.add(el.slot_id)
            elif el.kind in (SlotKind.PICTURE, SlotKind.ICON) and el.image_path:
                # битый файл — слот остаётся незаполненным и чистится ниже
                try:
                    self._fill_picture(slide, sp, el, shapes, slot_ids)
                except (OSError, ValueError) as e:
                    log.warning("слайд %s: картинка %s не вставлена (%s) — слот очищен", slide_ir.idx, el.image_path, e)
                    continue
                filled.add(el.slot_id)
            elif el.paragraphs:
                fill_text(sp, el.paragraphs, el.style_overrides)
                filled.add(el.slot_id)
        # незаполненные слоты очищаем, чтобы не остался текст образца
        unfilled: list[Box] = []
        for slot in exemplar.slots:
            if slot.id in filled or slot.kind in KEEP_IF_UNFILLED or slot.kind not in TEXT_KINDS | PICTURE_KINDS:
                continue
            sp = shapes.get(slot.id)
            if sp is None:
                continue
            if slot.kind in PICTURE_KINDS:
                self._clear_picture_captions(sp, el_box=slot.box, shapes=shapes, slot_ids=slot_ids)
                if sp.find("p:nvSpPr/p:nvPr/p:ph", NS) is not None:
                    _remove(sp)
                elif sp.find("p:txBody", NS) is not None:
                    clear_text(sp)
                continue
            unfilled.append(slot.box)
            if sp.find("p:nvSpPr/p:nvPr/p:ph", NS) is not None or _has_visible_frame(sp):
                # пустой плейсхолдер показывает подсказку, слот с рамкой — пустую карточку: удаляем целиком
                _remove(sp)
            elif sp.find("p:txBody", NS) is not None:
                clear_text(sp)
        fixed = set(exemplar.fixed)
        self._remove_empty_containers(shapes, unfilled, [e.box for e in slide_ir.elements if e.slot_id in filled],
                                      slot_ids, fixed)
        # текст образца вне слотов и фиксированных элементов не переносим; короткий декор вроде «01» оставляем
        for sid, sp in shapes.items():
            if sid in slot_ids or sp.getparent() is None or localname(sp) != "sp" or sp.find("p:txBody", NS) is None:
                continue
            text = shape_text(sp)
            # заглушка в нескольких абзацах — тоже заглушка, даже у «фиксированного» элемента
            if is_placeholder_text(" ".join(t.text or "" for t in sp.iter(A + "t"))) or sid not in fixed and _is_sample_text(text):
                clear_text(sp)

    def _remove_empty_containers(
        self, shapes: dict[str, etree._Element], unfilled: list[Box], filled: list[Box],
        slot_ids: set[str], fixed: set[str],
    ) -> None:
        """Убрать фон карточки образца, в которой не осталось заполненных слотов, вместе с её декором."""
        if not unfilled:
            return
        max_area = 0.5 * self.prs.slide_width * self.prs.slide_height
        centers = [(b.x + b.w / 2, b.y + b.h / 2) for b in unfilled]
        kept = [(b.x + b.w / 2, b.y + b.h / 2) for b in filled]
        for sid, sp in list(shapes.items()):
            if sid in slot_ids or sid in fixed or sp.getparent() is None or localname(sp) not in ("sp", "pic"):
                continue
            bb = absolute_bbox(sp)
            if bb is None or bb[2] * bb[3] > max_area:
                continue
            box = Box(x=bb[0], y=bb[1], w=bb[2], h=bb[3])
            inside = lambda c: box.x <= c[0] <= box.x2 and box.y <= c[1] <= box.y2
            if not any(inside(c) for c in centers) or any(inside(c) for c in kept):
                continue
            for oid, other in shapes.items():
                if oid != sid and oid not in slot_ids and oid not in fixed and other.getparent() is not None \
                        and _center_inside(other, box) and (ob := absolute_bbox(other)) and ob[2] * ob[3] < bb[2] * bb[3]:
                    _remove(other)
            _remove(sp)

    def _replace_with_native(
        self, slide: Slide, sp: etree._Element, el: Element, shapes: dict[str, etree._Element],
        slot_ids: set[str], fixed: set[str],
    ) -> None:
        """Убрать фигуру слота (для «нарисованной» диаграммы — всё в её боксе), вставить нативный объект."""
        keep = (slot_ids - {el.slot_id}) | fixed
        _remove(sp)
        for sid, other in shapes.items():
            if sid == el.slot_id or sid in keep or other.getparent() is None:
                continue
            if _center_inside(other, el.box):
                _remove(other)
        overrides = dict(el.style_overrides)
        lum = background_luminance(slide)
        if lum is not None and lum < DARK_BG_LUMINANCE:
            # на тёмном фоне текст палитры не виден — подписи белые
            overrides["text_color"] = LIGHT_TEXT
        if el.chart:
            add_chart(slide, el.chart, el.box, overrides)
        elif el.table:
            add_table(slide, el.table, el.box, overrides)

    def _fill_picture(
        self, slide: Slide, sp: etree._Element, el: Element, shapes: dict[str, etree._Element], slot_ids: set[str]
    ) -> None:
        image_part, rId = slide.part.get_or_add_image_part(el.image_path)
        img_w, img_h = image_part.image.size
        bb = absolute_bbox(sp)
        frame_w, frame_h = (bb[2], bb[3]) if bb else (el.box.w, el.box.h)
        src_rect = crop_rect(img_w, img_h, frame_w, frame_h)
        if localname(sp) == "pic":
            blip_fill = sp.find("p:blipFill", NS)
            if blip_fill is None:
                return
            _set_blip(blip_fill, rId, src_rect)
        elif sp.find("p:nvSpPr/p:nvPr/p:ph[@type='pic']", NS) is not None:
            # плейсхолдер картинки: как PowerPoint при вставке — p:sp становится p:pic, иначе LibreOffice
            # и аудит видят пустой плейсхолдер, а blipFill на нём не рисуется
            pic = _placeholder_to_pic(sp, rId, src_rect)
            self._clear_picture_captions(pic, el.box, shapes, slot_ids)
            shapes[shape_id(pic)] = pic
        else:
            sp_pr = sp.find("p:spPr", NS)
            if sp_pr is None:
                return
            for child in list(sp_pr):
                if localname(child) in FILL_TAGS:
                    sp_pr.remove(child)
            blip_fill = etree.Element(A + "blipFill")
            geom = next((c for c in sp_pr if localname(c) in ("prstGeom", "custGeom")), None)
            if geom is not None:
                geom.addnext(blip_fill)
            else:
                sp_pr.append(blip_fill)
            _set_blip(blip_fill, rId, src_rect)
            if sp.find("p:txBody", NS) is not None:
                clear_text(sp)  # «Вставить фото» внутри самой рамки
            self._clear_picture_captions(sp, el.box, shapes, slot_ids)

    @staticmethod
    def _clear_picture_captions(sp: etree._Element, el_box: Box, shapes: dict[str, etree._Element], slot_ids: set[str]) -> None:
        """Убрать подписи зоны картинки («Вставить фото») — короткие текстовые фигуры поверх рамки."""
        bb = absolute_bbox(sp)
        zone = Box(x=bb[0], y=bb[1], w=bb[2], h=bb[3]) if bb else el_box
        for sid, other in shapes.items():
            if sid == shape_id(sp) or sid in slot_ids or other.getparent() is None:
                continue
            text = shape_text(other).strip()
            if localname(other) == "sp" and 0 < len(text) <= 40 and _center_inside(other, zone):
                _remove(other)


# ──────────────────────────── текст ────────────────────────────


def fill_text(sp: etree._Element, paragraphs: list[Paragraph], overrides: dict | None = None) -> None:
    """Заменить абзацы p:txBody, унаследовав pPr/rPr от абзацев образца (по уровню lvl)."""
    tx = sp.find("p:txBody", NS)
    if tx is None:
        return
    overrides = overrides or {}
    old_paras = tx.findall("a:p", NS)
    by_lvl: dict[int, etree._Element] = {}
    for p in old_paras:
        if p.find("a:r", NS) is None:
            continue
        by_lvl.setdefault(_lvl(p), p)
    fallback = next(iter(by_lvl.values()), old_paras[0] if old_paras else None)
    for p in old_paras:
        tx.remove(p)
    for para in paragraphs:
        tmpl = by_lvl.get(para.level, fallback)
        tx.append(_build_paragraph(para, tmpl, overrides))


def clear_text(sp: etree._Element) -> None:
    """Оставить один пустой абзац с endParaRPr образца — фигура и её стиль сохраняются."""
    tx = sp.find("p:txBody", NS)
    if tx is None:
        return
    old = tx.findall("a:p", NS)
    first = old[0] if old else None
    for p in old:
        tx.remove(p)
    p = etree.SubElement(tx, A + "p")
    if first is not None:
        ppr = first.find("a:pPr", NS)
        if ppr is not None:
            p.append(copy.deepcopy(ppr))
        epr = first.find("a:endParaRPr", NS)
        if epr is None:
            rpr = first.find("a:r/a:rPr", NS)
            if rpr is not None:
                epr = copy.deepcopy(rpr)
                epr.tag = A + "endParaRPr"
        if epr is not None:
            p.append(copy.deepcopy(epr))


def _lvl(p: etree._Element) -> int:
    ppr = p.find("a:pPr", NS)
    return int(ppr.get("lvl", "0")) if ppr is not None else 0


def _build_paragraph(para: Paragraph, tmpl: etree._Element | None, overrides: dict) -> etree._Element:
    p = etree.Element(A + "p")
    ppr = copy.deepcopy(tmpl.find("a:pPr", NS)) if tmpl is not None and tmpl.find("a:pPr", NS) is not None else None
    if ppr is None:
        ppr = etree.Element(A + "pPr")
    if para.level:
        ppr.set("lvl", str(para.level))
    elif ppr.get("lvl"):
        del ppr.attrib["lvl"]
    _apply_bullet(ppr, para.bullet)
    if len(ppr) or ppr.attrib:
        p.append(ppr)

    base_rpr = _template_rpr(tmpl)
    for run in para.runs:
        rpr = copy.deepcopy(base_rpr)
        _apply_run_style(rpr, overrides)
        _apply_run_style(rpr, {k: v for k, v in run.model_dump().items() if k != "text" and v not in (None, False)})
        lines = run.text.split("\n")
        for i, line in enumerate(lines):
            if i:
                br = etree.SubElement(p, A + "br")
                br.append(copy.deepcopy(rpr))
            r = etree.SubElement(p, A + "r")
            r.append(copy.deepcopy(rpr))
            t = etree.SubElement(r, A + "t")
            t.text = line
    epr = tmpl.find("a:endParaRPr", NS) if tmpl is not None else None
    if epr is not None:
        p.append(copy.deepcopy(epr))
    return p


def _template_rpr(tmpl: etree._Element | None) -> etree._Element:
    if tmpl is not None:
        rpr = tmpl.find("a:r/a:rPr", NS)
        if rpr is None:
            rpr = tmpl.find("a:fld/a:rPr", NS)
        if rpr is None:
            epr = tmpl.find("a:endParaRPr", NS)
            if epr is not None:
                rpr = copy.deepcopy(epr)
                rpr.tag = A + "rPr"
        if rpr is not None:
            rpr = copy.deepcopy(rpr)
            # гиперссылки образца не переносим
            for h in rpr.findall("a:hlinkClick", NS) + rpr.findall("a:hlinkMouseOver", NS):
                rpr.remove(h)
            return rpr
    rpr = etree.Element(A + "rPr")
    rpr.set("lang", "ru-RU")
    return rpr


def _apply_bullet(ppr: etree._Element, bullet: bool) -> None:
    has = [c for c in ppr if localname(c) in BULLET_TAGS]
    if bullet:
        if not has or all(localname(c) == "buNone" for c in has):
            for c in has:
                ppr.remove(c)
            if int(ppr.get("marL") or 0) <= 0 or int(ppr.get("indent") or 0) >= 0:
                ppr.set("marL", str(DEFAULT_BULLET_MARL))
                ppr.set("indent", str(-DEFAULT_BULLET_MARL))
            bu = etree.Element(A + "buChar")
            bu.set("char", "•")
            _insert_bullet(ppr, bu)
    elif any(localname(c) != "buNone" for c in has):
        for c in has:
            ppr.remove(c)
        _insert_bullet(ppr, etree.Element(A + "buNone"))


def _insert_bullet(ppr: etree._Element, bu: etree._Element) -> None:
    """a:bu* идёт после lnSpc/spcBef/spcAft/buClr/buSzPct/buFont и перед tabLst/defRPr."""
    after = {"lnSpc", "spcBef", "spcAft", "buClrTx", "buClr", "buSzTx", "buSzPct", "buSzPts", "buFontTx", "buFont"}
    pos = 0
    for i, c in enumerate(ppr):
        if localname(c) in after:
            pos = i + 1
    ppr.insert(pos, bu)


def _apply_run_style(rpr: etree._Element, style: dict) -> None:
    """bold/italic/size_pt/color/font поверх rPr образца."""
    if "bold" in style:
        rpr.set("b", "1" if style["bold"] else "0")
    if "italic" in style:
        rpr.set("i", "1" if style["italic"] else "0")
    if style.get("size_pt"):
        rpr.set("sz", str(int(round(float(style["size_pt"]) * 100))))
    if style.get("color"):
        for c in list(rpr):
            if localname(c) in FILL_TAGS:
                rpr.remove(c)
        fill = etree.Element(A + "solidFill")
        clr = etree.SubElement(fill, A + "srgbClr")
        clr.set("val", str(style["color"]).lstrip("#").upper())
        _insert_fill(rpr, fill)
    if style.get("font"):
        for tag in ("latin", "ea", "cs"):
            el = rpr.find(f"a:{tag}", NS)
            if el is None:
                el = etree.Element(A + tag)
                _insert_font(rpr, el)
            el.set("typeface", str(style["font"]))


def _insert_fill(rpr: etree._Element, fill: etree._Element) -> None:
    """Порядок детей rPr: ln, fill, effect*, highlight, uLn*, uFill*, latin, ea, cs, sym, hlink*."""
    pos = 0
    for i, c in enumerate(rpr):
        if localname(c) == "ln":
            pos = i + 1
    rpr.insert(pos, fill)


def _insert_font(rpr: etree._Element, el: etree._Element) -> None:
    order = ["latin", "ea", "cs", "sym", "hlinkClick", "hlinkMouseOver", "rtl", "extLst"]
    idx = order.index(localname(el))
    pos = len(rpr)
    for i, c in enumerate(rpr):
        if localname(c) in order and order.index(localname(c)) > idx:
            pos = i
            break
    rpr.insert(pos, el)


# ──────────────────────────── картинки ────────────────────────────


def crop_rect(img_w: int, img_h: int, frame_w: int, frame_h: int) -> tuple[int, int, int, int] | None:
    """(l, t, r, b) в 1/1000 % для a:srcRect: center-crop картинки под аспект рамки. None — аспект совпал."""
    if not (img_w and img_h and frame_w and frame_h):
        return None
    img_ar, frame_ar = img_w / img_h, frame_w / frame_h
    if abs(img_ar - frame_ar) / frame_ar < 0.01:
        return None
    if img_ar > frame_ar:  # картинка шире — режем бока
        keep = frame_ar / img_ar
        cut = int(round((1 - keep) / 2 * 100000))
        return cut, 0, cut, 0
    keep = img_ar / frame_ar
    cut = int(round((1 - keep) / 2 * 100000))
    return 0, cut, 0, cut


def _placeholder_to_pic(sp: etree._Element, rId: str, src_rect: tuple[int, int, int, int] | None) -> etree._Element:
    """Заменить p:sp-плейсхолдер картинки на p:pic с тем же id, ph и геометрией (txBody и style у картинки нет)."""
    pic = etree.Element(P + "pic")
    nv = etree.SubElement(pic, P + "nvPicPr")
    nv.append(copy.deepcopy(sp.find("p:nvSpPr/p:cNvPr", NS)))
    locks = etree.SubElement(etree.SubElement(nv, P + "cNvPicPr"), A + "picLocks")
    locks.set("noGrp", "1")
    locks.set("noChangeAspect", "1")
    nv.append(copy.deepcopy(sp.find("p:nvSpPr/p:nvPr", NS)))
    blip_fill = etree.SubElement(pic, P + "blipFill")
    _set_blip(blip_fill, rId, src_rect)
    sp_pr = copy.deepcopy(sp.find("p:spPr", NS)) if sp.find("p:spPr", NS) is not None else etree.Element(P + "spPr")
    for child in list(sp_pr):
        if localname(child) in FILL_TAGS:
            sp_pr.remove(child)
    pic.append(sp_pr)
    sp.addnext(pic)
    sp.getparent().remove(sp)
    return pic


def _set_blip(blip_fill: etree._Element, rId: str, src_rect: tuple[int, int, int, int] | None) -> None:
    blip = blip_fill.find("a:blip", NS)
    if blip is None:
        blip = etree.Element(A + "blip")
        blip_fill.insert(0, blip)
    blip.set(R + "embed", rId)
    for c in list(blip):  # эффекты образца к новой картинке не относятся
        blip.remove(c)
    old = blip_fill.find("a:srcRect", NS)
    if old is not None:
        blip_fill.remove(old)
    if src_rect is not None:
        sr = etree.Element(A + "srcRect")
        for k, v in zip(("l", "t", "r", "b"), src_rect):
            if v:
                sr.set(k, str(v))
        blip.addnext(sr)
    if blip_fill.find("a:stretch", NS) is None and blip_fill.find("a:tile", NS) is None:
        st = etree.SubElement(blip_fill, A + "stretch")
        etree.SubElement(st, A + "fillRect")


# ──────────────────────────── фон слайда ────────────────────────────


def background_luminance(slide: Slide) -> float | None:
    """Яркость фона слайда 0..1 (свой `p:bg`, иначе лейаута, иначе мастера); None — определить нельзя."""
    for part in (slide.part, slide.slide_layout.part, slide.slide_layout.slide_master.part):
        bg = part._element.find(".//p:cSld/p:bg", NS)
        if bg is None:
            continue
        return _fill_luminance(bg, part)
    return None


def _fill_luminance(bg: etree._Element, part: Part) -> float | None:
    if (blip := bg.find(".//a:blipFill/a:blip", NS)) is not None and blip.get(R + "embed"):
        try:
            from io import BytesIO

            from PIL import Image, ImageStat

            image_part = part.related_part(blip.get(R + "embed"))
            with Image.open(BytesIO(image_part.blob)) as im:
                small = im.convert("L").resize((32, 18))
                return ImageStat.Stat(small).mean[0] / 255.0
        except Exception as e:  # noqa: BLE001
            log.debug("фон-картинка не прочитана: %s", e)
            return None
    clr = bg.find(".//a:solidFill/*", NS)
    if clr is None:
        clr = bg.find(".//a:gradFill/a:gsLst/a:gs/*", NS)
    if clr is None:
        clr = bg.find("p:bgRef/*", NS)
    if clr is None:
        return None
    if localname(clr) == "srgbClr":
        v = clr.get("val", "")
        if len(v) == 6:
            r, g, b = (int(v[i:i + 2], 16) / 255.0 for i in (0, 2, 4))
            return 0.2126 * r + 0.7152 * g + 0.0722 * b
        return None
    if localname(clr) == "schemeClr":
        role = clr.get("val", "")
        return 1.0 if role in SCHEME_LIGHT else 0.0 if role in SCHEME_DARK else None
    return None


# ──────────────────────────── геометрия и служебное ────────────────────────────


def _center_inside(sp: etree._Element, box: Box) -> bool:
    bb = absolute_bbox(sp)
    if bb is None:
        return False
    cx, cy = bb[0] + bb[2] / 2, bb[1] + bb[3] / 2
    return box.x <= cx <= box.x2 and box.y <= cy <= box.y2


def _insert_after_geom(sp_pr: etree._Element, el: etree._Element) -> None:
    """Вставить el в spPr на место заливки по схеме: после prstGeom/custGeom (или xfrm), иначе в начало."""
    anchor = next((c for c in sp_pr if localname(c) in ("prstGeom", "custGeom")), None)
    anchor = anchor if anchor is not None else sp_pr.find("a:xfrm", NS)
    if anchor is not None:
        anchor.addnext(el)
    else:
        sp_pr.insert(0, el)


def _has_visible_frame(sp: etree._Element) -> bool:
    """У фигуры своя заливка или обводка — без текста она останется пустой рамкой."""
    sp_pr = sp.find("p:spPr", NS)
    if sp_pr is None:
        return False
    if any(localname(c) in ("solidFill", "gradFill", "pattFill", "blipFill") for c in sp_pr):
        return True
    ln = sp_pr.find("a:ln", NS)
    if ln is not None:
        return ln.find("a:noFill", NS) is None and any(localname(c) in ("solidFill", "gradFill") for c in ln)
    # p:style: idx="0" — нет заливки/обводки
    style = sp.find("p:style", NS)
    if style is not None and not any(localname(c) == "noFill" for c in sp_pr):
        for ref in ("a:lnRef", "a:fillRef"):
            el = style.find(ref, NS)
            if el is not None and el.get("idx", "0") != "0":
                return True
    return False


def _is_sample_text(text: str) -> bool:
    """Текст-образец, а не декор: хотя бы два буквенных символа."""
    return sum(ch.isalpha() for ch in text) >= 2


def _remove(sp: etree._Element) -> None:
    parent = sp.getparent()
    if parent is None:
        return
    parent.remove(sp)
    # пустую группу после удаления детей тоже убираем
    if parent.tag == P + "grpSp" and not any(localname(c) in ("sp", "pic", "cxnSp", "graphicFrame", "grpSp") for c in parent):
        _remove(parent)


def _partname_template(partname: PackURI) -> str:
    """'/ppt/charts/chart3.xml' → '/ppt/charts/chart%d.xml'; без цифр — '/ppt/embeddings/x.xlsx' → 'x%d.xlsx'."""
    s = str(partname)
    m = re.search(r"(\d+)(\.[^./]+)$", s)
    if m:
        return s[: m.start(1)] + "%d" + m.group(2)
    stem, ext = s.rsplit(".", 1) if "." in s.rsplit("/", 1)[-1] else (s, "")
    return f"{stem}%d.{ext}" if ext else f"{s}%d"


__all__ = ["DeckWriter", "clear_text", "crop_rect", "fill_text", "render_deck", "render_pptx"]
