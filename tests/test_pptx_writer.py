"""render/pptx_writer: клонирование образцов, перенос rels, заполнение слотов."""

from __future__ import annotations

from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation
from pptx.parts.chart import ChartPart
from pptx.util import Inches, Pt

from deckforge.core.ir import (
    ChartSpec,
    DeckIR,
    Element,
    Paragraph,
    SlideIR,
    SlotKind,
    TextRun,
)
from deckforge.core.ooxml import NS, R, absolute_bbox, iter_shapes, shape_id, shape_text
from deckforge.parsing.layout_classifier import classify_template
from deckforge.render.pptx_writer import crop_rect, fill_text, render_pptx

# ──────────────────────────── fill_text на синтетическом слайде ────────────────────────────


def _textbox_slide():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
    p = tb.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "Образец"
    run.font.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = __import__("pptx.dml.color", fromlist=["RGBColor"]).RGBColor(0xFF, 0x00, 0x53)
    return prs, tb._element


def test_fill_text_keeps_formatting():
    _, sp = _textbox_slide()
    paras = [
        Paragraph(runs=[TextRun(text="Первый\nвторая строка")], bullet=True),
        Paragraph(runs=[TextRun(text="Второй", italic=True)], level=1, bullet=True),
    ]
    fill_text(sp, paras)
    ps = sp.findall("p:txBody/a:p", NS)
    assert len(ps) == 2
    rpr = ps[0].find("a:r/a:rPr", NS)
    assert rpr.get("b") == "1" and rpr.get("sz") == "2000"
    assert rpr.find("a:solidFill/a:srgbClr", NS).get("val") == "FF0053"
    assert ps[0].find("a:br", NS) is not None and shape_text(sp) == "Первыйвторая строкаВторой"
    assert ps[0].find("a:pPr/a:buChar", NS) is not None
    assert ps[1].find("a:pPr", NS).get("lvl") == "1"
    assert ps[1].find("a:r/a:rPr", NS).get("i") == "1"


def test_fill_text_overrides_and_no_bullet():
    _, sp = _textbox_slide()
    fill_text(sp, [Paragraph(runs=[TextRun(text="x")])], {"font": "Play", "size_pt": 14, "color": "#0077FF"})
    rpr = sp.find("p:txBody/a:p/a:r/a:rPr", NS)
    assert rpr.get("sz") == "1400"
    assert rpr.find("a:latin", NS).get("typeface") == "Play"
    assert rpr.find("a:solidFill/a:srgbClr", NS).get("val") == "0077FF"
    assert sp.find("p:txBody/a:p/a:pPr/a:buChar", NS) is None


def test_crop_rect():
    assert crop_rect(400, 300, 1600, 900) == (0, 12500, 0, 12500)  # 4:3 в 16:9
    assert crop_rect(1600, 900, 400, 400) == (21875, 0, 21875, 0)  # широкая в квадрат
    assert crop_rect(160, 90, 1600, 900) is None


# ──────────────────────────── клонирование на реальном шаблоне ────────────────────────────


def _deck(exemplars, slides):
    return DeckIR(template_id="t", strategy="test", slide_w=0, slide_h=0, slides=slides)


def _no_duplicate_entries(path: Path) -> bool:
    """Новая картинка не получает имя части образца."""
    import zipfile

    names = zipfile.ZipFile(path).namelist()
    return len(names) == len(set(names))


def _dangling(prs) -> int:
    n = 0
    for s in prs.slides:
        for el in s.part._element.iter():
            for attr in (R + "embed", R + "id", R + "link"):
                rid = el.get(attr)
                if rid and rid not in s.part.rels:
                    n += 1
    return n


def test_clone_every_exemplar(template_path, tmp_path: Path):
    tpl = template_path("VK Tech")
    src = Presentation(str(tpl))
    n_layouts = sum(len(m.slide_layouts) for m in src.slide_masters)
    n_images = sum(1 for p in src.part.package.iter_parts() if type(p).__name__ == "ImagePart")

    exemplars = [p.to_exemplar() for p in classify_template(tpl)]
    slides = []
    for i, e in enumerate(exemplars):
        title = next((s for s in e.slots if s.kind == SlotKind.TITLE), None)
        elements = []
        if title is not None:
            elements.append(Element(slot_id=title.id, kind=SlotKind.TITLE, box=title.box,
                                    paragraphs=[Paragraph(runs=[TextRun(text=f"Заголовок {i + 1}")])]))
        slides.append(SlideIR(idx=i, exemplar_id=e.id, archetype=e.archetype, elements=elements,
                              outline_ref=i, notes=f"заметка {i + 1}"))
    out = render_pptx(_deck(exemplars, slides), tpl, exemplars, tmp_path / "all.pptx")

    prs = Presentation(str(out))
    assert len(prs.slides) == len(exemplars)
    assert _dangling(prs) == 0
    assert _no_duplicate_entries(out)
    assert sum(len(m.slide_layouts) for m in prs.slide_masters) == n_layouts
    assert sum(1 for p in prs.part.package.iter_parts() if type(p).__name__ == "ImagePart") <= n_images
    assert prs.part._element.find("p:embeddedFontLst", NS) is not None
    for i, (s, e) in enumerate(zip(prs.slides, exemplars)):
        assert s.has_notes_slide and s.notes_slide.notes_text_frame.text == f"заметка {i + 1}"
        texts = {shape_id(sp): shape_text(sp) for sp in iter_shapes(s.part._element.find("p:cSld/p:spTree", NS))}
        title = next((sl for sl in e.slots if sl.kind == SlotKind.TITLE), None)
        if title is not None:
            assert texts[title.id] == f"Заголовок {i + 1}"
        # остальные текстовые слоты очищены от текста-заглушки
        for sl in e.slots:
            if sl.kind in (SlotKind.BODY, SlotKind.LABEL, SlotKind.CAPTION, SlotKind.SUBTITLE, SlotKind.NUMBER):
                assert texts.get(sl.id, "") == "", (e.id, sl.id, texts.get(sl.id))


def test_chart_slide_cloned_twice(template_path, tmp_path: Path):
    tpl = template_path("ЛЦТ2026")
    exemplars = [p.to_exemplar() for p in classify_template(tpl)]
    src = Presentation(str(tpl))
    chart_idx = next(i for i, s in enumerate(src.slides) if any(sh.has_chart for sh in s.shapes))
    e = exemplars[chart_idx]
    chart_slot = next(s for s in e.slots if s.kind == SlotKind.CHART)

    # 1) дважды без ChartSpec → диаграмма образца с его цифрами на слайд не попадает, а её chart-части и xlsx
    #    не остаются в пакете сиротами (связи переносятся только для того, на что ссылается итоговый XML)
    slides = [SlideIR(idx=i, exemplar_id=e.id, archetype=e.archetype, elements=[], outline_ref=i) for i in range(2)]
    prs = Presentation(str(render_pptx(_deck(exemplars, slides), tpl, exemplars, tmp_path / "twice.pptx")))
    assert not [sh for s in prs.slides for sh in s.shapes if sh.has_chart]
    assert not [p for p in prs.part.package.iter_parts() if isinstance(p, ChartPart)]
    assert _dangling(prs) == 0

    # 2) с ChartSpec → образцовой диаграммы нет, есть новая нативная; data_labels → c:dLbls, y_label → ось
    spec = ChartSpec(kind="column", title="Тест", categories=["a", "b"], series={"s1": [1, 2]}, y_label="часы")
    el = Element(slot_id=chart_slot.id, kind=SlotKind.CHART, box=chart_slot.box, chart=spec,
                 style_overrides={"data_labels": True})
    slides = [SlideIR(idx=0, exemplar_id=e.id, archetype=e.archetype, elements=[el], outline_ref=0)]
    prs = Presentation(str(render_pptx(_deck(exemplars, slides), tpl, exemplars, tmp_path / "spec.pptx")))
    charts = [sh for sh in prs.slides[0].shapes if sh.has_chart]
    assert len(charts) == 1 and charts[0].chart.chart_title.text_frame.text == "Тест"
    chart = charts[0].chart
    assert chart.plots[0].has_data_labels and chart.plots[0].data_labels.number_format == "0"
    assert chart.value_axis.has_title and chart.value_axis.axis_title.text_frame.text == "часы"
    assert _dangling(prs) == 0


def test_unfilled_chart_of_exemplar_is_removed(template_path, tmp_path: Path):
    """МТУСИ slide20: две нативные диаграммы, данные — у одной; вторая с цифрами шаблона («1 кв», «Основной») уходит."""
    tpl = template_path("МТУСИ")
    exemplars = [p.to_exemplar() for p in classify_template(tpl)]
    e = next((x for x in exemplars if sum(s.kind == SlotKind.CHART for s in x.slots) == 2), None)
    if e is None:
        pytest.skip("в шаблоне нет образца с двумя диаграммами")
    first = next(s for s in e.slots if s.kind == SlotKind.CHART)
    spec = ChartSpec(kind="column", title="", categories=["a", "b"], series={"s1": [1, 2]})
    el = Element(slot_id=first.id, kind=SlotKind.CHART, box=first.box, chart=spec)
    slides = [SlideIR(idx=0, exemplar_id=e.id, archetype=e.archetype, elements=[el], outline_ref=0)]
    prs = Presentation(str(render_pptx(_deck(exemplars, slides), tpl, exemplars, tmp_path / "one_of_two.pptx")))
    charts = [sh.chart for sh in prs.slides[0].shapes if sh.has_chart]
    assert len(charts) == 1 and list(charts[0].plots[0].categories) == ["a", "b"]
    assert _dangling(prs) == 0


def test_unfilled_text_placeholder_is_removed(template_path, tmp_path: Path):
    """Образцы на плейсхолдерах: незаполненный body-ph удаляется, а не остаётся пустым."""
    tpl = template_path("ЛЦТ2026")
    exemplars = [p.to_exemplar() for p in classify_template(tpl)]
    src = Presentation(str(tpl))
    with_body = [(i, e) for i, e in enumerate(exemplars)
                 if any(sh.is_placeholder and sh.placeholder_format.type is not None
                        and str(sh.placeholder_format.type).startswith("BODY") for sh in src.slides[i].shapes)
                 and any(s.kind == SlotKind.TITLE for s in e.slots)]
    assert with_body, "в holdout ожидался образец с body-плейсхолдером"
    idx, e = with_body[0]
    title = next(s for s in e.slots if s.kind == SlotKind.TITLE)
    el = Element(slot_id=title.id, kind=SlotKind.TITLE, box=title.box, paragraphs=[Paragraph(runs=[TextRun(text="Т")])])
    slides = [SlideIR(idx=0, exemplar_id=e.id, archetype=e.archetype, elements=[el], outline_ref=0)]
    prs = Presentation(str(render_pptx(_deck(exemplars, slides), tpl, exemplars, tmp_path / "ph.pptx")))
    empty_ph = [sh for sh in prs.slides[0].shapes if sh.is_placeholder and sh.has_text_frame and not sh.text_frame.text.strip()
                and str(sh.placeholder_format.type).startswith("BODY")]
    assert not empty_ph
    assert any(sh.has_text_frame and sh.text_frame.text == "Т" for sh in prs.slides[0].shapes)
    # подсказки плейсхолдеров лейаутов сняты, поля (номер слайда, дата) остаются
    for layout in prs.slide_layouts:
        for sh in layout.placeholders:
            ph_type = str(sh.placeholder_format.type)
            if sh.has_text_frame and not any(t in ph_type for t in ("SLIDE_NUMBER", "DATE", "FOOTER")):
                assert not sh.text_frame.text.strip(), f"{layout.name}: {sh.text_frame.text!r}"
    assert any(sh.text_frame.text.strip() for layout in prs.slide_layouts for sh in layout.placeholders
               if sh.has_text_frame and "SLIDE_NUMBER" in str(sh.placeholder_format.type))


def test_empty_placeholders_of_exemplar_are_dropped(template_path, tmp_path: Path):
    """Пустые плейсхолдеры образца (не слоты) писатель убирает; заголовок и номер слайда остаются."""
    tpl = template_path("VK Education")
    exemplars = [p.to_exemplar() for p in classify_template(tpl)]
    src = Presentation(str(tpl))
    def empty_picture_ph(slide) -> bool:
        return any(sh.is_placeholder and "PICTURE" in str(sh.placeholder_format.type)
                   and sh.element.find(".//a:blip", NS) is None for sh in slide.shapes)

    idx, e = next((i, e) for i, e in enumerate(exemplars)
                  if empty_picture_ph(src.slides[i]) and any(s.kind == SlotKind.TITLE for s in e.slots))
    title = next(s for s in e.slots if s.kind == SlotKind.TITLE)
    el = Element(slot_id=title.id, kind=SlotKind.TITLE, box=title.box, paragraphs=[Paragraph(runs=[TextRun(text="Т")])])
    slides = [SlideIR(idx=0, exemplar_id=e.id, archetype=e.archetype, elements=[el], outline_ref=0)]
    prs = Presentation(str(render_pptx(_deck(exemplars, slides), tpl, exemplars, tmp_path / "empty_ph.pptx")))
    for sh in prs.slides[0].shapes:
        if not sh.is_placeholder or "SLIDE_NUMBER" in str(sh.placeholder_format.type):
            continue
        assert sh.has_text_frame and sh.text_frame.text.strip(), f"пустой плейсхолдер остался: {sh.name}"
    assert any(sh.has_text_frame and sh.text_frame.text == "Т" for sh in prs.slides[0].shapes)


def test_picture_placeholder_becomes_pic(template_path, tmp_path: Path):
    """Плейсхолдер картинки (p:sp с ph type=pic) после заполнения — p:pic с тем же id, как делает PowerPoint."""
    tpl = template_path("ЛЦТ2026")
    exemplars = [p.to_exemplar() for p in classify_template(tpl)]
    src = Presentation(str(tpl))
    found = None
    for i, e in enumerate(exemplars):
        tree = src.slides[i].part._element.find("p:cSld/p:spTree", NS)
        for s in e.slots:
            if s.kind != SlotKind.PICTURE:
                continue
            sp = next((x for x in iter_shapes(tree) if shape_id(x) == s.id), None)
            if sp is not None and sp.find("p:nvSpPr/p:nvPr/p:ph[@type='pic']", NS) is not None:
                found = (e, s)
                break
        if found:
            break
    assert found, "в шаблоне нет sp-плейсхолдера картинки"
    e, slot = found
    img = tmp_path / "pic.png"
    Image.new("RGB", (300, 300), "#FF0053").save(img)
    el = Element(slot_id=slot.id, kind=SlotKind.PICTURE, box=slot.box, image_path=str(img))
    slides = [SlideIR(idx=0, exemplar_id=e.id, archetype=e.archetype, elements=[el], outline_ref=0)]
    prs = Presentation(str(render_pptx(_deck(exemplars, slides), tpl, exemplars, tmp_path / "ph_pic.pptx")))
    tree = prs.slides[0].part._element.find("p:cSld/p:spTree", NS)
    pic = next(x for x in iter_shapes(tree) if shape_id(x) == slot.id)
    assert pic.tag.endswith("}pic") and pic.find("p:nvPicPr/p:nvPr/p:ph", NS) is not None
    assert pic.find("p:blipFill/a:blip", NS).get(R + "embed") in prs.slides[0].part.rels
    assert pic.find("p:spPr/a:xfrm", NS) is not None and pic.find("p:txBody", NS) is None


def test_picture_fill_and_crop(template_path, tmp_path: Path):
    tpl = template_path("VK Tech")
    exemplars = [p.to_exemplar() for p in classify_template(tpl)]
    e, slot = next((e, s) for e in exemplars for s in e.slots if s.kind == SlotKind.PICTURE)
    img = tmp_path / "pic.png"
    Image.new("RGB", (400, 300), "#0077FF").save(img)
    el = Element(slot_id=slot.id, kind=SlotKind.PICTURE, box=slot.box, image_path=str(img))
    slides = [SlideIR(idx=0, exemplar_id=e.id, archetype=e.archetype, elements=[el], outline_ref=0)]
    out = render_pptx(_deck(exemplars, slides), tpl, exemplars, tmp_path / "pic.pptx")
    assert _no_duplicate_entries(out)
    prs = Presentation(str(out))
    root = prs.slides[0].part._element
    sp = next(s for s in iter_shapes(root.find("p:cSld/p:spTree", NS)) if shape_id(s) == slot.id)
    blip = sp.find(".//a:blipFill/a:blip", NS)
    assert blip is not None and blip.get(R + "embed") in prs.slides[0].part.rels
    expected = crop_rect(400, 300, slot.box.w, slot.box.h)
    sr = sp.find(".//a:blipFill/a:srcRect", NS)
    if expected is None:
        assert sr is None
    else:
        assert sr is not None and int(sr.get("t", 0)) == expected[1] and int(sr.get("l", 0)) == expected[0]


# ──────────────────────────── пустые рамки: аватар, карточка с иконкой, плашка-«чип» ────────────────────────────


def _text_el(slot, text: str) -> Element:
    return Element(slot_id=slot.id, kind=slot.kind, box=slot.box, paragraphs=[Paragraph(runs=[TextRun(text=text)])])


def _render_one(tpl, exemplars, e, elements, out: Path):
    slides = [SlideIR(idx=0, exemplar_id=e.id, archetype=e.archetype, elements=elements, outline_ref=0)]
    prs = Presentation(str(render_pptx(_deck(exemplars, slides), tpl, exemplars, out)))
    return {shape_id(sp): sp for sp in iter_shapes(prs.slides[0].part._element.find("p:cSld/p:spTree", NS))}


def test_speaker_avatar_removed(template_path, tmp_path: Path):
    """VK WorkSpace: серый кружок-аватар рядом с «Имя Спикера, должность» — заглушка фото, уходит всегда."""
    tpl = template_path("WorkSpace")
    exemplars = [p.to_exemplar() for p in classify_template(tpl)]
    e = next(e for e in exemplars if any(s.sample_text and "Спикер" in s.sample_text for s in e.slots)
             and any(s.kind == SlotKind.TITLE for s in e.slots))
    title = next(s for s in e.slots if s.kind == SlotKind.TITLE)
    speaker = next(s for s in e.slots if s.sample_text and "Спикер" in s.sample_text)
    shapes = _render_one(tpl, exemplars, e, [_text_el(title, "Т"), _text_el(speaker, "Подзаголовок")],
                         tmp_path / "avatar.pptx")
    for sp in shapes.values():
        geom = sp.find("p:spPr/a:prstGeom", NS)
        assert not (geom is not None and geom.get("prst") == "ellipse" and not shape_text(sp).strip()), "аватар остался"


def test_empty_card_takes_its_icon(template_path, tmp_path: Path):
    """Сетка, где body — сама карточка: пустая уходит вместе со своей иконкой; карточка с подписью — остаётся."""
    tpl = template_path("WorkSpace")
    exemplars = [p.to_exemplar() for p in classify_template(tpl)]
    e = next(e for e in exemplars if e.card_frames and sum(s.kind == SlotKind.BODY for s in e.slots) == 3
             and sum(s.kind == SlotKind.ICON for s in e.slots) == 3)
    title = next(s for s in e.slots if s.kind == SlotKind.TITLE)
    bodies = sorted((s for s in e.slots if s.kind == SlotKind.BODY), key=lambda s: s.box.x)
    icons = sorted((s for s in e.slots if s.kind == SlotKind.ICON), key=lambda s: s.box.x)
    labels = sorted((s for s in e.slots if s.kind == SlotKind.LABEL), key=lambda s: s.box.x)

    # две карточки из трёх: третья уходит, её иконка — тоже
    shapes = _render_one(tpl, exemplars, e, [_text_el(title, "Т"), _text_el(bodies[0], "а"), _text_el(bodies[1], "б")],
                         tmp_path / "two.pptx")
    assert bodies[2].id not in shapes and icons[2].id not in shapes
    assert icons[0].id in shapes and icons[1].id in shapes

    # заполнены только подписи внутри карточек: рамки карточек остаются
    if labels:
        shapes = _render_one(tpl, exemplars, e, [_text_el(title, "Т"), *(_text_el(lb, f"п{i}") for i, lb in enumerate(labels))],
                             tmp_path / "labels.pptx")
        assert all(b.id in shapes for b in bodies)


def test_title_chip_plate_grows(template_path, tmp_path: Path):
    """ЛЦТ2026: плашка-«чип» уже бокса заголовка растягивается под длинный заголовок, но не шире plate_max_w."""
    tpl = template_path("ЛЦТ2026")
    exemplars = [p.to_exemplar() for p in classify_template(tpl)]
    e, title = next((e, s) for e in exemplars for s in e.slots if s.kind == SlotKind.TITLE and s.plate_id)
    assert title.max_lines == 1 and title.hard_lines and title.plate_max_w
    src = {shape_id(sp): sp for sp in iter_shapes(
        Presentation(str(tpl)).slides[int(e.id.removeprefix("slide")) - 1].part._element.find("p:cSld/p:spTree", NS))}
    w0 = int(src[title.plate_id].find("p:spPr/a:xfrm/a:ext", NS).get("cx"))
    shapes = _render_one(tpl, exemplars, e, [_text_el(title, "Бесшовная интеграция с корпоративными системами")],
                         tmp_path / "chip.pptx")
    w = int(shapes[title.plate_id].find("p:spPr/a:xfrm/a:ext", NS).get("cx"))
    assert w0 < w <= title.plate_max_w
    # короткий заголовок плашку не трогает
    shapes = _render_one(tpl, exemplars, e, [_text_el(title, "Итоги")], tmp_path / "chip_short.pptx")
    assert int(shapes[title.plate_id].find("p:spPr/a:xfrm/a:ext", NS).get("cx")) == w0


def test_title_box_narrowed_left_of_cards(template_path, tmp_path: Path):
    """VK WorkSpace slide5: колонка карточек справа задевает первую строку заголовка — бокс сужается до карточек
    (Slot.wrap_w), длинный заголовок переносится левее них, а не уходит под карточки."""
    tpl = template_path("WorkSpace")
    exemplars = [p.to_exemplar() for p in classify_template(tpl)]
    e = next(e for e in exemplars if e.id == "slide5")
    title = next(s for s in e.slots if s.kind == SlotKind.TITLE)
    assert title.wrap_w and title.wrap_w < title.box.w
    shapes = _render_one(tpl, exemplars, e, [_text_el(title, "Этап 3: Масштабирование и передача полномочий")],
                         tmp_path / "narrow.pptx")
    assert int(shapes[title.id].find("p:spPr/a:xfrm/a:ext", NS).get("cx")) == title.wrap_w


def test_empty_card_removed_with_its_icon_decor(template_path, tmp_path: Path):
    """VK WorkSpace slide5 (разметка VLM из data/archetypes): значок иконки в карточке — декор, не слот, плашка
    под ним — тоже не слот. Пустая карточка уходит вместе с ними, у заполненной иконка остаётся."""
    from deckforge.parsing.exemplars import load_exemplars

    tpl = template_path("WorkSpace")
    exemplars = load_exemplars(tpl)
    e = next(e for e in exemplars if e.id == "slide5")
    title = next(s for s in e.slots if s.kind == SlotKind.TITLE)
    cards = sorted((s for s in e.slots if s.kind == SlotKind.BODY), key=lambda s: s.box.y)
    assert len(cards) == 3
    src = {shape_id(sp): sp for sp in iter_shapes(
        Presentation(str(tpl)).slides[4].part._element.find("p:cSld/p:spTree", NS))}
    slot_ids = {s.id for s in e.slots}
    inside = lambda card: {sid for sid, sp in src.items() if sid not in slot_ids and sid != card.id
                           and (bb := absolute_bbox(sp)) and card.box.x <= bb[0] and card.box.y <= bb[1]
                           and bb[0] + bb[2] <= card.box.x2 and bb[1] + bb[3] <= card.box.y2}
    decor = [inside(c) for c in cards]
    assert all(decor), "у каждой карточки плашка и значок"
    shapes = _render_one(tpl, exemplars, e, [_text_el(title, "Т"), _text_el(cards[0], "Первая карточка")],
                         tmp_path / "cards.pptx")
    assert decor[0] <= set(shapes), "у заполненной карточки иконка остаётся"
    assert not (decor[1] | decor[2]) & set(shapes), "иконки пустых карточек уходят вместе с ними"
