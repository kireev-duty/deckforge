"""Стресс-инварианты: крайние outline, мусор от модели, синтетические шаблоны — без исключений и потерь.

Полная матрица — `tools/stress_test.py`; здесь быстрые регрессии на VK Tech и синтетике.
"""

from __future__ import annotations

import copy
import math
import random
from pathlib import Path

import pytest
from pydantic import ValidationError

from deckforge.content import repair_outline
from deckforge.core.ir import Archetype, ChartSpec, DeckOutline, ImageSpec, OutlineSlide, TableSpec
from deckforge.core.strategy import load_strategy
from deckforge.layout import build_deck_ir
from deckforge.layout.builder import MAX_RETRIES
from deckforge.layout.fitting import normalize, xml_safe
from deckforge.layout.planner import sanitize_data, split_mixed_data
from deckforge.render import render_pptx
from tests.conftest import cassette
from tests.fixtures import stress_outlines as so
from tests.fixtures.stress_templates import TEMPLATES

REPO = Path(__file__).resolve().parents[1]
STRATEGIES = ("executive", "narrative", "visual")


@pytest.fixture(scope="module")
def vk_tech():
    from deckforge.pipeline import parse_template
    from tests.conftest import find_template

    p = find_template("VK Tech")
    if p is None:
        pytest.skip("VK Tech не найден (LFS?)")
    return parse_template(p)


def _build(parsed, outline: DeckOutline, strategy: str = "executive"):
    return build_deck_ir(outline, load_strategy(strategy), parsed.exemplars, parsed.tokens.template_id,
                         parsed.tokens.slide_w, parsed.tokens.slide_h, parsed.style)


def _per_ref_max(ir) -> int:
    counts: dict[int, int] = {}
    for s in ir.slides:
        counts[s.outline_ref] = counts.get(s.outline_ref, 0) + 1
    return max(counts.values(), default=0)


# ──────────────────────────── текст ────────────────────────────


def test_normalize_strips_xml_invalid_chars() -> None:
    assert normalize("до\x00после \x0bвер\x1fт \x08bs \ud800 \ufffe ок\tтаб\nстрока") == "допосле верт bs ок таб строка"
    assert xml_safe("a\x00b\nc\td") == "ab\nc\td"
    assert normalize(so.EMOJI) == so.EMOJI and normalize(so.RTL) == so.RTL


def test_control_chars_survive_layout_and_render(vk_tech, tmp_path: Path) -> None:
    res = _build(vk_tech, so.text_extremes())
    out = render_pptx(res.ir, vk_tech.template, vk_tech.exemplars, tmp_path / "text.pptx")
    assert out.exists() and out.stat().st_size > 100_000
    texts = [r.text for s in res.ir.slides for el in s.elements for p in el.paragraphs for r in p.runs]
    assert not any("\x00" in t or "\x0b" in t for t in texts)
    assert not any("\x00" in s.notes for s in res.ir.slides)
    assert not [c for c in res.choices if c.exemplar_id is None]


# ──────────────────────────── данные ────────────────────────────


def test_sanitize_data_aligns_series_and_drops_empty() -> None:
    w: list[str] = []
    short = OutlineSlide(idx=1, archetype=Archetype.CHART, title="т",
                         chart=ChartSpec(kind="column", title="", categories=["а", "б", "в"], series={"x": [1.0], "y": [1.0, 2.0, 3.0]}))
    s = sanitize_data(short, w)
    assert s.chart is not None and s.chart.categories == ["а"] and s.chart.series == {"x": [1.0], "y": [1.0]}
    nan = OutlineSlide(idx=2, archetype=Archetype.CHART, title="т",
                       chart=ChartSpec(kind="line", title="", categories=["а", "б"], series={"x": [1.0, math.nan]}))
    s = sanitize_data(nan, w)
    assert s.chart is None and s.archetype == Archetype.BULLETS
    empty = OutlineSlide(idx=3, archetype=Archetype.CHART, title="т", chart=ChartSpec(kind="bar", title="", categories=[], series={}))
    assert sanitize_data(empty, w).chart is None
    ragged = OutlineSlide(idx=4, archetype=Archetype.TABLE, title="т",
                          table=TableSpec(header=["а", "б", "в"], rows=[["1"], [], ["1", "2", "3", "4"]]))
    s = sanitize_data(ragged, w)
    assert s.table is not None and s.table.rows == [["1", "", ""], ["", "", ""], ["1", "2", "3"]]
    nocols = OutlineSlide(idx=5, archetype=Archetype.TABLE, title="т", table=TableSpec(header=[], rows=[[]]))
    s = sanitize_data(nocols, w)
    assert s.table is None and s.archetype == Archetype.BULLETS
    assert len(w) == 5
    clean = OutlineSlide(idx=6, archetype=Archetype.CHART, title="т", chart=so._chart(3))
    assert sanitize_data(clean, w) is clean and len(w) == 5


def test_plan_drops_empty_content_slides() -> None:
    """Готовый outline идёт мимо repair_outline — пустой контентный слайд снимает планировщик (иначе I03)."""
    from deckforge.layout.planner import plan

    slides = [
        OutlineSlide(idx=0, archetype=Archetype.TITLE, title="Т"),
        OutlineSlide(idx=1, archetype=Archetype.BULLETS, title="Пустой"),
        OutlineSlide(idx=2, archetype=Archetype.CARDS, title="С тезисами", bullets=["а", "б"]),
        OutlineSlide(idx=3, archetype=Archetype.CHART, title="Битая диаграмма",
                     chart=ChartSpec(kind="bar", title="", categories=[], series={})),
        OutlineSlide(idx=4, archetype=Archetype.CLOSING, title="Т"),
    ]
    res = plan(DeckOutline(title="т", purpose="other", slides=slides), load_strategy("executive"))
    titles = [s.title for s in res.slides]
    assert "Пустой" not in titles and "Битая диаграмма" not in titles  # chart без данных → bullets без тела
    assert titles[0] == "Т" and "С тезисами" in titles and titles[-1] == "Т"
    assert sum("без содержимого — пропущен" in w for w in res.warnings) == 2

    # если пустыми оказались все контентные слайды — колоду не выбрасываем
    only_empty = [slides[0], OutlineSlide(idx=1, archetype=Archetype.BULLETS, title="Пустой")]
    res2 = plan(DeckOutline(title="т", purpose="other", slides=only_empty), load_strategy("executive"))
    assert [s.title for s in res2.slides] == ["Т", "Пустой"]


def test_kpis_next_to_chart_go_to_own_slide() -> None:
    mixed = OutlineSlide(idx=1, archetype=Archetype.KPI, title="т", kpis=so._kpis(2), bullets=["а"], chart=so._chart(3))
    out = split_mixed_data([mixed])
    assert [s.archetype for s in out] == [Archetype.KPI, Archetype.KPI]
    assert out[0].kpis == [] and out[0].chart is not None and out[1].kpis == mixed.kpis and out[1].chart is None
    plain = OutlineSlide(idx=2, archetype=Archetype.KPI, title="т", kpis=so._kpis(2))
    assert split_mixed_data([plain]) == [plain]


@pytest.mark.parametrize("strategy", STRATEGIES)
def test_numbers_extremes_build_and_render(vk_tech, tmp_path: Path, strategy: str) -> None:
    res = _build(vk_tech, so.numbers_extremes(), strategy)
    assert not [c for c in res.choices if c.exemplar_id is None]
    render_pptx(res.ir, vk_tech.template, vk_tech.exemplars, tmp_path / f"{strategy}.pptx")


def test_step_badges_get_ordinals_not_leads() -> None:
    """Кружки «1 2 3» у шагов процесса: в label только порядковый номер, пункт целиком в тело."""
    from deckforge.layout.builder import is_badge
    from deckforge.pipeline import parse_template

    p = REPO / "data" / "wild" / "presentation_eng_dark.pptx"
    if not p.exists() or p.stat().st_size < 10_000:
        pytest.skip("нет presentation_eng_dark.pptx")
    parsed = parse_template(p)
    badges = {(e.id, s.id): s for e in parsed.exemplars for s in e.slots if is_badge(s)}
    assert badges, "в шаблоне ожидались label-слоты с номером-образцом"
    full = so.all_process()
    short = full.model_copy(update={"slides": full.slides[:4] + full.slides[-1:]})  # в объёме
    res = _build(parsed, short, "narrative")
    filled = 0
    for sl in res.ir.slides:
        for el in sl.elements:
            if (sl.exemplar_id, el.slot_id) in badges:
                text = "".join(r.text for par in el.paragraphs for r in par.runs)
                assert text.isdigit() and len(text) <= 2, (sl.exemplar_id, el.slot_id, text)
                filled += 1
    assert filled > 0


def test_numbers_extremes_on_vk_education() -> None:
    """Пустое значение KPI не делит на ноль в fit_number."""
    from deckforge.pipeline import parse_template
    from tests.conftest import find_template

    p = find_template("VK Education")
    if p is None:
        pytest.skip("VK Education не найден (LFS?)")
    parsed = parse_template(p)
    for strategy in STRATEGIES:
        res = _build(parsed, so.numbers_extremes(), strategy)
        assert not [c for c in res.choices if c.exemplar_id is None]


# ──────────────────────────── картинки ────────────────────────────


def test_missing_or_broken_image_does_not_break_render(vk_tech, tmp_path: Path) -> None:
    (tmp_path / "fake.png").write_text("не картинка", "utf-8")
    slides = [
        OutlineSlide(idx=0, archetype=Archetype.TITLE, title="т"),
        OutlineSlide(idx=1, archetype=Archetype.IMAGE_TEXT, title="нет файла", bullets=["тезис"],
                     image=ImageSpec(path=str(tmp_path / "missing.png"))),
        OutlineSlide(idx=2, archetype=Archetype.IMAGE_TEXT, title="битый файл", bullets=["тезис"],
                     image=ImageSpec(path=str(tmp_path / "fake.png"))),
        OutlineSlide(idx=3, archetype=Archetype.CLOSING, title="конец"),
    ]
    outline = DeckOutline(title="т", purpose="other", slides=slides)
    res = _build(vk_tech, outline, "visual")
    assert not any(el.image_path and el.image_path.endswith("missing.png") for s in res.ir.slides for el in s.elements)
    out = render_pptx(res.ir, vk_tech.template, vk_tech.exemplars, tmp_path / "img.pptx")
    assert out.exists()


# ──────────────────────────── образцы и остатки ────────────────────────────


def test_freeform_in_ready_outline_is_not_skipped(vk_tech) -> None:
    res = _build(vk_tech, so.rare_archetypes())
    assert not [c for c in res.choices if c.exemplar_id is None]


def test_nothing_placed_retries_other_exemplar(vk_tech) -> None:
    """Образец без текстовых слотов не получает слайд с тезисом: builder перебирает образцы, пока контент не ляжет."""
    res = _build(vk_tech, so.all_image_text_no_path(), "executive")
    for s in res.ir.slides[1:-1]:
        assert any(el.paragraphs for el in s.elements if el.kind.value in ("body", "label", "caption")), s.exemplar_id
    assert _per_ref_max(res.ir) <= MAX_RETRIES + 1


def test_structural_leftover_becomes_text_or_kpi_slide(vk_tech) -> None:
    """Буллеты/KPI на титульном слайде не капают по одному в подписи титулов."""
    slides = [OutlineSlide(idx=0, archetype=Archetype.TITLE, title="т", bullets=[f"п{i}" for i in range(6)],
                           kpis=so._kpis(4)),
              OutlineSlide(idx=1, archetype=Archetype.CLOSING, title="к")]
    res = _build(vk_tech, DeckOutline(title="т", purpose="other", slides=slides))
    assert not [c for c in res.choices if c.exemplar_id is None]
    assert _per_ref_max(res.ir) <= 3
    assert {s.archetype for s in res.ir.slides} & {Archetype.KPI, Archetype.CARDS, Archetype.BULLETS}


@pytest.mark.parametrize("case", sorted(so.CASES))
def test_every_stress_outline_builds_without_skips(vk_tech, case: str) -> None:
    for strategy in STRATEGIES:
        res = _build(vk_tech, so.CASES[case](), strategy)
        assert not [c for c in res.choices if c.exemplar_id is None], (case, strategy)
        assert _per_ref_max(res.ir) <= 4, (case, strategy)
        if case != "empty":
            assert res.ir.slides


_BLOATED = [  # 20 × 8 KPI / 12 шагов на шаблонах, где ни один образец столько не вмещает
    ("02_HSE", "all_kpi"), ("presentation_eng_dark", "all_kpi"), ("presentation_eng_dark", "all_process"),
    ("Теорема Пифагора", "all_process"), ("VK Education", "all_process"), ("VK Tech", "all_process"),
]


@pytest.mark.parametrize("name,case", _BLOATED)
def test_over_budget_kpi_process_do_not_bloat_on_real_templates(name: str, case: str) -> None:
    """Критерии те же, что в `tools/stress_test.py`: контент сворачивается в карточки/список, а не в каскад."""
    from deckforge.pipeline import parse_template
    from tests.conftest import find_template
    from tools.stress_test import bloat_problems

    p = find_template(name)
    if p is None:
        pytest.skip(f"{name} не найден (LFS?)")
    parsed = parse_template(p)
    outline = so.CASES[case]()
    for strategy in STRATEGIES:
        res = _build(parsed, outline, strategy)
        assert bloat_problems(res.ir.slides, len(outline.slides), load_strategy(strategy), case) == [], (name, case, strategy)
        assert not [c for c in res.choices if c.exemplar_id is None]
        texts = " ".join(r.text for s in res.ir.slides for el in s.elements for p_ in el.paragraphs for r in p_.runs)
        for s in outline.slides:
            for k in s.kpis:
                assert k.label in texts, (name, case, strategy, k.label)
            for step in s.steps:
                assert step.split(":")[0] in texts, (name, case, strategy, step)


# ──────────────────────────── фазз repair_outline ────────────────────────────

_GARBAGE = [None, "", "строка", 0, 1, -1, 3.5, True, [], {}, ["a", "b"], {"x": 1}, [[1, 2], [3]], [{"a": 1}], "1 250",
            "NaN", "x" * 5000, [None, None], {"name": "s", "values": ["1", "2", "3", "4", "5", "6"]}, "\x00\x0b",
            list(range(100))]
_AVAIL = {Archetype.TITLE, Archetype.CARDS, Archetype.KPI, Archetype.CHART, Archetype.TABLE, Archetype.PROCESS,
          Archetype.IMAGE_TEXT, Archetype.BULLETS, Archetype.CLOSING, Archetype.SECTION}


def _mutate(rng: random.Random, raw: dict) -> dict:
    d = copy.deepcopy(raw)
    for _ in range(rng.randint(1, 4)):
        kind = rng.choice(["top", "slide", "deep", "grow"])
        slides = d.get("slides")
        if kind == "top":
            d[rng.choice(["title", "purpose", "audience", "language", "slides", "extra"])] = rng.choice(_GARBAGE)
        elif not isinstance(slides, list) or not slides:
            continue
        else:
            s = slides[rng.randrange(len(slides))]
            if not isinstance(s, dict):
                continue
            if kind == "slide":
                key = rng.choice(list(s) + ["archetype", "bullets", "kpis", "chart", "table", "steps", "image", "quote",
                                            "cards", "sources", "idx", "section", "speaker_notes"])
                if rng.random() < 0.3:
                    s.pop(key, None)
                else:
                    s[key] = rng.choice(_GARBAGE)
            elif kind == "deep":
                for field in ("chart", "table"):
                    if isinstance(s.get(field), dict) and s[field]:
                        s[field][rng.choice(list(s[field]))] = rng.choice(_GARBAGE)
                for field in ("kpis", "bullets"):
                    if isinstance(s.get(field), list) and s[field]:
                        s[field][rng.randrange(len(s[field]))] = rng.choice(_GARBAGE)
            else:
                s["bullets"] = [f"пункт {k}" for k in range(rng.choice([0, 1, 7, 30, 100]))]
                s["kpis"] = [{"value": f"{k}%", "label": f"м{k}"} for k in range(rng.choice([0, 5, 50]))]
    return d


def test_repair_outline_fuzz_returns_outline_or_validation_error(vk_tech) -> None:
    """Любая мутация ответа модели → DeckOutline или ValidationError, никаких TypeError; outline собирается."""
    raw0 = cassette("outline_writer_pulse")
    rng = random.Random(42)
    ok = 0
    for _ in range(150):
        raw = _mutate(rng, raw0)
        try:
            outline, _ = repair_outline(raw, _AVAIL, {"brief"}, purpose="product")
        except ValidationError:
            continue
        ok += 1
        for strategy in STRATEGIES:
            res = _build(vk_tech, outline, strategy)
            assert not [c for c in res.choices if c.exemplar_id is None]
            assert len(res.ir.slides) <= 60 and _per_ref_max(res.ir) <= 4
    assert ok >= 80


def test_repair_outline_scalars_instead_of_lists() -> None:
    for raw in ({"slides": 3.5}, {"slides": 1}, {"slides": {"idx": 0}}, {"slides": "x"}):
        with pytest.raises(ValidationError):
            repair_outline(raw, _AVAIL, set())
    raw = {"title": "т", "slides": [
        {"idx": 0, "archetype": "table", "title": "таблица", "table": {"header": 1, "rows": 3.5}},
        {"idx": 1, "archetype": "chart", "title": "график", "chart": {"categories": "а", "series": {"x": 5}}},
        {"idx": 2, "archetype": "table", "title": "строки", "table": {"header": ["а", "б"], "rows": [None, "x", ["1", "2"]]}},
    ]}
    outline, warnings = repair_outline(raw, _AVAIL, set())
    kinds = [s.archetype for s in outline.slides]
    assert kinds == [Archetype.TITLE, Archetype.TABLE, Archetype.CHART, Archetype.TABLE, Archetype.CLOSING]
    assert outline.slides[1].table.rows == [["3.5"]] and outline.slides[2].chart.series == {"x": [5.0]}
    assert outline.slides[3].table.rows == [["", ""], ["x", ""], ["1", "2"]]


# ──────────────────────────── синтетические шаблоны ────────────────────────────


@pytest.mark.parametrize("name", sorted(TEMPLATES))
def test_synthetic_templates_survive_pipeline(tmp_path: Path, name: str, monkeypatch: pytest.MonkeyPatch) -> None:
    import importlib

    from deckforge.audit import audit_deck
    from deckforge.export.html import export_html
    from deckforge.pipeline import parse_template

    monkeypatch.setattr(importlib.import_module("deckforge.pipeline.run"), "PREPARED_DIR", tmp_path / "prepared")
    pptx = TEMPLATES[name](tmp_path)
    parsed = parse_template(pptx)
    outline = so.base()
    for strategy in STRATEGIES:
        res = _build(parsed, outline, strategy)
        # parsed.template, а не pptx: шаблон без слайдов разбирается из копии с образцами из лейаутов
        out = render_pptx(res.ir, parsed.template, parsed.exemplars, tmp_path / f"{name}_{strategy}.pptx")
        report = audit_deck(out, parsed.dna, res.ir)
        export_html(out, out.with_suffix(".html"), ir=res.ir)
        skipped = [c for c in res.choices if c.exemplar_id is None]
        if name in ("four_by_three", "a4_portrait", "placeholders_only"):
            # колода собирается целиком
            assert not skipped and len(res.ir.slides) >= 8, (name, strategy)
            assert report.errors == 0 or name == "placeholders_only"
        if name == "nested_groups":  # растянутые картинки самого шаблона — предупреждение
            assert not [f for f in report.findings if f.check_id == "L07_picture_stretched" and f.severity == "error"]


def _run_executive(pptx: Path, tmp_path: Path):
    from deckforge.pipeline import RunConfig, run

    cfg = RunConfig(template=pptx, content_pack=REPO / "examples" / "content_pack", strategies=["executive"],
                    output_dir=tmp_path / "run", images="off",
                    audit={"deterministic": True, "contextual": False, "autofix": True})
    return run(cfg, outline=so.base())


def test_template_without_slides_uses_layouts(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """Презентация без слайдов (как .potx): образцы — из лейаутов с плейсхолдерами, колода собирается."""
    import importlib

    monkeypatch.setattr(importlib.import_module("deckforge.pipeline.run"), "PREPARED_DIR", tmp_path / "prepared")
    res = _run_executive(TEMPLATES["empty_presentation"](tmp_path), tmp_path)
    assert res.decks[0].stats["slides"] >= 8 and res.decks[0].stats["skipped"] == 0
    assert any("нет слайдов" in w for w in res.warnings)


def test_template_without_exemplars_gives_empty_deck_not_crash(tmp_path: Path,
                                                               monkeypatch: pytest.MonkeyPatch) -> None:
    """Ни слайдов, ни плейсхолдеров в лейаутах — образцов нет: пустая колода с предупреждениями, без падения."""
    import importlib

    from pptx import Presentation

    monkeypatch.setattr(importlib.import_module("deckforge.pipeline.run"), "PREPARED_DIR", tmp_path / "prepared")
    prs = Presentation(str(TEMPLATES["empty_presentation"](tmp_path)))
    for layout in prs.slide_layouts:
        for ph in list(layout.placeholders):
            ph.element.getparent().remove(ph.element)
    prs.save(str(pptx := tmp_path / "bare.pptx"))
    res = _run_executive(pptx, tmp_path)
    assert res.decks[0].stats["slides"] == 0 and res.decks[0].stats["skipped"] == len(res.decks[0].choices)
    assert any("пропущен" in w for w in res.warnings)
