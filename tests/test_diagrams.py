"""Шаги процесса настоящим SmartArt «Простой процесс»: планировщик по process_form, части SmartArt, геометрия
по правилам макета (сверено с PowerPoint), чтение отрисовки для аудита и HTML, чистый аудит."""

from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu

from deckforge.audit import audit_deck
from deckforge.core.ir import Archetype, Box, DeckOutline, DiagramSpec, OutlineSlide, SlideIR
from deckforge.core.strategy import load_strategy
from deckforge.layout import layout_deck
from deckforge.layout.planner import apply_process_form
from deckforge.parsing.dna import build_dna
from deckforge.render import render_pptx
from deckforge.render.diagrams import caption_size, on_fill
from deckforge.render.smartart import LAYOUT_URN, add_smartart, layout_steps

STEPS = ["Аудит текущей инфраструктуры", "Проектирование архитектуры", "Пилотное внедрение",
         "Масштабирование и поддержка"]
FIXTURE = Path(__file__).parent / "fixtures" / "smartart_process1.json"
DGM = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
RELTYPES = ("diagramData", "diagramLayout", "diagramQuickStyle", "diagramColors", "diagramDrawing")


def _process(n: int = 4) -> OutlineSlide:
    return OutlineSlide(idx=1, archetype=Archetype.PROCESS, title="Этапы внедрения", steps=STEPS[:n])


def test_process_form_modes() -> None:
    visual, narrative, executive = (load_strategy(n) for n in ("visual", "narrative", "executive"))
    with_process = set(Archetype)
    without_process = set(Archetype) - {Archetype.PROCESS}
    # diagram — всегда; template — только без process-образца; list — никогда
    assert apply_process_form(_process(), visual, with_process).diagram is not None
    assert apply_process_form(_process(), narrative, with_process).diagram is None
    s = apply_process_form(_process(), narrative, without_process)
    assert s.diagram == DiagramSpec(kind="process", items=STEPS) and not s.steps
    assert apply_process_form(_process(), executive, without_process).diagram is None
    # один шаг или слишком много — остаётся списком
    assert apply_process_form(_process(1), visual, with_process).diagram is None
    many = _process()
    many.steps = [f"Шаг {i}" for i in range(9)]
    assert apply_process_form(many, visual, with_process).diagram is None


def test_caption_size_and_text_color() -> None:
    # длинное слово не рвётся: кегль уменьшается, пока слово не влезет в строку
    big = caption_size(["Масштабирование"], w=914400 * 2, h=914400)
    small = caption_size(["Масштабирование"], w=914400, h=914400)
    assert small < big
    assert on_fill("520977") == "FFFFFF" and on_fill("D6E8FF") == "212121"
    assert on_fill("D6E8FF", dark="1A1A1A") == "1A1A1A"  # тёмный текст — цвет палитры шаблона


def test_layout_steps_match_powerpoint() -> None:
    """Геометрия отрисовки — правила layoutDef process1; с тем, что строит сам PowerPoint (tools/make_smartart_assets),
    расходится меньше чем на 0,5 % ширины бокса — при правке в PowerPoint схема не «прыгнет»."""
    cases = json.loads(FIXTURE.read_text("utf-8"))
    assert len(cases) == 10
    for name, case in cases.items():
        w, h = case["box"]
        ours = layout_steps(case["n"], w, h)
        assert [g for g, _ in ours] == [s["geom"] for s in case["shapes"]], name
        for (_, b), s in zip(ours, case["shapes"]):
            for mine, theirs, size in zip((b.x, b.y, b.w, b.h), s["box"], (w, h, w, h)):
                assert abs(mine - theirs) <= 0.005 * w, (name, (b.x, b.y, b.w, b.h), s["box"])


def _smartart_on_blank(out: Path, items: list[str], overrides: dict) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = Box(x=Emu(457200), y=Emu(1600200), w=Emu(8229600), h=Emu(3000000))
    add_smartart(slide, DiagramSpec(kind="process", items=items), box, overrides)
    prs.save(str(out))
    return out


def _parts(pptx: Path) -> dict[str, str]:
    with zipfile.ZipFile(pptx) as z:
        return {n: z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("ppt/diagrams/")
                or re.fullmatch(r"ppt/slides/_rels/slide\d+\.xml\.rels", n) or n == "[Content_Types].xml"}


def test_smartart_parts_colors_and_determinism(tmp_path: Path) -> None:
    overrides = {"palette": "0077FF,00AEE8", "font": "Play", "text_color": "1A1A1A"}
    parts = _parts(_smartart_on_blank(tmp_path / "a.pptx", STEPS[:3], overrides))
    data = next(v for k, v in parts.items() if re.fullmatch(r"ppt/diagrams/data\d+\.xml", k))
    drawing = next(v for k, v in parts.items() if re.fullmatch(r"ppt/diagrams/drawing\d+\.xml", k))
    rels = next(v for k, v in parts.items() if k.startswith("ppt/slides/_rels/"))
    # пять частей со связями от слайда и своими content types
    assert all(f"/{t}" in rels for t in RELTYPES)
    assert parts["[Content_Types].xml"].count("drawingml.diagram") == 5
    # данные: макет process1, узлы с текстами шагов по порядку, цвет и шрифт шаблона — пользовательское оформление
    assert f'loTypeId="{LAYOUT_URN}"' in data
    assert re.findall(r"<a:t>([^<]+)</a:t>", data) == STEPS[:3]
    assert data.count('custT="1"') == 3 and 'typeface="Play"' in data and '<a:srgbClr val="0077FF"/>' in data
    # стрелка — заливка точки sibTrans (так её хранит PowerPoint), у последнего шага стрелки нет
    assert len(re.findall(r'type="sibTrans"[^>]*><dgm:prSet/><dgm:spPr><a:solidFill>', data)) == 2
    # отрисовка: узлы и стрелки, тексты те же, текст контрастный к заливке, ссылки на стиль (без них PowerPoint
    # собирает фигуры без обводки и заливки)
    assert re.findall(r'prst="(\w+)"', drawing) == ["roundRect", "rightArrow", "roundRect", "rightArrow", "roundRect"]
    assert re.findall(r"<a:t>([^<]+)</a:t>", drawing) == STEPS[:3]
    assert drawing.count("<dsp:style>") == 5
    first_text = re.search(r'<a:rPr[^>]*><a:solidFill><a:srgbClr val="(\w+)"', drawing).group(1)
    assert first_text == on_fill("0077FF", "1A1A1A")
    # data-часть называет отрисовку id связи слайда
    rid = re.search(r'dataModelExt[^>]*relId="(rId\d+)"', data).group(1)
    assert re.search(rf'Id="{rid}"[^>]*diagramDrawing', rels)
    # один вход — один файл (modelId — детерминированные GUID)
    again = _parts(_smartart_on_blank(tmp_path / "b.pptx", STEPS[:3], overrides))
    assert next(v for k, v in again.items() if re.fullmatch(r"ppt/diagrams/data\d+\.xml", k)) == data


def test_smartart_rendered_read_back_and_audit_clean(template_path, tmp_path: Path) -> None:
    from deckforge.core.deck_reader import read_shapes
    from deckforge.core.package import Package, PartCtx
    from deckforge.export.html import export_html

    tpl = template_path("VK Tech")
    dna = build_dna(tpl)
    outline = DeckOutline(title="Тест", purpose="product", slides=[
        OutlineSlide(idx=0, archetype=Archetype.TITLE, title="Тест схемы"),
        _process(),
        OutlineSlide(idx=2, archetype=Archetype.CLOSING, title="Спасибо"),
    ])
    res = layout_deck(outline, dna, load_strategy("visual"))
    diagram_slides = [s for s in res.ir.slides if any(e.diagram for e in s.elements)]
    assert len(diagram_slides) == 1
    idx = diagram_slides[0].idx
    out = render_pptx(res.ir, tpl, dna.exemplars, tmp_path / "diagram.pptx")

    # на слайде — одна рамка SmartArt, картинок и групп-заменителей нет
    prs = Presentation(str(out))
    frames = [sh for sh in prs.slides[idx].shapes if sh._element.find(f".//{{{DGM}}}relIds") is not None]
    assert len(frames) == 1 and frames[0].name == "Шаги процесса"

    # читатель колоды видит отрисовку: узлы, стрелки и текст шагов
    pkg = Package(out)
    shapes = read_shapes(PartCtx.for_slide(pkg, pkg.slides[idx]))
    [frame] = [s for s in shapes if s.children]
    assert [c.geom for c in frame.children].count("roundRect") == len(STEPS)
    assert frame.text.split("\n") == STEPS and frame.is_content

    report = audit_deck(out, dna, res.ir)
    on_slide = [f for f in report.findings if f.slide_idx == idx]
    assert not [f for f in on_slide if f.severity == "error"], on_slide
    assert not [f for f in on_slide if f.check_id[:3] in ("L05", "L06", "L02", "L03", "T01", "T02", "T03", "I03")], on_slide

    html = export_html(out, tmp_path / "diagram.html").read_text("utf-8")
    assert all(t in html for t in STEPS) and html.count("polygon(") >= len(STEPS) - 1


def test_template_smartart_keeps_its_drawing(template_path, tmp_path: Path) -> None:
    """Образец со своим SmartArt (МТУСИ slide8, пять схем): при клонировании переносятся все пять частей каждой
    схемы, включая отрисовку — её называет data-часть, а не XML слайда (иначе LibreOffice рисовал бы пустоту)."""
    from deckforge.parsing.layout_classifier import classify_template

    tpl = template_path("МТУСИ")
    exemplars = [p.to_exemplar() for p in classify_template(tpl)]
    e = next((x for x in exemplars if x.id == "slide8"), None)
    if e is None:
        pytest.skip("в шаблоне нет slide8")
    from tests.test_pptx_writer import _deck

    out = render_pptx(_deck(exemplars, [SlideIR(idx=0, exemplar_id=e.id, archetype=e.archetype, elements=[],
                                                outline_ref=0)]), tpl, exemplars, tmp_path / "mtusi8.pptx")
    parts = _parts(out)
    rels = next(v for k, v in parts.items() if k.startswith("ppt/slides/_rels/"))
    datas = [v for k, v in parts.items() if re.fullmatch(r"ppt/diagrams/data\d+\.xml", k)]
    assert len(datas) == 5 and all(rels.count(f"/{t}\"") == 5 for t in RELTYPES)
    for data in datas:
        rid = re.search(r'dataModelExt[^>]*relId="(rId\d+)"', data).group(1)
        assert re.search(rf'Id="{rid}"[^>]*diagramDrawing', rels), rid
