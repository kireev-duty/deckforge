"""Синтетические колоды-нарушения для каждой детерминированной проверки (docs/AUDIT.md) + «чистая» колода.

Генерируются python-pptx на лету (в tmp_path) — бинарники в git не храним. Режим без DeckIR: «чужая»
колода, шаблон для T04/T05 — сама чистая колода (те же лейауты). `tools/make_fixtures.py` выгружает
их в папку, чтобы посмотреть глазами через render-deck.

Использование:
    case = FIXTURES["L03"](tmp_path)   # Case(pptx, dna, ir, check, slide_idx)
"""

from __future__ import annotations

import re
import zipfile
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt

from deckforge.core.ir import Box, ColorToken, DeckIR, FixedElement, GridSpec, TemplateDNA, TypeToken
from deckforge.core.ooxml import A, NS

SW, SH = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.5)
COL2 = Inches(7)
LOGO_BOX = Box(x=Inches(12), y=Inches(6.8), w=Inches(1), h=Inches(0.5))
TEXT, ACCENT, MUTED, BG = "212121", "0077FF", "8F8F8F", "FFFFFF"
FONT = "Arial"


@dataclass
class Case:
    pptx: Path
    dna: TemplateDNA
    ir: DeckIR | None
    check: str  # префикс проверки, напр. "L03"
    slide_idx: int  # где ожидается находка (-1 — файл)


# ──────────────────────────── DNA ────────────────────────────


def fixture_dna(source_path: Path | str = "") -> TemplateDNA:
    return TemplateDNA(
        template_id="fixture", source_path=str(source_path), slide_w=SW, slide_h=SH,
        colors=[ColorToken(hex=BG, role="background"), ColorToken(hex=TEXT, role="text"),
                ColorToken(hex=ACCENT, role="accent"), ColorToken(hex=MUTED, role="muted")],
        typography=[TypeToken(role="h1", font=FONT, size_pt=32, bold=True), TypeToken(role="h2", font=FONT, size_pt=24),
                    TypeToken(role="body", font=FONT, size_pt=18), TypeToken(role="caption", font=FONT, size_pt=12)],
        fonts=[FONT],
        grid=GridSpec(slide_w=SW, slide_h=SH, margin_left=MARGIN, margin_right=MARGIN, margin_top=MARGIN,
                      margin_bottom=MARGIN, columns_x=[MARGIN, COL2], rows_y=[MARGIN, Inches(1.5)]),
        fixed_elements=[FixedElement(kind="logo", box=LOGO_BOX, occurrences=10, is_picture=True)],
        exemplars=[],
    )


# ──────────────────────────── строитель ────────────────────────────


class DeckBuilder:
    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = SW, SH
        self._imgs: dict[tuple[int, int], Path] = {}
        self.tmp: Path | None = None

    def slide(self, layout: int = 6):
        return self.prs.slides.add_slide(self.prs.slide_layouts[layout])

    def text(self, slide, x, y, w, h, paragraphs, size=18, font=FONT, color=TEXT, bold=False, bullets=False,
             fill: str | None = None, wrap=True):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = wrap
        tf.auto_size = MSO_AUTO_SIZE.NONE  # python-pptx по умолчанию ставит spAutoFit — фикстуры должны быть «жёсткими»
        if isinstance(paragraphs, str):
            paragraphs = [paragraphs]
        for i, text in enumerate(paragraphs):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            run = p.add_run()
            run.text = text
            run.font.size, run.font.name, run.font.bold = Pt(size), font, bold
            run.font.color.rgb = RGBColor.from_string(color)
            if bullets:
                ppr = p._p.get_or_add_pPr()
                ppr.set("marL", "285750")
                ppr.set("indent", "-285750")
                bu = etree.SubElement(ppr, A + "buChar")
                bu.set("char", "•")
        if fill:
            tb.fill.solid()
            tb.fill.fore_color.rgb = RGBColor.from_string(fill)
        return tb

    def title(self, slide, text, y=MARGIN, size=32):
        return self.text(slide, MARGIN, y, SW - 2 * MARGIN, Inches(1.3), text, size=size, bold=True)

    def image(self, px_w: int, px_h: int) -> Path:
        assert self.tmp is not None
        key = (px_w, px_h)
        if key not in self._imgs:
            p = self.tmp / f"img_{px_w}x{px_h}.png"
            Image.new("RGB", (px_w, px_h), "#0077FF").save(p)
            self._imgs[key] = p
        return self._imgs[key]

    def picture(self, slide, x, y, w, h, px=(400, 300)):
        return slide.shapes.add_picture(str(self.image(*px)), x, y, w, h)

    def chart(self, slide, x, y, w, h, n_series=2, legend=True, axis_title=True, categories=3):
        data = CategoryChartData()
        data.categories = [f"к{i}" for i in range(categories)]
        for s in range(n_series):
            data.add_series(f"серия {s + 1}", [s + i + 1 for i in range(categories)])
        frame = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, w, h, data)
        chart = frame.chart
        chart.has_legend = legend
        if axis_title:
            chart.value_axis.has_title = True
            chart.value_axis.axis_title.text_frame.text = "млн ₽"
        return frame

    def table(self, slide, x, y, w, h, rows=3, cols=3):
        frame = slide.shapes.add_table(rows, cols, x, y, w, h)
        for r in range(rows):
            for c in range(cols):
                cell = frame.table.cell(r, c)
                cell.text = f"r{r}c{c}"
                for p in cell.text_frame.paragraphs:
                    for run in p.runs:
                        run.font.name, run.font.size = FONT, Pt(12)
                        run.font.color.rgb = RGBColor.from_string(TEXT)
        return frame

    def logo(self, slide, x=LOGO_BOX.x, y=LOGO_BOX.y):
        return self.picture(slide, x, y, LOGO_BOX.w, LOGO_BOX.h, px=(200, 100))

    def save(self, path: Path) -> Path:
        self.prs.save(str(path))
        return path


def _builder(tmp: Path) -> DeckBuilder:
    b = DeckBuilder()
    b.tmp = tmp
    return b


def _one_slide_deck(tmp: Path, name: str, fill: Callable[[DeckBuilder, object], None]) -> tuple[Path, TemplateDNA]:
    """Колода из одного слайда с заголовком; шаблон для T04/T05 — чистая колода."""
    clean = make_clean(tmp).pptx
    b = _builder(tmp)
    s = b.slide()
    b.title(s, "Заголовок слайда с выводом")
    fill(b, s)
    return b.save(tmp / f"{name}.pptx"), fixture_dna(clean)


# ──────────────────────────── чистая колода ────────────────────────────


def make_clean(tmp: Path) -> Case:
    path = tmp / "clean.pptx"
    if path.exists():
        return Case(path, fixture_dna(path), None, "", -1)
    b = _builder(tmp)
    s1 = b.slide()  # титул: два блока → «разреженный» без IR
    b.text(s1, MARGIN, Inches(2.5), Inches(9), Inches(1.3), "Пульс команды: перегрузка видна за неделю", size=32, bold=True)
    # 24 pt — «крупный» текст по WCAG: muted #8F8F8F на белом (3.2:1) допустим при пороге 3.0
    b.text(s1, MARGIN, Inches(3.9), Inches(9), Inches(0.6), "Итоги пилота и план запуска", size=24, color=MUTED)
    b.logo(s1)
    s2 = b.slide()
    b.title(s2, "Команды теряют треть времени на согласования")
    b.text(s2, MARGIN, Inches(2), Inches(6), Inches(4), [
        "Статусы вместо работы — до 9,5 часов в неделю", "Перегрузка скрыта — трекер показывает задачи",
        "Поздние сигналы — о выгорании узнают на ретро", "Ручные отчёты — данные устаревают за день",
    ], bullets=True)
    b.picture(s2, COL2, Inches(2), Inches(5), Inches(3.75))
    b.logo(s2)
    s3 = b.slide()
    b.title(s3, "Время на согласования упало вдвое за месяц")
    b.chart(s3, MARGIN, Inches(2), Inches(8), Inches(4.5))
    b.logo(s3)
    s4 = b.slide()
    b.title(s4, "До и после пилота по трём метрикам")
    b.table(s4, MARGIN, Inches(2), Inches(8), Inches(3))
    b.logo(s4)
    b.save(path)
    return Case(path, fixture_dna(path), None, "", -1)


# ──────────────────────────── вёрстка ────────────────────────────


def make_L01(tmp: Path) -> Case:
    p, dna = _one_slide_deck(tmp, "L01", lambda b, s: b.picture(s, SW - Inches(1), Inches(2), Inches(3), Inches(2.25)))
    return Case(p, dna, None, "L01", 0)


def make_L02(tmp: Path) -> Case:
    def fill(b, s):
        b.text(s, MARGIN, Inches(2), Inches(6), Inches(2), "Первый блок текста, довольно длинный")
        b.text(s, Inches(3), Inches(3), Inches(6), Inches(2), "Второй блок текста поверх первого")
    p, dna = _one_slide_deck(tmp, "L02", fill)
    return Case(p, dna, None, "L02", 0)


def make_L03(tmp: Path) -> Case:
    words = " ".join(["согласование"] * 40)
    p, dna = _one_slide_deck(tmp, "L03", lambda b, s: b.text(s, MARGIN, Inches(2), Inches(3), Inches(0.5), words))
    return Case(p, dna, None, "L03", 0)


def make_L04(tmp: Path) -> Case:
    p, dna = _one_slide_deck(tmp, "L04", lambda b, s: b.text(s, SW - Inches(2), Inches(2), Inches(4), Inches(1), "Текст у края"))
    return Case(p, dna, None, "L04", 0)


def make_L05(tmp: Path) -> Case:
    p, dna = _one_slide_deck(tmp, "L05", lambda b, s: b.text(s, Inches(3.3), Inches(2), Inches(4), Inches(1), "Блок мимо сетки"))
    return Case(p, dna, None, "L05", 0)


def make_L06(tmp: Path) -> Case:
    p, dna = _one_slide_deck(tmp, "L06", lambda b, s: b.text(s, Inches(0.1), Inches(2), Inches(4), Inches(1), "Блок в поле"))
    return Case(p, dna, None, "L06", 0)


def make_L07(tmp: Path) -> Case:
    p, dna = _one_slide_deck(tmp, "L07", lambda b, s: b.picture(s, MARGIN, Inches(2), Inches(6), Inches(2)))
    return Case(p, dna, None, "L07", 0)


# ──────────────────────────── шаблон ────────────────────────────


def make_T01(tmp: Path) -> Case:
    p, dna = _one_slide_deck(tmp, "T01", lambda b, s: b.text(s, MARGIN, Inches(2), Inches(6), Inches(1), "Чужой шрифт",
                                                            font="Comic Sans MS"))
    return Case(p, dna, None, "T01", 0)


def make_T02(tmp: Path) -> Case:
    p, dna = _one_slide_deck(tmp, "T02", lambda b, s: b.text(s, MARGIN, Inches(2), Inches(6), Inches(1), "Кегль 21", size=21))
    return Case(p, dna, None, "T02", 0)


def make_T03(tmp: Path) -> Case:
    p, dna = _one_slide_deck(tmp, "T03", lambda b, s: b.text(s, MARGIN, Inches(2), Inches(6), Inches(1), "Чужой цвет",
                                                            color="FF00AA"))
    return Case(p, dna, None, "T03", 0)


def make_T04(tmp: Path) -> Case:
    p, dna = _one_slide_deck(tmp, "T04", lambda b, s: None)
    _rename_layouts(p, "Чужой макет")
    return Case(p, dna, None, "T04", 0)


def _rename_layouts(pptx: Path, new_name: str) -> None:
    """Переименовать все лейауты в файле (p:cSld/@name) — колода «не на макете шаблона»."""
    src = zipfile.ZipFile(pptx)
    items = [(i, src.read(i.filename)) for i in src.infolist()]
    src.close()
    with zipfile.ZipFile(pptx, "w", zipfile.ZIP_DEFLATED) as dst:
        for info, data in items:
            if re.fullmatch(r"ppt/slideLayouts/slideLayout\d+\.xml", info.filename):
                root = etree.fromstring(data)
                csld = root.find("p:cSld", NS)
                if csld is not None:
                    csld.set("name", new_name)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            dst.writestr(info.filename, data)


def make_T05(tmp: Path) -> Case:
    p, dna = _one_slide_deck(tmp, "T05", lambda b, s: b.logo(s, x=Inches(1)))  # логотип не на своём месте
    return Case(p, dna, None, "T05", 0)


def make_T06(tmp: Path) -> Case:
    p, dna = _one_slide_deck(tmp, "T06", lambda b, s: b.text(s, MARGIN, Inches(2), Inches(6), Inches(1), "Бледный текст",
                                                            color="BBBBBB"))
    return Case(p, dna, None, "T06", 0)


# ──────────────────────────── плотность ────────────────────────────


def make_D01(tmp: Path) -> Case:
    items = [f"Пункт номер {i + 1} списка" for i in range(8)]
    p, dna = _one_slide_deck(tmp, "D01", lambda b, s: b.text(s, MARGIN, Inches(2), Inches(8), Inches(4.5), items, bullets=True))
    return Case(p, dna, None, "D01", 0)


def make_D02(tmp: Path) -> Case:
    long = " ".join(f"слово{i}" for i in range(20))
    p, dna = _one_slide_deck(tmp, "D02", lambda b, s: b.text(s, MARGIN, Inches(2), Inches(8), Inches(3), [long, "Короткий"],
                                                            bullets=True))
    return Case(p, dna, None, "D02", 0)


def make_D03(tmp: Path) -> Case:
    p, dna = _one_slide_deck(tmp, "D03", lambda b, s: b.table(s, MARGIN, Inches(2), Inches(8), Inches(4.5), rows=11, cols=3))
    return Case(p, dna, None, "D03", 0)


def make_D04(tmp: Path) -> Case:
    p, dna = _one_slide_deck(tmp, "D04", lambda b, s: b.chart(s, MARGIN, Inches(2), Inches(8), Inches(4.5), n_series=7))
    return Case(p, dna, None, "D04", 0)


def make_D05(tmp: Path) -> Case:
    def fill(b, s):  # три блока (не «разреженный»), но контент занимает < 25 %
        b.text(s, MARGIN, Inches(2), Inches(2), Inches(0.4), "мало", size=12)
        b.text(s, COL2, Inches(2), Inches(2), Inches(0.4), "текста", size=12)
    p, dna = _one_slide_deck(tmp, "D05", fill)
    return Case(p, dna, None, "D05", 0)


# ──────────────────────────── целостность ────────────────────────────


def make_I01(tmp: Path) -> Case:
    p = tmp / "I01.pptx"
    p.write_bytes(b"this is not a pptx at all")
    return Case(p, fixture_dna(make_clean(tmp).pptx), None, "I01", -1)


def make_I02(tmp: Path) -> Case:
    p, dna = _one_slide_deck(tmp, "I02", lambda b, s: b.text(s, MARGIN, Inches(2), Inches(6), Inches(1),
                                                            "Lorem ipsum dolor sit amet"))
    return Case(p, dna, None, "I02", 0)


def make_I03(tmp: Path) -> Case:
    clean = make_clean(tmp).pptx
    b = _builder(tmp)
    s = b.slide(layout=5)  # Title Only: настоящий плейсхолдер заголовка
    s.shapes.title.text = "Только заголовок"
    b.slide()  # и совсем пустой
    return Case(b.save(tmp / "I03.pptx"), fixture_dna(clean), None, "I03", 0)


def make_I04(tmp: Path) -> Case:
    clean = make_clean(tmp).pptx
    b = _builder(tmp)
    s = b.slide()
    b.picture(s, 0, 0, SW, SH, px=(1600, 900))
    return Case(b.save(tmp / "I04.pptx"), fixture_dna(clean), None, "I04", 0)


def make_I05(tmp: Path) -> Case:
    p, dna = _one_slide_deck(tmp, "I05", lambda b, s: b.chart(s, MARGIN, Inches(2), Inches(8), Inches(4.5), legend=False,
                                                             axis_title=False))
    return Case(p, dna, None, "I05", 0)


def make_I06(tmp: Path) -> Case:
    clean = make_clean(tmp).pptx
    b = _builder(tmp)
    for _ in range(2):
        s = b.slide()
        b.title(s, "Одинаковый заголовок на двух слайдах")
        b.text(s, MARGIN, Inches(2), Inches(8), Inches(3), ["Первый пункт про статусы", "Второй пункт про перегрузку"],
               bullets=True)
    return Case(b.save(tmp / "I06.pptx"), fixture_dna(clean), None, "I06", 1)


FIXTURES: dict[str, Callable[[Path], Case]] = {
    name[5:]: fn for name, fn in globals().items() if name.startswith("make_") and name != "make_clean"
}

__all__ = ["Case", "FIXTURES", "DeckBuilder", "fixture_dna", "make_clean"]
