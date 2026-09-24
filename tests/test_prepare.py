"""Подготовка шаблона: .potx / .pptm и шаблон без слайдов приводятся к .pptx; VLM-разметка — в кэш, API и UI."""

from __future__ import annotations

import importlib
import zipfile
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

from deckforge.core.ir import DeckOutline
from deckforge.parsing import exemplars
from deckforge.parsing.normalize import PRESENTATION_MAIN, needs_normalize, normalize_template
from deckforge.pipeline import RunConfig, prepare_template, run
from tests.conftest import SAMPLE_OUTLINE, FakeClient

REPO = Path(__file__).resolve().parents[1]
run_mod = importlib.import_module("deckforge.pipeline.run")  # в пакете имена run и prepare_template — функции
prepare_mod = importlib.import_module("deckforge.pipeline.prepare")
TEMPLATE_MAIN = "application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
MACRO_MAIN = "application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml"
TAGGER = {"archetype": "cards", "confidence": 0.9, "tags": ["vlm"]}


def _repack(src: Path, dest: Path, content_type: str, vba: bool = False) -> Path:
    """Копия пакета с другим типом основной части; `vba` — плюс vbaProject.bin со связью, как у .pptm."""
    with zipfile.ZipFile(src) as z, zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as out:
        for info in z.infolist():
            data = z.read(info.filename)
            if info.filename == "[Content_Types].xml":
                data = data.replace(PRESENTATION_MAIN.encode(), content_type.encode())
                if vba:
                    data = data.replace(b"</Types>", b'<Override PartName="/ppt/vbaProject.bin" '
                                        b'ContentType="application/vnd.ms-office.vbaProject"/></Types>')
            elif vba and info.filename == "ppt/_rels/presentation.xml.rels":
                data = data.replace(b"</Relationships>", b'<Relationship Id="rIdVba" Type="http://schemas.'
                                    b'microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin"/>'
                                    b"</Relationships>")
            out.writestr(info, data)
        if vba:
            out.writestr("ppt/vbaProject.bin", b"macro")
    return dest


def _without_slides(src: Path, dest: Path) -> Path:
    """Шаблон, как обычный .potx: мастера и лейауты есть, слайдов нет."""
    prs = Presentation(str(src))
    lst = prs.slides._sldIdLst
    for sid in list(lst):
        prs.part.drop_rel(sid.rId)
        lst.remove(sid)
    prs.save(str(dest))
    return dest


@pytest.fixture
def prepared_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    d = tmp_path / "prepared"
    monkeypatch.setattr(run_mod, "PREPARED_DIR", d)
    return d


@pytest.fixture
def vlm_cache(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """Кэш разметки во временной папке; бандл data/archetypes не подмешивается."""
    d = tmp_path / "archetypes"
    monkeypatch.setattr(exemplars, "ARCHETYPES_CACHE", d)
    monkeypatch.setattr(exemplars, "ARCHETYPES_BUNDLED", tmp_path / "bundled")
    monkeypatch.setattr(prepare_mod, "RENDER_DIR", tmp_path / "render")
    return d


def _fake_pngs(out_dir: Path, n: int) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    for i in range(1, n + 1):
        Image.new("RGB", (32, 18), "white").save(out_dir / f"slide_{i:02d}.png")


def test_potx_is_parsed_like_pptx(template_path, tmp_path: Path, prepared_dir: Path) -> None:
    src = template_path("ЛЦТ2026")
    potx = _repack(src, tmp_path / "шаблон.potx", TEMPLATE_MAIN)
    assert ".potx" in needs_normalize(potx)
    with pytest.raises(ValueError):
        Presentation(str(potx))  # сам python-pptx .potx не открывает

    parsed = run_mod.parse_template(potx)
    assert parsed.template.parent == prepared_dir and parsed.template.suffix == ".pptx"
    assert parsed.meta["source"] == "шаблон.potx" and ".potx" in parsed.meta["normalized"]
    assert needs_normalize(parsed.template) is None
    Presentation(str(parsed.template))
    assert len(parsed.exemplars) == len(run_mod.parse_template(src).exemplars)
    mtime = parsed.template.stat().st_mtime_ns
    assert run_mod.parse_template(potx).template.stat().st_mtime_ns == mtime  # второй разбор — из кэша


def test_pptx_with_slides_is_untouched(template_path, prepared_dir: Path) -> None:
    src = template_path("ЛЦТ2026")
    assert needs_normalize(src) is None
    parsed = run_mod.parse_template(src)
    assert parsed.template == src and "normalized" not in parsed.meta
    assert not prepared_dir.exists()


def test_macro_parts_are_dropped(template_path, tmp_path: Path) -> None:
    pptm = _repack(template_path("ЛЦТ2026"), tmp_path / "m.pptm", MACRO_MAIN, vba=True)
    assert ".pptm" in needs_normalize(pptm)
    out = normalize_template(pptm, tmp_path / "m.pptx")
    with zipfile.ZipFile(out) as z:
        assert "ppt/vbaProject.bin" not in z.namelist()
        assert b"vbaProject" not in z.read("[Content_Types].xml")
        assert b"vbaProject" not in z.read("ppt/_rels/presentation.xml.rels")
        assert PRESENTATION_MAIN.encode() in z.read("[Content_Types].xml")
    assert needs_normalize(out) is None
    Presentation(str(out))


def test_template_without_slides_gets_exemplars_from_layouts(template_path, tmp_path: Path,
                                                             prepared_dir: Path) -> None:
    """Обычный .potx — только лейауты: образцы из них, колода собирается, а не выходит пустой."""
    empty = _repack(_without_slides(template_path("ЛЦТ2026"), tmp_path / "e.pptx"), tmp_path / "e.potx",
                    TEMPLATE_MAIN)
    assert needs_normalize(empty) is not None
    cfg = RunConfig(template=empty, strategies=["executive"], output_dir=tmp_path / "run", images="off",
                    export=["pptx"], audit={"deterministic": True, "contextual": False, "autofix": True})
    res = run(cfg, outline=DeckOutline.model_validate_json(SAMPLE_OUTLINE.read_text("utf-8")))
    assert len(res.parsed.exemplars) >= 10
    assert {e.archetype.value for e in res.parsed.exemplars} >= {"title", "bullets", "cards"}
    assert any("нет слайдов" in w for w in res.warnings)
    deck = res.decks[0]
    assert deck.stats["slides"] >= 8 and len(Presentation(str(deck.pptx)).slides) == deck.stats["slides"]
    assert deck.load_manifest()["template"]["normalized"]


def test_prepare_template_caches_vlm_answers(template_path, tmp_path: Path, vlm_cache: Path) -> None:
    src = tmp_path / "lct.pptx"
    src.write_bytes(template_path("ЛЦТ2026").read_bytes())  # своё имя: не пересечься с кэшем датасета
    n = len(Presentation(str(src)).slides)
    _fake_pngs(prepare_mod.RENDER_DIR / src.stem, n)
    lines: list[str] = []
    report = prepare_template(src, FakeClient(by_skill={"template_tagger": [TAGGER]}), progress=lines.append,
                              do_render=False)
    assert report.thumbnails == n and report.vlm_calls == report.ambiguous > 0 and report.vlm_errors == 0
    assert report.cache == vlm_cache / "lct.json" and report.cache.exists()
    assert report.vlm_changed and lines and report.summary()["vlm_calls"] == report.vlm_calls

    profiles = exemplars.load_profiles(src)  # следующий разбор берёт ответы из кэша
    vlm = [p for p in profiles if p.source == "vlm"]
    assert len(vlm) == report.ambiguous and all(p.archetype.value == "cards" for p in vlm)
    assert exemplars.has_vlm_labels(src) and run_mod.parse_template(src).meta["labels"] == "rules+vlm"


def test_prepare_without_vlm_writes_no_cache(template_path, tmp_path: Path, vlm_cache: Path) -> None:
    src = tmp_path / "rules.pptx"
    src.write_bytes(template_path("ЛЦТ2026").read_bytes())
    report = prepare_template(src, None, do_render=False)
    assert report.cache is None and report.vlm_calls == 0 and report.profiles
    assert run_mod.parse_template(src).meta["labels"] == "rules"


def test_api_prepare_and_potx_upload(template_path, tmp_path: Path, vlm_cache: Path, prepared_dir: Path) -> None:
    from fastapi.testclient import TestClient

    from deckforge.api.app import create_app

    src = template_path("ЛЦТ2026")
    app = create_app(root=tmp_path / "api", executor=None,
                     client_factory=lambda: FakeClient(by_skill={"template_tagger": [TAGGER]}))
    with TestClient(app) as api:
        potx = _repack(src, tmp_path / "свой.potx", TEMPLATE_MAIN)
        r = api.post("/templates", files={"file": ("свой.potx", potx.read_bytes())})
        assert r.status_code == 200, r.text
        t = r.json()
        assert t["path"].endswith(".pptx") and t["summary"]["slides"] > 0
        _fake_pngs(prepare_mod.RENDER_DIR / Path(t["path"]).stem, len(Presentation(t["path"]).slides))

        r = api.post(f"/templates/{t['id']}/prepare")
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["report"]["vlm_calls"] > 0 and body["report"]["cache"]
        assert body["template"]["summary"]["slides"] == t["summary"]["slides"]
        assert api.post("/templates/nope/prepare").status_code == 404


@pytest.mark.filterwarnings("ignore::DeprecationWarning")
def test_ui_prepare_button_only_for_uploads(template_path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import streamlit as st
    from streamlit.testing.v1 import AppTest

    src = template_path("ЛЦТ2026")
    up = tmp_path / "templates" / "abc123abc123" / "свой__abc123abc123.pptx"
    up.parent.mkdir(parents=True)
    up.write_bytes(src.read_bytes())
    monkeypatch.setenv("DECKFORGE_UI_ROOT", str(tmp_path))
    st.cache_resource.clear()  # реестр шаблонов — cache_resource: не взять корень от прошлого теста
    at = AppTest.from_file(str(REPO / "deckforge" / "ui" / "app.py"), default_timeout=120).run()
    assert not at.exception, at.exception
    assert not [b for b in at.button if b.label.startswith("Уточнить разметку VLM")]  # датасет размечен заранее
    at.sidebar.selectbox(key="template_choice").set_value("📤 свой").run()
    assert not at.exception, at.exception
    assert [b for b in at.button if b.label.startswith("Уточнить разметку VLM")]
    assert any(c.value.startswith("Разметка образцов: правила по геометрии") for c in at.caption)
    st.cache_resource.clear()
