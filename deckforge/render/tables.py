"""Нативные таблицы python-pptx из TableSpec: шапка в цвете accent, тело без Office-стиля по умолчанию.

style_overrides: font, size_pt, text_color, accent, header_fill, header_text_color.
"""

from __future__ import annotations

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.slide import Slide
from pptx.util import Pt

from deckforge.core.ir import Box, TableSpec
from deckforge.core.ooxml import A, NS

NO_STYLE_NO_GRID = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
ROW_LINES = 2  # запас высоты строки в строках текста
LINE_SPACING = 1.2


def _rgb(value: object, default: str) -> RGBColor:
    return RGBColor.from_string(str(value or default).lstrip("#").upper())


def add_table(slide: Slide, spec: TableSpec, box: Box, overrides: dict | None = None):
    overrides = overrides or {}
    rows, cols = 1 + len(spec.rows), len(spec.header)
    frame = slide.shapes.add_table(rows, cols, box.x, box.y, box.w, box.h)
    table = frame.table
    tbl = frame._element.find(".//a:tbl", NS)
    tbl_pr = tbl.find("a:tblPr", NS)
    style_id = tbl_pr.find("a:tableStyleId", NS)
    if style_id is None:
        style_id = etree.SubElement(tbl_pr, A + "tableStyleId")
    style_id.text = NO_STYLE_NO_GRID
    table.first_row = True
    table.horz_banding = False

    font_name = str(overrides["font"]) if overrides.get("font") else None
    size = Pt(float(overrides.get("size_pt") or 12))
    header_fill = _rgb(overrides.get("header_fill") or overrides.get("accent"), "0077FF")
    header_text = _rgb(overrides.get("header_text_color"), "FFFFFF")
    body_text = _rgb(overrides.get("text_color"), "212121")
    # высота строки по кеглю с запасом на перенос, а не по боксу: редактор сам увеличит строку при нужде
    row_h = min(box.h // rows, int(size * ROW_LINES * LINE_SPACING) + Pt(6))

    def fill_cell(cell, text: str, bold: bool, color: RGBColor) -> None:
        cell.text = text
        cell.margin_left = cell.margin_right = Pt(6)
        cell.margin_top = cell.margin_bottom = Pt(3)
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.size = size
                r.font.bold = bold
                r.font.color.rgb = color
                if font_name:
                    r.font.name = font_name

    for c, h in enumerate(spec.header):
        cell = table.cell(0, c)
        fill_cell(cell, h, True, header_text)
        cell.fill.solid()
        cell.fill.fore_color.rgb = header_fill
    for r, row in enumerate(spec.rows, start=1):
        for c in range(cols):
            cell = table.cell(r, c)
            fill_cell(cell, row[c] if c < len(row) else "", False, body_text)
            cell.fill.background()
    for r in range(rows):
        table.rows[r].height = row_h
    return frame
