"""Контент-пакет: фрагменты с идентификаторами из brief.md / *.md|txt|docx|pdf / data/*.json|csv|xlsx."""

import zipfile
from pathlib import Path

import openpyxl
import pymupdf

from deckforge.content import load_content_pack

REPO = Path(__file__).resolve().parents[1]
PACK = REPO / "examples" / "content_pack"


def test_example_pack_fragments() -> None:
    pack = load_content_pack(PACK)
    by_id = {f.id: f for f in pack.fragments}
    assert "brief" in by_id and by_id["brief"].kind == "text" and "Пульс команды" in pack.brief
    assert by_id["m:coordination_hours"].kind == "series"
    assert by_id["m:time_structure"].kind == "table"
    assert by_id["m:quote"].kind == "quote"
    assert by_id["m:overload_share"].kind == "metric"
    assert by_id["m:response_days"].kind == "metric"  # before/after без value
    assert "doc:product" in by_id and any(i.startswith("doc:product:") for i in by_id)
    assert any(i.startswith("brief:") for i in by_id)  # секции брифа по **жирным** заголовкам


def test_prompt_text_has_all_data_ids_but_not_brief() -> None:
    pack = load_content_pack(PACK)
    text, warnings = pack.to_prompt_text()
    assert not warnings
    for frag in pack.fragments:
        if frag.id == "brief" or frag.id.startswith("brief:"):
            assert f"[{frag.id}]" not in text  # бриф идёт отдельным входом скилла
        else:
            assert f"[{frag.id}]" in text
    assert '"series"' in text and "Пилотные команды" in text


def test_prompt_text_truncates_with_warning() -> None:
    pack = load_content_pack(PACK)
    text, warnings = pack.to_prompt_text(limit=300)
    assert len(text) <= 300 and warnings and "обрезан" in warnings[0]


def test_csv_and_plain_md(tmp_path: Path) -> None:
    (tmp_path / "brief.md").write_text("# Бриф\n\nТекст брифа.\n", "utf-8")
    (tmp_path / "notes.md").write_text("# Заметки\n\n## Раздел A\n\nодин\n\n## Раздел B\n\nдва\n", "utf-8")
    (tmp_path / "data").mkdir()
    rows = "\n".join(["месяц,значение"] + [f"m{i},{i}" for i in range(40)])
    (tmp_path / "data" / "sales.csv").write_text(rows, "utf-8")
    pack = load_content_pack(tmp_path)
    ids = pack.ids()
    assert {"brief", "doc:notes", "doc:notes:раздел-a", "doc:notes:раздел-b", "csv:sales"} <= ids
    csv_frag = pack.get("csv:sales")
    assert csv_frag and csv_frag.kind == "table" and csv_frag.data["header"] == ["месяц", "значение"]
    assert len(csv_frag.data["rows"]) == 30 and "первые 30" in csv_frag.text
    assert pack.get("brief").text.startswith("# Бриф")


_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _docx(path: Path, paragraphs: list[tuple[str, str]], table: list[list[str]] | None = None) -> None:
    """Минимальный .docx: абзацы (styleId, текст) и одна таблица; без python-docx."""
    def p(style: str, text: str) -> str:
        ppr = f'<w:pPr><w:pStyle w:val="{style}"/></w:pPr>' if style else ""
        return f"<w:p>{ppr}<w:r><w:t>{text}</w:t></w:r></w:p>"

    body = "".join(p(s, t) for s, t in paragraphs)
    if table:
        body += "<w:tbl>" + "".join(
            "<w:tr>" + "".join(f"<w:tc><w:p><w:r><w:t>{c}</w:t></w:r></w:p></w:tc>" for c in row) + "</w:tr>"
            for row in table) + "</w:tbl>"
    doc = f'<?xml version="1.0" encoding="UTF-8"?><w:document xmlns:w="{_W}"><w:body>{body}</w:body></w:document>'
    types = ('<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
             '<Default Extension="xml" ContentType="application/xml"/>'
             '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("[Content_Types].xml", types)
        z.writestr("word/document.xml", doc)


def test_docx_headings_and_tables(tmp_path: Path) -> None:
    (tmp_path / "brief.md").write_text("# Бриф\n\nТекст.\n", "utf-8")
    _docx(tmp_path / "spec.docx", [("Title", "Спецификация"), ("Heading1", "Проблема"), ("", "Текст проблемы."),
                                   ("Heading1", "Решение"), ("", "Текст решения.")],
          table=[["Метрика", "До", "После"], ["Согласования", "33%", "17%"]])
    pack = load_content_pack(tmp_path)
    assert not pack.warnings
    assert {"doc:spec", "doc:spec:проблема", "doc:spec:решение", "doc:spec:table-1"} <= pack.ids()
    spec = pack.get("doc:spec")
    assert spec.title == "Спецификация" and "## Проблема" in spec.text and "Текст решения." in spec.text
    assert pack.get("doc:spec:проблема").text == "Текст проблемы."
    t = pack.get("doc:spec:table-1")
    assert t.kind == "table" and t.data == {"header": ["Метрика", "До", "После"], "rows": [["Согласования", "33%", "17%"]]}
    assert "[doc:spec:table-1]" in pack.to_prompt_text()[0]


def test_xlsx_sheets(tmp_path: Path) -> None:
    (tmp_path / "brief.md").write_text("# Бриф\n", "utf-8")
    (tmp_path / "data").mkdir()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Продажи"
    ws.append(["месяц", "значение", None, None])  # хвост пустых колонок
    for i in range(40):
        ws.append([f"m{i}", float(i)])
    wb.create_sheet("Пустой")
    wb.save(tmp_path / "data" / "sales.xlsx")
    pack = load_content_pack(tmp_path)
    frag = pack.get("xlsx:sales")  # единственный непустой лист → без суффикса листа
    assert frag and frag.kind == "table" and frag.title == "Продажи"
    assert frag.data["header"] == ["месяц", "значение"] and frag.data["rows"][3] == ["m3", "3"]
    assert len(frag.data["rows"]) == 30 and "первые 30 из 40" in frag.text

    wb = openpyxl.Workbook()
    wb.active.append(["a", "b"])
    wb.active.append([1, 2])
    wb.create_sheet("Риски").append(["риск", "мера"])
    wb.save(tmp_path / "data" / "two.xlsx")
    ids = load_content_pack(tmp_path).ids()
    assert {"xlsx:two:sheet", "xlsx:two:риски"} <= ids


def test_pdf_pages_and_txt_sections(tmp_path: Path) -> None:
    (tmp_path / "brief.md").write_text("# Бриф\n", "utf-8")
    with pymupdf.open() as doc:
        for n in (1, 2):
            doc.new_page().insert_text((72, 72), f"Report page {n}\nfact {n}")
        doc.save(tmp_path / "report.pdf")
    (tmp_path / "notes.txt").write_text("**Раздел A.** один\n\n**Раздел B.** два\n", "utf-8")
    pack = load_content_pack(tmp_path)
    assert {"doc:report", "doc:report:p1", "doc:report:p2", "doc:notes", "doc:notes:раздел-a"} <= pack.ids()
    assert pack.get("doc:report").title == "Report page 1" and "fact 2" in pack.get("doc:report:p2").text


def test_broken_file_is_skipped_with_warning(tmp_path: Path) -> None:
    (tmp_path / "brief.md").write_text("# Бриф\n", "utf-8")
    (tmp_path / "bad.docx").write_bytes(b"not a zip")
    pack = load_content_pack(tmp_path)
    assert pack.ids() == {"brief"} and pack.warnings and pack.warnings[0].startswith("bad.docx: не прочитан")
    assert pack.to_prompt_text()[1] == pack.warnings
