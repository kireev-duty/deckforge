"""Схема из нативных автофигур по DiagramSpec — замена SmartArt, когда в шаблоне нет process-образца.

Шаги слева направо: первый — «стрелка-табличка» (homePlate), остальные — шевроны, с зазором (без наложений,
L02); в фигуре номер шага, под ней подпись. Всё одной группой: в PowerPoint схема двигается целиком,
каждая фигура и подпись редактируются. Цвета — палитра шаблона, шрифт — шрифт шаблона.

style_overrides: font, size_pt, text_color, accent, palette ("FF0053,520977,…").
"""

from __future__ import annotations

from collections.abc import Callable

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.slide import Slide
from pptx.util import Emu, Pt

from deckforge.core.colors import contrast_ratio
from deckforge.core.ir import Box, DiagramSpec
from deckforge.render.charts import palette_from

GAP_SHARE = 0.02  # зазор между фигурами, доля ширины бокса
CHEVRON_H_SHARE = 0.32  # высота фигуры шага — доля высоты бокса…
CHEVRON_H_MAX = Emu(914400)  # …но не выше дюйма
CHEVRON_ASPECT = 0.45  # …и не выше доли своей ширины: у узкого шеврона остриё съедает место под номер
CAPTION_GAP = Emu(91440)  # 0,1" между фигурой и подписью
NUMBER_SHARE = 0.45  # кегль номера — доля высоты фигуры
CAPTION_MAX_PT, CAPTION_MIN_PT = 18.0, 10.0
CHAR_WIDTH = 0.62  # ширина знака в долях кегля: с запасом, иначе длинное слово рвётся («Масштабировани/е»)
TOP_SHARE = 0.4  # свободная высота бокса: 40 % над схемой, 60 % под ней — чуть выше центра
LINE_SPACING = 1.2
EMU_PER_PT = 12700
LIGHT, DARK = "FFFFFF", "212121"


def caption_size(texts: list[str], w: int, h: int, base_pt: float | None = None) -> float:
    """Наибольший кегль из [CAPTION_MIN_PT, CAPTION_MAX_PT], при котором самая длинная подпись влезает в w×h."""
    size = min(CAPTION_MAX_PT, base_pt or CAPTION_MAX_PT)
    while size > CAPTION_MIN_PT:
        cpl = max(1, int(w / EMU_PER_PT / (CHAR_WIDTH * size)))
        lines = max(1, int(h / EMU_PER_PT / (LINE_SPACING * size)))
        if all(_lines_needed(t, cpl) <= lines and max(map(len, t.split()), default=0) <= cpl for t in texts):
            return size
        size -= 1.0
    return CAPTION_MIN_PT


def _lines_needed(text: str, cpl: int) -> int:
    """Строк при переносе по словам."""
    lines, cur = 1, 0
    for word in text.split():
        add = len(word) + (1 if cur else 0)
        if cur and cur + add > cpl:
            lines, cur = lines + 1, len(word)
        else:
            cur += add
    return lines


def on_fill(fill: str, dark: str = DARK) -> str:
    """Цвет текста на заливке фигуры: белый или тёмный текст палитры — что контрастнее."""
    return LIGHT if contrast_ratio(LIGHT, fill) >= contrast_ratio(dark, fill) else dark


def snap_size(size: float, scale: list[float], floor: float) -> float:
    """Ближайший не больший кегль из шкалы шаблона (не ниже floor); шкалы нет — как есть."""
    fits = [s for s in scale if floor <= s <= size]
    return max(fits) if fits else size


def _scale(overrides: dict) -> list[float]:
    out = []
    for v in str(overrides.get("type_scale") or "").split(","):
        try:
            out.append(float(v))
        except ValueError:
            continue
    return sorted(out)


def add_diagram(slide: Slide, spec: DiagramSpec, box: Box, overrides: dict | None = None,
                caption_color: Callable[[Box], str | None] | None = None):
    """Нарисовать схему в боксе; возвращает группу.

    `caption_color(бокс подписи)` — цвет подписи по тому, что под ней (белая карточка шаблона под одной
    подписью, тёмный фон под другой); None — тёмный текст палитры. Без функции — `text_color` для всех."""
    overrides = overrides or {}
    n = len(spec.items)
    if n == 0:
        return None
    font = str(overrides["font"]) if overrides.get("font") else None
    text_color = str(overrides.get("text_color") or DARK).lstrip("#").upper()
    # тёмный текст на светлой фигуре — тёмный цвет палитры (подписи на тёмном фоне могли стать белыми)
    dark = str(overrides.get("palette_text") or overrides.get("text_color") or DARK).lstrip("#").upper()
    if dark == LIGHT:
        dark = DARK
    scale = _scale(overrides)
    colors = palette_from(overrides, n)
    gap = int(box.w * GAP_SHARE)
    w = (box.w - gap * (n - 1)) // n
    h = int(min(box.h * CHEVRON_H_SHARE, CHEVRON_H_MAX, w * CHEVRON_ASPECT))
    room = max(0, box.h - h - CAPTION_GAP)
    cap_pt = caption_size(spec.items, w, room, float(overrides["size_pt"]) if overrides.get("size_pt") else None)
    cap_pt = snap_size(cap_pt, scale, CAPTION_MIN_PT)
    # подписи — по нужной высоте, а не до низа бокса; схема целиком — чуть выше центра области
    cpl = max(1, int(w / EMU_PER_PT / (CHAR_WIDTH * cap_pt)))
    lines = max(_lines_needed(t, cpl) for t in spec.items)
    cap_h = min(room, int((lines + 0.5) * LINE_SPACING * cap_pt * EMU_PER_PT))
    top = box.y + int((room - cap_h) * TOP_SHARE)
    cap_y = top + h + CAPTION_GAP
    num_pt = snap_size(max(12.0, round(h / EMU_PER_PT * NUMBER_SHARE)), scale, 12.0)

    group = slide.shapes.add_group_shape()
    group.name = "deckforge-diagram"
    for i, text in enumerate(spec.items):
        x = box.x + i * (w + gap)
        fill = colors[i % len(colors)].lstrip("#").upper()
        shape = group.shapes.add_shape(MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON, x, top, w, h)
        shape.name = f"Шаг {i + 1}"
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        shape.line.fill.background()
        shape.shadow.inherit = False
        tf = shape.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.word_wrap = False
        para = tf.paragraphs[0]
        para.alignment = PP_ALIGN.CENTER
        run = para.add_run()
        run.text = f"{i + 1:02d}"
        _style(run.font, font, num_pt, on_fill(fill, dark), bold=True)

        cap = group.shapes.add_textbox(x, cap_y, w, cap_h)
        cap.name = f"Подпись шага {i + 1}"
        ctf = cap.text_frame
        ctf.word_wrap = True
        ctf.vertical_anchor = MSO_ANCHOR.TOP
        for attr in ("margin_left", "margin_right"):
            setattr(ctf, attr, Emu(0))
        cpara = ctf.paragraphs[0]
        cpara.alignment = PP_ALIGN.LEFT
        crun = cpara.add_run()
        crun.text = text
        color = text_color if caption_color is None else (caption_color(Box(x=x, y=cap_y, w=w, h=cap_h)) or dark)
        _style(crun.font, font, cap_pt, color.lstrip("#").upper())
    return group


def _style(f, font: str | None, size_pt: float, color: str, bold: bool = False) -> None:
    if font:
        f.name = font
    f.size = Pt(size_pt)
    f.bold = bold
    f.color.rgb = RGBColor.from_string(color)


__all__ = ["add_diagram", "caption_size", "on_fill"]
