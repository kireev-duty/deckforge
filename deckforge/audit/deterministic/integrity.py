"""Целостность: файл открывается, нет заглушек, пустых и растровых слайдов, подписи диаграмм, дубли (I01–I06)."""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

from pptx import Presentation

from deckforge.audit.context import SPARSE_ARCHETYPES, AuditContext
from deckforge.audit.deterministic.base import finding
from deckforge.core.ir import Finding, Severity, SlotKind
from deckforge.core.placeholders import PLACEHOLDER_PATTERNS, PLACEHOLDER_WHOLE

EMPTY_PH_TYPES = {"body", "obj", "subTitle", "title", "ctrTitle", "pic", "chart", "tbl"}
RASTER_MIN_AREA = 0.9
DUPLICATE_JACCARD = 0.8
DUPLICATE_MIN_WORDS = 5
_WORD = re.compile(r"[\w%]{2,}", re.UNICODE)


# ──────────────────────────── I01 ────────────────────────────


def probe_file(pptx: str | Path) -> Finding | None:
    """Файл не открывается: битый zip, дубли имён частей, python-pptx падает."""
    pptx = Path(pptx)
    try:
        with zipfile.ZipFile(pptx) as z:
            names = z.namelist()
            bad = z.testzip()
        if bad:
            return _i01(f"Повреждённая часть архива: {bad}")
        if len(names) != len(set(names)):
            dup = sorted({n for n in names if names.count(n) > 1})[:3]
            return _i01(f"В архиве дублируются имена частей: {', '.join(dup)}")
        Presentation(str(pptx))
    except Exception as e:  # noqa: BLE001 — любая ошибка и есть находка
        return _i01(f"Файл не открывается: {type(e).__name__}: {str(e)[:120]}")
    return None


def _i01(msg: str) -> Finding:
    return Finding(check_id="I01_file", kind="deterministic", severity=Severity.ERROR, slide_idx=-1, message=msg)


def check_I01(ctx: AuditContext) -> list[Finding]:
    f = probe_file(ctx.pptx)
    return [f] if f else []


# ──────────────────────────── I02 ────────────────────────────


def check_I02(ctx: AuditContext) -> list[Finding]:
    """Остался текст-заглушка."""
    out: list[Finding] = []
    for slide in ctx.slides:
        for sh in slide.shapes:
            if not sh.has_text:
                if sh.is_placeholder and sh.ph_type in EMPTY_PH_TYPES and sh.tag == "sp":
                    # пустой плейсхолдер в редакторе показывает подсказку лейаута
                    out.append(finding("I02_placeholder_text", slide, Severity.WARNING,
                                       f"Пустой плейсхолдер «{sh.ph_type}» — в редакторе покажется подсказка лейаута",
                                       sh, autofix="drop_shape", pattern="empty_placeholder"))
                continue
            text = sh.text
            hit = next((p.pattern for p in PLACEHOLDER_PATTERNS if p.search(text)), None)
            if hit is None and re.sub(r"\s+", " ", text).strip().lower() in PLACEHOLDER_WHOLE:
                hit = "whole"
            if hit is None:
                continue
            out.append(finding("I02_placeholder_text", slide, Severity.ERROR,
                               f"Текст-заглушка: «{text.strip()[:50]}»", sh, autofix="refill_slot", pattern=hit))
    return out


# ──────────────────────────── I03 ────────────────────────────


def check_I03(ctx: AuditContext) -> list[Finding]:
    """Пустой слайд или слайд только с заголовком."""
    out: list[Finding] = []
    for slide in ctx.slides:
        content = slide.content()
        if not content:
            out.append(finding("I03_empty_slide", slide, Severity.ERROR, "Пустой слайд", autofix="drop_slide", blocks=0))
            continue
        if slide.archetype in SPARSE_ARCHETYPES:
            continue
        if len(content) == 1 and content[0].has_text and _is_title(slide, content[0]):
            sev = Severity.ERROR if slide.ir is not None else Severity.WARNING
            out.append(finding("I03_empty_slide", slide, sev, "На слайде только заголовок", content[0],
                               autofix="drop_slide", blocks=1))
    return out


def _is_title(slide, sh) -> bool:
    if sh.ph_type in ("title", "ctrTitle"):
        return True
    slot = slide.slot_of(sh.id)
    return slot is not None and slot.kind == SlotKind.TITLE


# ──────────────────────────── I04 ────────────────────────────


def check_I04(ctx: AuditContext) -> list[Finding]:
    """Слайд — одна растровая картинка вместо редактируемых объектов."""
    out: list[Finding] = []
    area = ctx.slide_w * ctx.slide_h
    for slide in ctx.slides:
        if any(sh.has_text for sh in slide.shapes) or any(sh.chart or sh.table for sh in slide.shapes):
            continue
        big = [sh for sh in slide.shapes if sh.is_picture and sh.area >= RASTER_MIN_AREA * area]
        if big:
            out.append(finding("I04_raster_slide", slide, Severity.ERROR,
                               "Слайд состоит из одной картинки — нет редактируемых объектов", big[0],
                               coverage=round(big[0].area / area, 3)))
    return out


# ──────────────────────────── I05 ────────────────────────────


def check_I05(ctx: AuditContext) -> list[Finding]:
    """У диаграммы нет подписей осей / единиц / легенды."""
    out: list[Finding] = []
    for slide in ctx.slides:
        for sh in slide.shapes:
            c = sh.chart
            if c is None or c.kind == "unknown":
                continue
            problems = []
            if c.kind not in ("pieChart", "doughnutChart", "pie3DChart") and not c.val_axis_title and not c.has_data_labels:
                problems.append("нет подписи оси значений / единиц измерения и подписей данных")
            if c.n_series > 1 and not c.has_legend:
                problems.append(f"{c.n_series} серий без легенды")
            if problems:
                out.append(finding("I05_chart_labels", slide, Severity.WARNING, "Диаграмма: " + "; ".join(problems), sh,
                                   autofix="add_chart_labels", series=c.n_series, legend=int(c.has_legend),
                                   axis_title=int(c.val_axis_title)))
    return out


# ──────────────────────────── I06 ────────────────────────────


def check_I06(ctx: AuditContext) -> list[Finding]:
    """Два слайда дублируют друг друга (Jaccard по словам)."""
    out: list[Finding] = []
    words = [(s, {w.lower() for w in _WORD.findall(s.all_text)}) for s in ctx.slides]
    for i, (a, wa) in enumerate(words):
        if len(wa) < DUPLICATE_MIN_WORDS:
            continue
        for b, wb in words[i + 1:]:
            if len(wb) < DUPLICATE_MIN_WORDS:
                continue
            j = len(wa & wb) / len(wa | wb)
            if j > DUPLICATE_JACCARD:
                out.append(finding("I06_duplicate_slides", b, Severity.WARNING,
                                   f"Слайд повторяет слайд {a.idx + 1} (сходство {j:.0%})", autofix="drop_slide",
                                   other_slide=a.idx, jaccard=round(j, 3)))
    return out


__all__ = ["check_I01", "check_I02", "check_I03", "check_I04", "check_I05", "check_I06", "probe_file"]
